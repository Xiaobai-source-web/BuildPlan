# -*- coding: utf-8 -*-
r"""域 5 接线 + 节拍提示词死路径显式化 + 「未入树清单」交付物 —— 数据层回归测试。

本文件钉三件事（对应域 5 任务的三条）：

① **接线**（设计 §4 / §14.1 裁决 #2、#7）：主链 26 → 27 个节点，
   `quantity_fill` 必须夹在 `audit_wbs`（R1 门已过）与 `norm_bind`（定额分母要在量定稿后
   才选）**之间**；界面名 `PIPELINE_TITLES["quantity_fill"]` 与节点契约的 name/title
   逐字一致；`delivery.py` 里**不再**写死节点数叙述（数会变，写死就是下一处对不上）。

② **§14.1.1 静默死路径**（真缺陷）：`prompts/beat_config.txt` **不在仓库里**，
   `prompts_loader.load` 抛 `FileNotFoundError`（OSError 子类），原来被
   `except (LLMError, Exception): pass` 吞掉 ⇒ 「节拍展开的 LLM 细化路径 100% 不生效、
   量永远来自 BASE 基线」在终端 / 计划 JSON / 交付物里**一个字都不留**。
   现在 `except OSError` 走既有节点告警通道（`emit("warning")` → 引擎 collect →
   `ctx["node_warnings"]` → `meta.node_warnings`）；**其它异常保持静默降级**（行为不变）。
   ⚠️ 本文件**不建** `beat_config.txt`：那会突然激活一条沉睡的 LLM 路径。

③ **「未入树清单」**（§14.1 裁决 #6，硬要求）：`delivery.confidence_section_blocks` 新增的
   一节必须在**截断之前**给出 ① 总数 ② 按 L3 工种分布 ③ 一句人话解释；明细前 50 条 +
   其余指向 `meta.quantity_coverage.not_in_tree`；键不存在 ⇒ **整节优雅缺席**，
   不抛异常、不印 "None"（`add_kv` 只兜得住真 None，兜不住被 f-string 串化的 "None"）。

运行（**必须**带 `--basetemp`；本文件不落盘，只是仓库纪律）：
  cd backend; python -m pytest tests/test_domain5_wiring.py -q -p no:cacheprovider \
      --basetemp=_test_tmp\d5wire
"""

import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

import pipeline.builder as B                              # noqa: E402
import pipeline.nodes.beat_node as BN                      # noqa: E402
import pipeline.nodes.delivery as D                        # noqa: E402
from pipeline.base import BaseNode                         # noqa: E402
from pipeline.engine import NODE_WARNINGS_CTX_KEY, Pipeline  # noqa: E402
from pipeline.llm import LLMError                          # noqa: E402


# ══════════════════ ① 接线：主链 27 个节点 ══════════════════

class TestQuantityFillWiring:
    def test_主链27个节点且quantity_fill夹在audit_wbs与norm_bind之间(self):
        steps = B.pipeline_steps()
        names = [n for n, _t in steps]

        assert len(names) == 27, "主链节点数应为 27：%r" % (names,)
        assert "quantity_fill" in names, names
        i = names.index("quantity_fill")
        assert names[i - 1] == "audit_wbs", "补量必须在 R1 审计门之后：%r" % (names,)
        assert names[i + 1] == "norm_bind", "补量必须在定额锚定之前：%r" % (names,)

    def test_界面名覆盖表有quantity_fill且随run_plan下发(self):
        assert B.PIPELINE_TITLES.get("quantity_fill") == "补全各工序工程量"
        titles = dict(B.pipeline_steps())
        assert titles["quantity_fill"] == B.PIPELINE_TITLES["quantity_fill"], \
            "run_plan 下发的界面名必须取自这张表（唯一真源）"

    def test_节点契约的name与title与设计一致(self):
        cls = B.QuantityAgentNode
        assert cls.name == "quantity_fill"
        assert cls.title == "补全各工序工程量"
        assert cls.__module__ == "pipeline.nodes.quantity_agent", \
            "必须是真模块里的类，不是 builder 里的占位类"

    def test_接线是硬import没有静默降级通道(self):
        """⚠️ 【第 3 批收口 · 父代理亲改】原来的
        `try: import ... except ImportError:` + 同名占位类**已删除**。

        理由（不是"文件已落地所以不需要"，而是**留着它本身就是缺陷**）：
        那条容错等于给"模块缺失"（改名 / 打包漏文件 / 循环 import）留了一条
        **静默降级**通道 —— 计划会照常跑完，只是一条工程量都不补，而没有任何人
        看得见。这与本批正在修的 `beat_config.txt` 静默死路径同族
        （见 `nodes/beat_node.py` 的 `BEAT_CONFIG_UNUSABLE_MSG` 与设计 §14.1.1）。
        宁可 import 时炸，也不要静默出一份没补量的计划。
        """
        src = (BACKEND / "pipeline" / "builder.py").read_text(encoding="utf-8")
        assert "QUANTITY_AGENT_IMPORT" not in src, \
            "容错接线常量必须删干净（父代理收口要求）"
        # ⚠️ 判定必须走**语法树**，不许数子串 —— 上面那段收口说明的注释里
        # 就含 `except ImportError` 字样，子串判定会命中自己的注释而假红。
        # （同一条纪律见 `backend\_probe_tmp\q_batch2_accept.py` 的 strip_code。）
        import ast as _ast
        tree = _ast.parse(src)
        wrapped = []
        for n in _ast.walk(tree):
            if isinstance(n, _ast.Try):
                for sub in _ast.walk(n):
                    if (isinstance(sub, _ast.ImportFrom)
                            and "quantity_agent" in (sub.module or "")):
                        wrapped.append(n.lineno)
        assert not wrapped, \
            "域 5 节点接线被 try 包住了（静默降级通道），行号 %s" % wrapped
        mods = [n for n in _ast.walk(tree)
                if isinstance(n, _ast.ImportFrom)
                and "quantity_agent" in (n.module or "")]
        assert mods, "必须有一条裸 import 引入 QuantityAgentNode"
        assert not hasattr(B, "QUANTITY_AGENT_IMPORT"), \
            "容错期的模块常量不该再存在"
        assert B.QuantityAgentNode.__module__ == "pipeline.nodes.quantity_agent"

    def test_真节点输出能直接喂进第八节(self):
        """契约握手（§4.3 出参 → §9.2 交付物输入）：两边的键名必须对得上。

        `run({})`（无 `kb_scope`）是节点的**空闭集早退**路径：不调模型、不落盘、不 `_stop`，
        返回 `{"quantity_coverage": …, "quantity_warnings": […]}` —— 正好用来钉契约形状。
        """
        assert getattr(B.QuantityAgentNode, "warning_ctx_key", "") == "quantity_warnings"
        out = B.QuantityAgentNode().run({})
        assert isinstance(out, dict), out
        assert isinstance(out.get("quantity_coverage"), dict), out
        assert isinstance(out.get("quantity_warnings"), list), out

        # 空闭集（closed_total == 0、无逐条数据）⇒ 第八节**优雅缺席**（不是空壳、不印 0 行表）
        assert D.quantity_coverage_blocks({"meta": {"quantity_coverage": out["quantity_coverage"]}}) == []

    def test_delivery里不再写死节点数(self):
        src = (BACKEND / "pipeline" / "nodes" / "delivery.py").read_text(encoding="utf-8")
        assert "26 个节点" not in src, "节点数叙述写死就是下一处对不上的地方"


# ══════════════════ ② §14.1.1：beat_config.txt 静默死路径 ══════════════════

PHASE = next(iter(BN.BASE_BEAT_CONFIGS))          # 真实节拍阶段名（取配置表第一个）


class _FakeLLM:
    """假模型：只记调用次数；可以指定"一调就抛"。"""

    def __init__(self, exc=None):
        self.exc = exc
        self.calls = 0

    def chat_json(self, system, user, **kw):
        self.calls += 1
        if self.exc is not None:
            raise self.exc
        return {}


def _raise_missing_prompt(name):
    """`prompts_loader.load("beat_config.txt")` 的真实行为：文件不在 → FileNotFoundError。"""
    raise FileNotFoundError(2, "No such file or directory", str(name))


def _resolve(node):
    """跑一次配置解析，返回 (结果, [(事件, 载荷)])。"""
    seen = []
    node._emit = lambda ev, d: seen.append((ev, d))
    cfg = node._resolve_config({}, PHASE, "1", {})
    return cfg, seen


class _BeatProbe(BaseNode):
    """把 BeatExpandNode 塞进真流水线，验证告警落进 `ctx["node_warnings"]`。"""

    name = "beat_build"
    title = "节拍展开（探针）"

    def __init__(self, beat):
        BaseNode.__init__(self)
        self._beat = beat

    def run(self, ctx):
        self._beat._emit = self.emit
        self._beat._resolve_config(ctx, PHASE, "1", {})
        self.done_summary = "探针跑完"
        return None


class TestBeatConfigDeadPath:
    def test_提示词文件缺失时发一条告警而不是静默pass(self, monkeypatch):
        monkeypatch.setattr(BN, "load", _raise_missing_prompt)
        llm = _FakeLLM()
        cfg, seen = _resolve(BN.BeatExpandNode(llm=llm))

        warns = [d for ev, d in seen if ev == "warning"]
        assert len(warns) == 1, "配置错误必须恰好留一条告警：%r" % (seen,)
        assert warns[0]["node"] == "beat_build"
        assert warns[0]["message"] == BN.BEAT_CONFIG_UNUSABLE_MSG
        assert "beat_config.txt" in warns[0]["detail"], warns[0]
        assert PHASE in warns[0]["detail"], "告警要说明是哪个节拍阶段：%r" % (warns[0],)

        # **行为一字不变**：照旧降级到 BASE 基线（不新建 beat_config.txt、不改工程量）
        assert cfg, cfg
        assert cfg["node_id"] == "1"
        assert llm.calls == 0, "文件都读不到，不许去调模型"

    def test_其它异常仍静默降级不发告警(self, monkeypatch):
        # 提示词**能读到**、是模型调用本身失败 —— 这一档必须保持原样静默降级
        monkeypatch.setattr(BN, "load", lambda name: "系统提示词")
        for exc in (LLMError("LLM 调用失败（共尝试 3 次）：timeout"),
                    RuntimeError("模型返回畸形结构")):
            cfg, seen = _resolve(BN.BeatExpandNode(llm=_FakeLLM(exc=exc)))
            assert [d for ev, d in seen if ev == "warning"] == [], \
                "非文件缺失的失败仍按原样静默降级：%r" % (exc,)
            assert cfg, cfg

    def test_真实的beat_config确实缺失(self):
        """§14.1.1 的前提：`prompts/beat_config.txt` 不在仓库里（不许新建它）。

        真跑一次 `load`，钉住"它是 OSError"这件事 —— 万一以后有人把文件落进来，
        这条会失败，逼着重新评估「突然激活沉睡的 LLM 路径」这件事。
        """
        if (BACKEND / "prompts" / "beat_config.txt").exists():
            pytest.skip("beat_config.txt 已落地：须按 §14.1.1 重新评估这条 LLM 路径")
        with pytest.raises(OSError):
            BN.load("beat_config.txt")

    def test_告警经引擎落进ctx的node_warnings(self, monkeypatch):
        monkeypatch.setattr(BN, "load", _raise_missing_prompt)
        probe = _BeatProbe(BN.BeatExpandNode(llm=_FakeLLM()))
        ctx = {}
        pipe = Pipeline(run_id="t")
        pipe.add_node(probe)
        pipe.run(ctx, emit=lambda ev, d: None)

        items = ctx.get(NODE_WARNINGS_CTX_KEY) or []
        assert len(items) == 1, items
        assert items[0]["node"] == "beat_build"
        assert items[0]["message"] == BN.BEAT_CONFIG_UNUSABLE_MSG
        assert items[0]["count"] == 1

    def test_模块文档已与实现一致(self):
        doc = BN.__doc__ or ""
        assert "LLM(beat_config.txt) → ②" not in doc, "旧失败链原文必须删掉"
        assert "当前不可用" in doc and "代码推算" in doc, doc[:400]


# ══════════════════ ③ 交付物：未入树清单 ══════════════════

def _qc_row(aid, wt, in_tree, **extra):
    r = {"activity_id": aid, "activity_name": "工序" + aid, "work_type_id": wt,
         "work_type_name": wt, "in_tree": in_tree}
    r.update(extra)
    return r


def _nit(r, **extra):
    d = {"activity_id": r["activity_id"], "activity_name": r["activity_name"],
         "work_type_id": r["work_type_id"], "work_type_name": r["work_type_name"],
         "unit": "m2", "quantity": 12.5, "status": "derived", "source": "ratio",
         "reason": "节拍引擎未展开"}
    d.update(extra)
    return d


def _coverage_plan():
    """闭集 70 个 L4：脚手架 31（进树 1）、桩基 39（进树 9）→ 未入树 60 条（>50，触发截断）。"""
    rows, nits = [], []
    for i in range(31):
        it = (i == 0)
        r = _qc_row("JG%02d" % i, "脚手架", it)
        rows.append(r)
        if not it:
            nits.append(_nit(r))
    for i in range(39):
        it = (i < 9)
        r = _qc_row("ZJ%02d" % i, "桩基", it)
        rows.append(r)
        if not it:
            nits.append(_nit(r))
    qc = {"source": "ratio", "structure_type_id": "frame_shear",
          "l4_rows": rows, "l4": {r["activity_id"]: r for r in rows},
          "not_in_tree": nits,
          "summary": {"closed_total": 70, "in_tree": 10, "unit_unresolved": 2,
                      "by_source": {"ratio": 40, "llm": 25, "user": 5}}}
    return {"meta": {"quantity_coverage": qc}}


def _blocks(plan):
    return D.confidence_section_blocks(plan, D._compute_view(plan))


def _my_blocks(blocks):
    """只取域 5 那一节（从它的 h3 到下一个 h3 之前）—— 章节里 kv/grid 不止一处。"""
    out, on = [], False
    for b in blocks:
        if b[0] == "h3":
            on = (b[1] == D.QUANTITY_COVERAGE_TITLE)
            if on:
                out.append(b)
            continue
        if on:
            out.append(b)
    return out


def _cells(blocks):
    """所有 kv / grid 格子（含表头）摊平成字符串表 —— 用来查 "None" 漏网。"""
    out = []
    for kind, payload in blocks:
        if kind == "kv":
            out += [str(c) for row in payload for c in row]
        elif kind == "grid":
            headers, rows = payload[0], payload[1]
            out += [str(h) for h in headers]
            out += [str(c) for row in rows for c in row]
        else:
            out.append(str(payload))
    return out


def _find(blocks, pred):
    return [i for i, b in enumerate(blocks) if pred(b)]


class TestQuantityCoverageDeliverable:
    def test_三样东西都在截断之前(self):
        plan = _coverage_plan()
        mine = _my_blocks(_blocks(plan))

        i_title = _find(mine, lambda b: b[0] == "h3")[0]

        # ① 总数：一句话给全
        kv = [b[1] for b in mine if b[0] == "kv"]
        assert kv, mine
        total = kv[0][0][1]
        assert total == "闭集 70 个 L4，进树 10 个，未入树 60 个", total

        # ② 按 L3 工种分组
        grows = [b[1] for b in mine if b[0] == "grid" and b[1][0][0] == "工种（L3）"]
        assert len(grows) == 1, [b[1][0] for b in mine if b[0] == "grid"]
        assert ["脚手架", "31", "1", "30"] in grows[0][1], grows[0][1]
        assert ["桩基", "39", "9", "30"] in grows[0][1], grows[0][1]

        # ③ 一句人话解释
        assert _find(mine, lambda b: b[0] == "para" and b[1] == D.QUANTITY_ABSENCE_WHY), \
            "缺「为什么没进树」的人话解释"

        # 顺序：标题 → 总数 → 分布 → 解释 → 明细（截断必须最后）
        i_group = _find(mine, lambda b: b[0] == "grid" and b[1][0][0] == "工种（L3）")[0]
        i_why = _find(mine, lambda b: b[1] == D.QUANTITY_ABSENCE_WHY)[0]
        i_detail = _find(mine, lambda b: b[0] == "grid" and b[1][0][0] == "工序 ID")[0]
        assert i_title < i_group < i_why < i_detail, (i_title, i_group, i_why, i_detail)

    def test_明细前50条其余指向meta(self):
        mine = _my_blocks(_blocks(_coverage_plan()))
        detail = [b[1] for b in mine
                  if b[0] == "grid" and b[1][0][0] == "工序 ID"][0]
        assert len(detail[1]) == D.QUANTITY_NOT_IN_TREE_LIMIT == 50, len(detail[1])

        tails = [b[1] for b in mine if b[0] == "para" and "not_in_tree" in str(b[1])]
        assert tails, [b[1] for b in mine if b[0] == "para"]
        assert "其余 10 条" in tails[0] and "meta.quantity_coverage.not_in_tree" in tails[0], \
            tails[0]

    def test_来源分档与单位未落实按数据出(self):
        kv = [b[1] for b in _my_blocks(_blocks(_coverage_plan())) if b[0] == "kv"][0]
        txt = "\n".join("%s=%s" % (k, v) for k, v in kv)
        assert "模型补量=25 个 L4（35.7%）" in txt, txt
        assert "单位未落实=2 个 L4" in txt, txt

    def test_字典单位与定额单位不一致的条数要露出来(self):
        plan = _coverage_plan()
        rows = plan["meta"]["quantity_coverage"]["l4_rows"]
        rows[3]["unit_evidence"] = "dict+norm_differs"
        rows[4]["unit_evidence"] = "dict+norm_differs"
        kv = [b[1] for b in _my_blocks(_blocks(plan)) if b[0] == "kv"][0]
        txt = "\n".join("%s=%s" % (k, v) for k, v in kv)
        assert "2 个 L4" in txt and "字典单位与定额单位不一致" in txt, txt

    def test_字典单位条数优先取summary里的那个键(self):
        """节点的 `summary.dict_norm_differs` 是它的官方口径；逐条表现数是兜底。"""
        plan = _coverage_plan()
        plan["meta"]["quantity_coverage"]["summary"]["dict_norm_differs"] = 3
        kv = [b[1] for b in _my_blocks(_blocks(plan)) if b[0] == "kv"][0]
        txt = "\n".join("%s=%s" % (k, v) for k, v in kv)
        assert "字典单位与定额单位不一致=3 个 L4" in txt, txt

    def test_键缺失时这一节优雅缺席(self):
        assert D.quantity_coverage_blocks({}) == []
        for qc in (None, {}, "bad", [], {"summary": {"closed_total": 0}},
                   {"not_in_tree": []}, {"l4_rows": []},
                   {"not_in_tree": "nope"}, {"l4": {}}):
            plan = {"meta": {"quantity_coverage": qc}}
            assert D.quantity_coverage_blocks(plan) == [], qc
            assert D.QUANTITY_COVERAGE_TITLE not in [
                b[1] for b in _blocks(plan) if b[0] == "h3"], qc

    def test_老计划整章不受影响且不出现本节(self):
        plan = {"meta": {"norm_coverage": {"total": 1, "bound": 1}}}
        blocks = _blocks(plan)
        assert blocks, blocks
        assert D.QUANTITY_COVERAGE_TITLE not in [b[1] for b in blocks if b[0] == "h3"]

    def test_缺字段的行一个None都不印(self):
        plan = {"meta": {"quantity_coverage": {
            "l4_rows": [{"activity_id": "X1", "in_tree": False}],
            "not_in_tree": [{"activity_id": "X1"},
                            {"activity_id": "X2", "activity_name": None, "unit": None,
                             "quantity": None, "source": None, "reason": None,
                             "work_type_name": "", "work_type_id": None}],
            "summary": {"closed_total": 2, "in_tree": 0}}}}
        blocks = D.quantity_coverage_blocks(plan)
        cells = _cells(blocks)
        assert cells, blocks
        assert all(isinstance(c, str) for c in cells), cells
        for bad in ("None", "nan", "NaN"):
            assert not any(bad == c for c in cells), (bad, cells)
        assert "None" not in "".join(cells), cells
        assert "闭集 2 个 L4，进树 0 个，未入树 2 个" in "".join(cells)

    def test_门进CONFIDENCE_META_KEYS且单独也能开章(self):
        assert "quantity_coverage" in D.CONFIDENCE_META_KEYS
        plan = _coverage_plan()
        assert D.has_confidence_meta(plan) is True
        assert D.has_confidence_section(plan, D._compute_view(plan)) is True
        assert D.has_confidence_meta({"meta": {}}) is False
        assert D.has_confidence_meta({"meta": {"quantity_coverage": {}}}) is False

    def test_看板卡片与Word同一真源(self):
        plan = _coverage_plan()
        html = D._confidence_section_html(plan, D._compute_view(plan))
        assert D.QUANTITY_COVERAGE_TITLE in html
        assert "闭集 70 个 L4，进树 10 个，未入树 60 个" in html
        assert "脚手架" in html
        assert "None" not in html

        other = D._confidence_section_html(
            {"meta": {"norm_coverage": {"total": 1, "bound": 1}}},
            D._compute_view({"meta": {"norm_coverage": {"total": 1, "bound": 1}}}))
        assert D.QUANTITY_COVERAGE_TITLE not in other

    def test_人工覆盖节编号让到9(self):
        assert D.QUANTITY_COVERAGE_TITLE.startswith("8.")
        assert D.NORM_OVERRIDE_TITLE.startswith("9.")

    def test_Word渲染路径吃得下这一节的块(self):
        """块是 Word 与看板的**同一真源**：这里用真 python-docx 走一遍 add_confidence_section。

        `add_kv` / `add_grid` 只兜得住**真 None**，所以这条顺带钉死"格子全是字符串、
        一个 None 都不漏"。
        """
        from docx import Document

        plan = _coverage_plan()
        doc = Document()

        def add_h(text, level=1):
            doc.add_heading(text, level=level)

        def add_kv(rows):
            t = doc.add_table(rows=0, cols=2)
            for k, v in rows:
                c = t.add_row().cells
                c[0].text = str(k)
                c[1].text = "" if v is None else str(v)

        def add_grid(header, rows):
            t = doc.add_table(rows=1, cols=len(header))
            for j, h in enumerate(header):
                t.rows[0].cells[j].text = str(h)
            for row in rows:
                c = t.add_row().cells
                for j, v in enumerate(row):
                    c[j].text = "" if v is None else str(v)

        assert D.add_confidence_section(doc, plan, D._compute_view(plan),
                                        add_h, add_kv, add_grid) is True
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for want in (D.QUANTITY_COVERAGE_TITLE, "闭集 70 个 L4，进树 10 个，未入树 60 个",
                     "脚手架", "工序JG01", "meta.quantity_coverage.not_in_tree"):
            assert want in text, want
        assert "None" not in text, text[-500:]
