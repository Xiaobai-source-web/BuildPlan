# -*- coding: utf-8 -*-
"""交付物对「AI 经验估算定额」的逐条标注 —— 政策变更（用户 2026-09-20）护栏。

政策变更前后（判据层在 `norm_defaults.py` / `scheduler.py`，消费层在 `delivery.py`）：

  旧政策：AI 凭经验编的定额（KB `sources.AI_ESTIMATE_V1`）"只作参考，不用来算班组"，
          交付物里写着「⚠ 无可用定额：AI 估算定额（只作参考，不用来算班组）」。
  新政策：AI 经验估算定额**照用**，与真人定额同等参与工期与班组计算，状态名
          `released_ai`，标注文案 `AI 经验估算定额（无规范依据，待审）`。
          仍然拦：人工否决 / 单位不可换算 / 定额口径不符 —— 那批任务照旧写
          「⚠ 无可用定额：…（工期沿用 WBS 估算，未计算班组）」，一个字都不许删。

本次政策变更**唯一不能省**的部分就是本文件钉的两件事：
  ① 交付物逐条标出"这条的工期/班组依据来自 AI 经验估算"（看板与 Word 同一真源）；
  ② 置信度章节给出**条数**，且条数必须从数据来 —— 数不出来整行不出，绝不用 0 充数。

运行：python -m pytest tests/test_delivery_ai_norm_label.py -q
"""

import re
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from pipeline.nodes import delivery as D          # noqa: E402

FORBIDDEN_OLD_PHRASE = "只作参考，不用来算班组"


# ══════════════════════ 公共样例 ══════════════════════

def _rd(tid, name, **kw):
    """一条 `resource_demand.tasks[*]`（只带本文件关心的字段）。"""
    base = {"task_id": tid, "task_name": name, "quantity": 100.0,
            "planned_duration_days": 5}
    base.update(kw)
    return base


def _ai_applied(tid, name):
    """AI 经验估算定额**真的照用**（新政策下的典型形态：`_norm_applied` 是 AI 来源）。"""
    return _rd(tid, name,
               resources={"瓦工": {"per_day": 9, "total_days": 45.0}},
               _resource_source={"瓦工": {"origin": "ai_estimate",
                                          "ref": "AI_ESTIMATE_V1"}},
               _norm_applied={"mode": "labor", "norm_value": 0.1,
                              "source_code": "AI_ESTIMATE_V1",
                              "match_type": "ai", "origin": "ai"})


def _ai_legacy_flag(tid, name):
    """旧计划落盘的 AI 拦截原文（政策变更前的产物）→ 交付物必须换文案。"""
    return _rd(tid, name,
               _norm_flagged="AI 估算定额（只作参考，不用来算班组）",
               _warning="定额不可作证据，未计算班组：AI 估算定额（只作参考，不用来算班组）")


def _real_norm(tid, name, ai_site_equipment=False):
    """真人定额（可另带一个 AI 来源的场地级设备台数）。"""
    t = _rd(tid, name,
            resources={"钢筋工": {"per_day": 20, "total_days": 100.0}},
            _norm_applied={"mode": "labor", "norm_value": 0.05,
                           "source_code": "GD_2018_A1_3",
                           "match_type": "exact", "origin": "kb"},
            _resource_source={"钢筋工": {"origin": "kb", "ref": "GD_2018_A1_3"}})
    if ai_site_equipment:
        t["resources"]["塔吊"] = {"per_day": 1, "total_days": 5.0}
        t["_resource_source"]["塔吊"] = {"origin": "ai_estimate",
                                         "ref": "AI_ESTIMATE_V1"}
    return t


def _unit_blocked(tid, name):
    """真降级：单位不可换算（**不是** AI 来源）。"""
    return _rd(tid, name,
               _norm_flagged="单位不可用：单位不一致且不可换算：任务「根」（count:根） vs "
                             "定额分母「m³」（volume）")


def _plan(rd_tasks, meta=None):
    """一份最小计划：WBS 叶子与排程行跟着 `rd_tasks` 一一对应。"""
    def _leaf(t):
        leaf = {"id": t["task_id"], "name": t["task_name"], "duration_days": 5,
                "quantity": t.get("quantity") or 1, "unit": "m³",
                "work_type": "混凝土工程"}
        return leaf

    def _sched(t):
        return {"task_id": t["task_id"], "task_name": t["task_name"],
                "start_date": "2026-09-01", "finish_date": "2026-09-05",
                "duration_days": 5,
                "assigned_resources": {k: (v.get("per_day") if isinstance(v, dict) else v)
                                       for k, v in (t.get("resources") or {}).items()}}

    plan = {
        "plan_id": "plan_ai_norm_label_test",
        "overview": {"project_name": "AI 定额标注测试", "total_duration_days": 10,
                     "planned_start_date": "2026-09-01",
                     "planned_end_date": "2026-09-11", "critical_path_length": 1},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "1.1", "name": "结构", "sub_packages": [_leaf(t) for t in rd_tasks]}]}]},
        "dependencies": [],
        "cpm_result": {"total_duration_days": 10, "critical_path": ["1.1.1"],
                       "schedule": [{"task_id": t["task_id"], "es": 0, "ef": 4}
                                    for t in rd_tasks]},
        "resource_demand": {"tasks": rd_tasks},
        "key_milestones": [{"name": "开工", "date": "2026-09-01",
                            "task_id": rd_tasks[0]["task_id"] if rd_tasks else "1.1.1",
                            "description": "开工"}],
        "critical_path_tasks": [_sched(rd_tasks[0])] if rd_tasks else [],
        "all_tasks_schedule": [_sched(t) for t in rd_tasks],
        "resource_plan": {"total_manpower_days": 145.0, "peak_manpower": 9,
                          "peak_manpower_source": "resource_curve",
                          "curve_peak_manpower": 9, "equipment_peak": {},
                          "material_summary": []},
        "report": "# 报告",
    }
    if meta is not None:
        plan["meta"] = meta
    return plan


def _ev(rd):
    return D._evidence_text(rd)


def _board_text(plan, tmp_path, monkeypatch):
    """渲染看板（隔离到 tmp，避免往真实 `输出结果/` 写东西），返回 HTML 文本。"""
    from pipeline import config
    monkeypatch.setattr(config, "DELIVERABLES_DIR", tmp_path)
    monkeypatch.setattr(D.config, "DELIVERABLES_DIR", tmp_path)
    return Path(D.build_plan_html(plan)).read_text(encoding="utf-8")


def _docx_text(plan, tmp_path, monkeypatch):
    """渲染 Word，返回「段落 + 表格单元格」的全部文本。"""
    from pipeline import config
    monkeypatch.setattr(config, "DELIVERABLES_DIR", tmp_path)
    monkeypatch.setattr(D.config, "DELIVERABLES_DIR", tmp_path)
    from docx import Document
    doc = Document(D.build_plan_docx(plan))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parts.extend(c.text for c in r.cells)
    return "\n".join(parts)


# ═══════════════ ① 逐条标注：依据来源是 AI 定额 ═══════════════

class TestAiNormLabelPerTask:
    def test_AI定额照用的任务标注新文案并给班组(self):
        """新政策形态：`_norm_applied` 是 AI 来源 → 标 AI 文案 + 正常显示班组/工日。"""
        text = _ev(_ai_applied("1.1.1", "柱浇筑"))
        assert D.AI_NORM_LABEL in text, text
        assert "班组 瓦工 9 人（45 工日）" in text, text
        assert "未计算班组" not in text, "真的照用算了班组，不许再说没算"
        assert FORBIDDEN_OLD_PHRASE not in text

    def test_旧计划AI拦截原文换文案且不谎称算了班组(self):
        """旧计划形态：`_norm_flagged` 点名 AI 定额但没有 AI 的 `_norm_applied`。

        交付物必须：① 换成政策文案；② 如实说这一回没据此算班组（不替它编班组）；
        ③ 一个字都不复述旧政策那句（fixture 里就是原文）。
        """
        text = _ev(_ai_legacy_flag("1.1.2", "梁浇筑"))
        assert D.AI_NORM_LABEL in text, text
        assert "本计划未按该定额计算班组" in text and "工期沿用 WBS 估算" in text, text
        assert FORBIDDEN_OLD_PHRASE not in text, text

    def test_AI定额行同时带单位假定换算时两者都显示(self):
        """`_unit_assumed`（写明来源的换算参数）与 AI 定额标注是两件事，都要显示。

        P3/D5：换算参数的**来源**由 `ctx_source` 分流 —— 这里显式给 `ai_estimate`，
        文案才是「AI估算换算参数」；来源键缺失时会如实写「来源未记录」，不默认归到 AI。
        """
        rd = _ai_applied("1.1.3", "ALC 墙板安装")
        rd["_unit_assumed"] = "按 AI 假定墙厚 200mm 换算：1420 m² × 0.2 m = 284 m³"
        rd["ctx_source"] = "ai_estimate"
        text = _ev(rd)
        assert "单位换算按AI估算换算参数" in text and "墙厚 200mm" in text, text
        assert D.AI_NORM_LABEL in text, text
        assert "班组 瓦工 9 人" in text, text

    def test_换算参数来源是定额条件时不许说成AI(self):
        """P3/D5 的核心：来源是 `norm_condition`（定额行适用条件档位）→ **非 AI 估算**。"""
        rd = _ai_applied("1.1.3", "ALC 墙板安装")
        rd["_unit_assumed"] = "按定额条件档位墙厚 200mm 换算：1420 m² × 0.2 m = 284 m³"
        rd["ctx_source"] = "norm_condition"
        text = _ev(rd)
        assert "定额条件档位换算参数" in text and "非 AI 估算" in text, text
        assert "AI估算换算参数" not in text, text

    def test_换算参数来源缺失时如实写未记录(self):
        """来源键缺失 → 「换算参数来源未记录」，**绝不默认归到 AI**。"""
        rd = _ai_applied("1.1.3", "ALC 墙板安装")
        rd["_unit_assumed"] = "按墙厚 200mm 换算：1420 m² × 0.2 m = 284 m³"
        text = _ev(rd)
        assert "换算参数来源未记录" in text, text
        assert "AI估算换算参数" not in text, text

    def test_真降级仍然写无可用定额与未计算班组(self):
        """单位不可换算 / 人工否决 / 无定额绑定 —— 措辞一个字没删（政策另一半）。"""
        text = _ev(_unit_blocked("1.1.4", "截（凿）桩头"))
        assert "⚠ 无可用定额" in text and "未计算班组" in text, text
        assert "工期沿用 WBS 估算" in text, text
        assert D.AI_NORM_LABEL not in text, "真降级不是 AI 来源，不许贴 AI 标注"

    def test_人工否决与无绑定的原因原文照旧(self):
        for reason in ("人工否决：评审认定不适用", "无定额绑定"):
            rd = _rd("9.9.1", "某任务", _norm_flagged=reason)
            text = _ev(rd)
            assert reason in text, text
            assert "（工期沿用 WBS 估算，未计算班组）" in text, text

    def test_真人定额加AI来源资源时弱标注指名道姓(self):
        """定额是真人定额、但塔吊台数走 AI 默认口径 → 弱标注必须点出是哪个资源。"""
        text = _ev(_real_norm("1.1.5", "钢筋绑扎", ai_site_equipment=True))
        assert "定额 GD_2018_A1_3" in text, text
        assert "含 AI 经验估算来源：塔吊（AI_ESTIMATE_V1，无规范依据，待审）" in text, text
        assert FORBIDDEN_OLD_PHRASE not in text
        # 没有 AI 来源的真人定额行不许被贴上任何 AI 标注
        plain = _ev(_real_norm("1.1.6", "钢筋绑扎"))
        assert "AI" not in plain, plain

    def test_判据是数据不是任务名(self):
        """把任务名换成完全无关的字，判据仍然成立（不按名字硬编码）。"""
        rd = _ai_legacy_flag("1.1.7", "随便改个名字")
        assert D.AI_NORM_LABEL in _ev(rd)
        rd2 = _rd("1.1.8", "名字里带 AI 估算定额")
        rd2["_norm_flagged"] = "单位不可用：单位不一致且不可换算"
        assert D.AI_NORM_LABEL not in _ev(rd2)

    def test_组行也要分AI与真降级两类(self):
        """合并展示行（组）：AI 经验估算与"真的没算出来"必须分开计数。"""
        rd_map = {"1.1.1": _ai_applied("1.1.1", "柱浇筑"),
                  "1.1.2": _ai_legacy_flag("1.1.2", "梁浇筑"),
                  "1.1.3": _real_norm("1.1.3", "钢筋绑扎", ai_site_equipment=True),
                  "1.1.4": _unit_blocked("1.1.4", "截（凿）桩头")}
        text = D._group_evidence(list(rd_map), rd_map)
        assert "2 项依据为 %s" % D.AI_NORM_LABEL in text, text
        assert "1 项含 AI 经验估算来源" in text, text
        assert "⚠ 其中 1 项无可用定额（工期沿用 WBS 估算，未计算班组）" in text, text
        assert FORBIDDEN_OLD_PHRASE not in text

    def test_拿不到字段时退回来源未记录不猜(self):
        assert _ev({}) == "来源未记录（该任务没有资源记录）"
        assert _ev(_rd("1.1.9", "空记录")) == "来源未记录（无资源记录）"


# ═══════════════ ② 交付物全文不许再出现旧政策措辞 ═══════════════

class TestNoOldPolicyWording:
    def test_看板与Word都不出现旧口径(self, tmp_path, monkeypatch):
        plan = _plan([_ai_applied("1.1.1", "柱浇筑"),
                      _ai_legacy_flag("1.1.2", "梁浇筑（旧计划落盘）"),
                      _real_norm("1.1.3", "钢筋绑扎", ai_site_equipment=True),
                      _unit_blocked("1.1.4", "截（凿）桩头")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 4, "bound": 3, "bound_pct": 75.0,
                                             "unbound": 1, "unbound_pct": 25.0}})
        html = _board_text(plan, tmp_path, monkeypatch)
        assert FORBIDDEN_OLD_PHRASE not in html, "看板复述了已废除的旧口径"
        assert D.AI_NORM_LABEL in html
        assert "含 AI 经验估算来源：塔吊" in html
        text = _docx_text(plan, tmp_path, monkeypatch)
        assert FORBIDDEN_OLD_PHRASE not in text, "Word 复述了已废除的旧口径"
        assert D.AI_NORM_LABEL in text
        # 真降级那句话仍然在（不许因为改口径把它一起删掉）
        assert "工期沿用 WBS 估算，未计算班组" in text

    def test_降级清单引导句在含AI行时不说仅作参考(self):
        assert D._degraded_lead([]) == D.NORM_DEGRADED_LEAD
        lead = D._degraded_lead([{"ai": True, "reason": "x"}])
        assert lead != D.NORM_DEGRADED_LEAD and D.AI_NORM_LABEL in lead, lead

    def test_降级原因把AI旧文案换掉(self):
        got = D._degraded_reason("AI 估算定额（只作参考，不用来算班组）")
        assert got.startswith(D.AI_NORM_LABEL), got
        assert FORBIDDEN_OLD_PHRASE not in got
        assert D._degraded_reason("无定额绑定") == "无定额绑定"

    def test_旧口径原始记录逐条挂澄清(self):
        """`meta.norm_coverage.by_reason` 是计划原始记录：一字不改 + 挂澄清。"""
        for raw in ("L4默认定额行仅有 AI 估算，不作工期证据",
                    "AI估算定额：KB 无定额行，只作参考"):
            got = D._by_reason_display(raw)
            assert got.startswith(raw), got
            assert "政策变更 2026-09-20：AI 经验估算定额已照用" in got, got
        # 与 AI 无关的原因、以及"AI 定额被单位拦下"的真实拦截都不该被贴上"已照用"
        assert D._by_reason_display("单位不一致") == "单位不一致"
        assert "已照用" not in D._by_reason_display("AI 估算定额单位不可换算")

    def test_单位拦截优先于AI来源不放行(self):
        """带明确单位/人工否决拦截词的原文 → 仍然写「无可用定额」，不许说已放行。"""
        rd = _rd("1.1.10", "AI 定额但单位不可换算",
                 _norm_flagged="单位不可用：单位不一致且不可换算：任务「根」 vs "
                               "定额分母「m³」（AI 估算定额）")
        text = _ev(rd)
        assert "⚠ 无可用定额" in text and "工期沿用 WBS 估算，未计算班组" in text, text
        assert D.AI_NORM_LABEL not in text, text


# ═══════════════ ③ 置信度章节的 AI 条数（唯一不能省） ═══════════════

def _blocks(plan):
    return D.confidence_section_blocks(plan, D._compute_view(plan))


def _kv(plan):
    """把 §1 的 kv 行汇总成 {标签: 值}。"""
    out = {}
    for kind, payload in _blocks(plan):
        if kind == "kv":
            for label, value in payload:
                out[str(label)] = value
    return out


def _paras(plan):
    return "\n".join(str(p) for k, p in _blocks(plan) if k == "para")


class TestAiNormCountInConfidence:
    def test_条数从任务数据算_给不同条数断言跟着变(self):
        for n in (1, 3):
            plan = _plan([_ai_applied("1.1.%d" % i, "AI 任务 %d" % i) for i in range(1, n + 1)],
                         meta={"audit_status": "未审计",
                               "norm_coverage": {"total": n, "bound": n, "bound_pct": 100.0,
                                                 "unbound": 0, "unbound_pct": 0.0}})
            kv = _kv(plan)
            assert D.AI_NORM_ROW_LABEL_LEGACY in kv, kv
            assert kv[D.AI_NORM_ROW_LABEL_LEGACY] == "%d 条（本次运行已据此算出班组与工日）" % n
            # 旧计划形态（没有 AI 的 `_norm_applied`）时条数照样数得出来，但状态如实写
            legacy = _plan([_ai_legacy_flag("1.1.%d" % i, "旧 AI 任务 %d" % i)
                            for i in range(1, n + 1)],
                           meta={"audit_status": "未审计"})
            kv2 = _kv(legacy)
            assert kv2[D.AI_NORM_ROW_LABEL_LEGACY] == "%d 条（本次运行未据此计算班组）" % n

    def test_上游released_ai优先且用政策标签(self):
        plan = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 400, "bound": 350, "bound_pct": 87.5,
                                             "unbound": 50, "unbound_pct": 12.5,
                                             "released_ai": 9, "released_ai_pct": 2.3}})
        kv = _kv(plan)
        assert kv.get(D.AI_NORM_ROW_LABEL) == "9（2.3%）", kv
        assert D.AI_NORM_ROW_LABEL_LEGACY not in kv

    def test_数不出来时整行不出绝不用0充数(self):
        # 没有任何 AI 任务
        plain = _plan([_real_norm("1.1.1", "钢筋绑扎")],
                      meta={"audit_status": "未审计",
                            "norm_coverage": {"total": 1, "bound": 1, "bound_pct": 100.0,
                                              "unbound": 0, "unbound_pct": 0.0}})
        kv = _kv(plain)
        assert D.AI_NORM_ROW_LABEL not in kv and D.AI_NORM_ROW_LABEL_LEGACY not in kv
        assert "0 条" not in "".join(str(v) for v in kv.values()), kv
        # 任务级数据整块拿不到（空 tasks）→ 也数不出来，不许猜 0
        empty = _plan([], meta={"audit_status": "未审计",
                                "norm_coverage": {"total": 5, "bound": 5, "bound_pct": 100.0,
                                                  "unbound": 0, "unbound_pct": 0.0}})
        assert D.AI_NORM_ROW_LABEL_LEGACY not in _kv(empty)

    def test_说明段写明AI定额已参与且导入规范后清退(self):
        plan = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 1, "bound": 1, "bound_pct": 100.0,
                                             "unbound": 0, "unbound_pct": 0.0}})
        para = _paras(plan)
        assert "AI 经验估算定额（无规范依据，待审）" in para, para
        assert "与真人定额同等参与工期与班组计算" in para, para
        assert "导入真实规范后应整体清退" in para, para
        assert "本次运行中这些任务均已据此算出班组与工日" in para, para
        # 旧句（AI 一律不参与计算）从今天起是假话，必须消失
        assert "一律不参与计算" not in para, para

    def test_旧计划形态说明段也如实且不谎称已参与(self):
        plan = _plan([_ai_legacy_flag("1.1.1", "旧 AI 任务 1")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 2, "bound": 0, "bound_pct": 0.0,
                                             "unbound": 2, "unbound_pct": 100.0}})
        para = _paras(plan)
        assert "本次运行的这些任务未据此计算班组" in para, para
        assert "均已据此算出班组" not in para, para

    def test_没有norm_coverage时AI条数照样如实给出(self):
        plan = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                     meta={"audit_status": "未审计"})
        assert D.confidence_section_blocks(plan, D._compute_view(plan)), "有 AI 数据就该有内容"
        kv = _kv(plan)
        assert kv.get(D.AI_NORM_ROW_LABEL_LEGACY) == "1 条（本次运行已据此算出班组与工日）"

    def test_estimated出现在by_confidence时如实显示(self):
        plan = _plan([_real_norm("1.1.1", "钢筋绑扎")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 400, "bound": 350, "bound_pct": 87.5,
                                             "unbound": 50, "unbound_pct": 12.5,
                                             "released_unapproved": 85,
                                             "released_unapproved_pct": 21.3,
                                             "by_confidence": {"parsed": 60, "verified": 25,
                                                               "estimated": 12}}})
        kv = _kv(plan)
        assert "其中：未经人工审定的真人定额（已放行）" in kv, kv
        assert "estimated 12 条" in kv["其中：未经人工审定的真人定额（已放行）"], kv
        assert "85（21.3%）" in kv["其中：未经人工审定的真人定额（已放行）"]

    def test_by_confidence的estimated在AI行里也如实显示(self):
        """`by_confidence` 有 `estimated` 而在放行行里看不到时，AI 那一行要带出来。"""
        plan = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 400, "bound": 350, "bound_pct": 87.5,
                                             "unbound": 50, "unbound_pct": 12.5,
                                             "released_ai": 12, "released_ai_pct": 3.0,
                                             "by_confidence": {"parsed": 60, "estimated": 12}}})
        kv = _kv(plan)
        assert "estimated（AI 经验估算）12 条" in kv[D.AI_NORM_ROW_LABEL], kv
        # 放行行已经列出了整张 by_confidence 时不重复
        plan2 = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                      meta={"audit_status": "未审计",
                            "norm_coverage": {"total": 400, "bound": 350, "bound_pct": 87.5,
                                              "unbound": 50, "unbound_pct": 12.5,
                                              "released_unapproved": 85,
                                              "released_unapproved_pct": 21.3,
                                              "released_ai": 12, "released_ai_pct": 3.0,
                                              "by_confidence": {"parsed": 60, "estimated": 12}}})
        kv2 = _kv(plan2)
        assert "estimated（AI 经验估算）" not in kv2[D.AI_NORM_ROW_LABEL], kv2
        assert "estimated 12 条" in kv2["其中：未经人工审定的真人定额（已放行）"]

    def test_未绑定原因里的AI原始记录挂上澄清(self):
        plan = _plan([_real_norm("1.1.1", "钢筋绑扎")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 10, "bound": 6, "bound_pct": 60.0,
                                             "unbound": 4, "unbound_pct": 40.0,
                                             "by_reason": {"AI估算定额：KB 无定额行，只作参考": 4},
                                             "by_reason_pct": {"AI估算定额：KB 无定额行，只作参考": 40.0}}})
        cells = [c for k, p in _blocks(plan) if k == "grid" for row in p[1] for c in row]
        hit = [c for c in cells if "AI估算定额" in str(c)]
        assert hit, cells
        assert "政策变更 2026-09-20：AI 经验估算定额已照用" in str(hit[0]), hit
        # 原始记录本身一字不改（可追溯），只追加澄清
        assert str(hit[0]).startswith("AI估算定额：KB 无定额行，只作参考"), hit

    def test_看板卡片里AI行在可见处_Word同源(self, tmp_path, monkeypatch):
        plan = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 1, "bound": 1, "bound_pct": 100.0,
                                             "unbound": 0, "unbound_pct": 0.0}})
        card = D._confidence_section_html(plan, D._compute_view(plan))
        visible = re.sub(r"<details\b.*?</details>", "", card, flags=re.S)
        assert D.AI_NORM_ROW_LABEL_LEGACY in visible, visible[:800]
        assert "1 条（本次运行已据此算出班组与工日）" in visible
        html = _board_text(plan, tmp_path, monkeypatch)
        assert card in html
        text = _docx_text(plan, tmp_path, monkeypatch)
        assert D.AI_NORM_ROW_LABEL_LEGACY in text, "Word 与看板必须同一真源"
        assert "数据来源与置信度" in text

    def test_facts把新政策口径给模型_不许它复述旧口径(self):
        """LLM 编排看板时也拿得到政策口径（否则它会在正文里复述已废除的旧话）。"""
        plan = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                     meta={"audit_status": "未审计",
                           "norm_coverage": {"total": 1, "bound": 1, "bound_pct": 100.0,
                                             "unbound": 0, "unbound_pct": 0.0}})
        pol = D._facts_bundle(plan, D._compute_view(plan))["ai_norm_policy"]
        assert pol["label"] == D.AI_NORM_LABEL and pol["state"] == "released_ai"
        assert pol["count"] == 1
        assert "不许" in pol["how_to_write"]
        # 数不出来 → count 为 None（模型据此一个字都不写，不许写 0 充数）
        plain = _plan([_real_norm("1.1.1", "钢筋绑扎")], meta={"audit_status": "未审计"})
        assert D._facts_bundle(plain, D._compute_view(plain))["ai_norm_policy"]["count"] is None


# ═══════════ ④ 看板与 Word 的「进出门」同一判据（政策变更 2026-09-20 追加） ═══════════
# 同一份计划的两个交付物口径不一致是用户实测投诉过的病。本章的进出门在两个渲染器里
# 必须是**同一个函数**（`delivery.has_confidence_meta`）：
#   · Word：`add_confidence_section` 第一道判据；
#   · 看板：`_confidence_section_html` 第一道判据。
# 政策变更点：门要把「AI 经验估算定额条数」也算作"有数据"—— 否则"有 AI 定额任务、
# 但计划没带那六项 meta"的计划两边都整章消失，条数无处可写（本次变更唯一不能省的部分）。

class TestConfidenceGateParity:
    def test_有AI任务但无六项meta时两边都出(self, tmp_path, monkeypatch):
        plan = _plan([_ai_applied("1.1.1", "AI 任务 1")],
                     meta={"audit_status": "未审计"})       # 六项 meta 一个都没有
        view = D._compute_view(plan)
        assert D.has_confidence_meta(plan) is True, "AI 条数也算'有置信度数据'"
        assert D.has_confidence_section(plan, view) is True
        html = _board_text(plan, tmp_path, monkeypatch)
        assert "数据来源与置信度" in html, "看板必须出这一章（条数要有处可写）"
        assert D.AI_NORM_ROW_LABEL_LEGACY in html
        text = _docx_text(plan, tmp_path, monkeypatch)
        assert "数据来源与置信度" in text, "Word 必须出这一章（与看板同进）"
        assert D.AI_NORM_ROW_LABEL_LEGACY in text

    def test_无AI任务且无六项meta时两边都不出(self, tmp_path, monkeypatch):
        plan = _plan([_real_norm("1.1.1", "钢筋绑扎")],
                     meta={"audit_status": "未审计"})
        # §4 峰值口径几乎每份计划都算得出来 —— 只有峰值一行也不许开出这一章空壳
        plan["resource_plan"]["declared_peak_manpower"] = 120
        plan["resource_plan"]["declared_peak_manpower_source"] = "model"
        view = D._compute_view(plan)
        assert D.has_confidence_meta(plan) is False
        assert D.confidence_section_blocks(plan, view), "块算得出来（§4 峰值）——所以门不能只判块"
        html = _board_text(plan, tmp_path, monkeypatch)
        assert "数据来源与置信度" not in html, "看板与 Word 必须同出同不出"
        text = _docx_text(plan, tmp_path, monkeypatch)
        assert "数据来源与置信度" not in text, "看板与 Word 必须同出同不出"
        assert D.AI_NORM_ROW_LABEL not in html and D.AI_NORM_ROW_LABEL_LEGACY not in html
