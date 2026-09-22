"""第 39 轮「让交付物说实话」——交付侧可见性护栏。

用户原话：「让交付物说实话：看板和 Word 上写明『人数被工作面容量压到 X 人，
来源=AI 估算/低置信度』；把 120 人改标成『模型估算，非用户输入』。」

实测缺陷（本轮修复前）：计划 JSON 里有 508 处「工作面容量」、73 条封顶记录，
而生成的看板 HTML 与 Word 里「工作面容量」出现 **0 次** —— 容量在起作用，
交付物一个字都不提，用户只能质问"为什么资源这么少"。

本文件从**真产物**（build_plan_html / build_plan_docx 写出的文件）上钉四件事：
  · 工作面容量口径（条数 + 封顶条数 + 逐条例子）确实印出来；
  · 峰值人数带**口径**：user → 用户给定上限；model_estimate → 模型估算，非用户输入；
  · 设备对账把"未生效"的用户申报限额显形；
  · 缺数据时整段降级（一个字不印、绝不抛异常），不编数。

运行：python -m pytest backend/tests/test_delivery_workface_visible.py -q -p no:cacheprovider
"""

import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from pipeline.nodes.delivery import (          # noqa: E402
    _equipment_binding_rows, _peak_caliber, _workface_sentence,
    build_plan_docx, build_plan_html,
)

# 交付物目录的隔离由 tests/conftest.py 全局负责（DELIVERABLES_DIR / PLANS_DIR → 临时目录）。


# ---------------------------------------------------------------- 最小 plan
def _tasks():
    """两条各 5 天的任务：一条被工作面容量封顶、一条只是按公式算出来。"""
    return [
        {"task_id": "1.1.1", "task_name": "柱浇筑", "start_date": "2026-09-01",
         "finish_date": "2026-09-06", "duration_days": 5,
         "assigned_resources": {"钢筋工": 7},
         "_workface_note": ("（工作面容量：按工程量与每施工段算出 7 人 / None 台；"
                            "source_type=ai_estimate / confidence=LOW）"),
         "_workface_capped": [{"resource": "钢筋工", "kind": "labor",
                               "original_per_day": 36, "capped_per_day": 7,
                               "unit_basis": "每施工段",
                               "reason": "工程量按公式需 36 人，超过每施工段上限 7 人"}]},
        {"task_id": "1.1.2", "task_name": "梁浇筑", "start_date": "2026-09-01",
         "finish_date": "2026-09-06", "duration_days": 5,
         "assigned_resources": {"混凝土工": 4},
         "_workface_note": ("（工作面容量：按工程量与每施工段算出 4 人 / None 台；"
                            "source_type=ai_estimate / confidence=LOW）")},
    ]


def _plan(**over):
    plan = {
        "plan_id": "plan_workface_visible",
        "overview": {"project_name": "工作面可见性测试", "total_duration_days": 10,
                     "planned_start_date": "2026-09-01",
                     "planned_end_date": "2026-09-11", "critical_path_length": 2},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "1.1", "name": "结构", "sub_packages": [
                {"id": "1.1.1", "name": "柱浇筑", "duration_days": 5, "quantity": 10,
                 "unit": "m³", "work_type": "混凝土工程"},
                {"id": "1.1.2", "name": "梁浇筑", "duration_days": 5, "quantity": 12,
                 "unit": "m³", "work_type": "混凝土工程"}]}]}]},
        "dependencies": [],
        "cpm_result": {"total_duration_days": 10, "critical_path": ["1.1.1"],
                       "schedule": [{"task_id": "1.1.1", "es": 0, "ef": 5, "ls": 0, "lf": 5}]},
        "resource_demand": {"tasks": _tasks()},
        "key_milestones": [{"name": "开工", "date": "2026-09-01", "task_id": "1.1.1",
                            "description": "开工"}],
        "critical_path_tasks": [{"task_id": "1.1.1", "task_name": "柱浇筑",
                                 "start_date": "2026-09-01", "finish_date": "2026-09-06",
                                 "duration_days": 5, "assigned_resources": {"钢筋工": 7}}],
        "all_tasks_schedule": _tasks(),
        "resource_plan": {"total_manpower_days": 55, "peak_manpower": 11,
                          "equipment_peak": {"混凝土输送泵车": 1},
                          "material_summary": []},
        "risks": [],
        "report": "# 监督报告",
    }
    plan.update(over)
    return plan


def _html_text(plan):
    return Path(build_plan_html(plan)).read_text(encoding="utf-8")


def _doc_text(plan):
    """段落 + 所有表格单元格文本，拼成一个大字符串。"""
    from docx import Document
    doc = Document(build_plan_docx(plan))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parts.extend(c.text for c in r.cells)
    return "\n".join(parts)


# ================================================================
# D1 / D2：工作面容量口径（看板 + Word）
# ================================================================
def test_看板与Word都写明工作面容量口径():
    plan = _plan()
    html = _html_text(plan)
    word = _doc_text(plan)
    # 两个产物都必须出现"工作面容量"这四个字（修复前是 0 次）
    assert "工作面容量" in html, "看板必须写明人数是被工作面容量算出来的"
    assert "工作面容量" in word, "Word 必须写明人数是被工作面容量算出来的"
    # 说明里必须带**条数**与**封顶条数**（2 条按公式算，其中 1 条顶到上限）
    assert "共 2 条任务的班组人数按本施工段工程量用标定公式算出" in html
    assert "其中 1 条顶到上限" in html
    assert "共 2 条任务的班组人数按本施工段工程量用标定公式算出" in word
    assert "其中 1 条顶到上限" in word
    # 逐条例子（任务（资源）| 原始 → 上限 | 原因）也要真的落进去
    for txt in ("柱浇筑（钢筋工）", "36 → 7"):
        assert txt in html, txt
        assert txt in word, txt
    # 上限的语义必须写出来，否则用户还是不知道"顶到上限"意味着什么
    assert "工程量再大也不加人，只能延长工期" in html
    assert "工程量再大也不加人，只能延长工期" in word
    # 来源与置信度也必须跟着一起写（用户原话："来源=AI 估算/低置信度"）
    assert "工作面容量标定来源：ai_estimate / LOW" in html
    assert "工作面容量标定来源：ai_estimate / LOW" in word


def test_缺工作面容量数据时整段降级不抛异常():
    """没有 `_workface_note` → 「工作面容量」一个字都不印（不许编 0 条）。"""
    plain_tasks = [dict(t) for t in _tasks()]
    for t in plain_tasks:
        t.pop("_workface_note", None)
        t.pop("_workface_capped", None)
    plan = _plan(resource_demand={"tasks": plain_tasks})
    assert _workface_sentence(plan) == ""
    html = _html_text(plan)                     # 不许抛异常
    word = _doc_text(plan)
    assert "工作面容量" not in html
    assert "工作面容量" not in word
    # 空 resource_demand 同样退化（不是每份计划都有这个字段）
    empty = _plan(resource_demand={})
    assert _workface_sentence(empty) == ""
    assert "工作面容量" not in _html_text(empty)
    assert "工作面容量" not in _doc_text(empty)


# ================================================================
# D3：峰值人数的口径标注
# ================================================================
def test_模型估算的峰值标注为非用户输入():
    plan = _plan(resource_plan={"total_manpower_days": 55, "peak_manpower": 120,
                                "declared_peak_manpower": 120,
                                "curve_peak_manpower": 11,
                                "peak_manpower_source": "model_estimate",
                                "equipment_peak": {}})
    cal = _peak_caliber(plan, _view(plan))
    assert cal["source"] == "model_estimate"
    assert cal["label"] == "模型估算，非用户输入"
    html = _html_text(plan)
    word = _doc_text(plan)
    assert "模型估算，非用户输入" in html, "看板必须点明这个数不是用户给的"
    assert "模型估算，非用户输入" in word, "Word 必须点明这个数不是用户给的"
    # 绝不许再把它叫成"用户限额"
    assert "用户限额口径" not in word
    assert "120" in html and "120" in word
    # 逐日曲线峰值必须**并列**保留（藏掉它就是另一种撒谎）
    assert "每日用工峰值（按任务叠加）" in html
    assert "每日用工峰值（按任务叠加）" in word


def test_用户给定的峰值标注为用户给定上限():
    plan = _plan(resource_plan={"total_manpower_days": 55, "peak_manpower": 30,
                                "curve_peak_manpower": 11,
                                "peak_manpower_source": "user",
                                "equipment_peak": {}})
    assert _peak_caliber(plan, _view(plan))["label"] == "用户给定上限"
    html = _html_text(plan)
    word = _doc_text(plan)
    assert "用户给定上限" in html
    assert "用户给定上限" in word
    assert "模型估算" not in word, "用户真给了上限，就不许说成模型估算"


def test_资源曲线口径与认不出的来源都不编口径名():
    curve = _plan(resource_plan={"total_manpower_days": 55, "peak_manpower": 11,
                                 "peak_manpower_source": "resource_curve"})
    assert _peak_caliber(curve, _view(curve))["label"] == "资源曲线口径"
    assert "资源曲线口径" in _html_text(curve)
    # 认不出的来源原样打出来，不硬套一个口径名
    odd = _plan(resource_plan={"total_manpower_days": 55, "peak_manpower": 11,
                               "peak_manpower_source": "future_source"})
    assert _peak_caliber(odd, _view(odd))["label"] == "口径 future_source"
    assert "口径 future_source" in _html_text(odd)


def test_缺口径键时退回老显示不报错():
    """旧计划没有 `peak_manpower_source` → 退回"资源曲线口径 / 用户限额口径"，不炸。"""
    plan = _plan(meta={"boundary_conditions": {"labor": {"peak_total": 3}}},
                 resource_plan={"total_manpower_days": 55, "peak_manpower": 3})
    assert _peak_caliber(plan, _view(plan)) is None
    html = _html_text(plan)
    word = _doc_text(plan)
    assert "用户限额口径 3 人" in word
    assert "资源曲线口径 11 人" in word
    assert "<b>峰值人数：</b>3 人" in html
    assert "每日用工峰值（按任务叠加）" in html


# ================================================================
# E1（用户 2026-09-21 **重新裁定**）：**申报峰值不再是展示项**
# ------------------------------------------------------------
# 更早一轮的要求是「把 120 人改标成『模型估算，非用户输入』」（当时的用户原话，见
# 旧版本文件 §:213）；本次用户已重新裁定为**「连源头一起删、不再有申报峰值这一展示项」**
# （重构方案 §0.3 病根 3 / §5 阶段 5）—— 模型替用户补的 120 不再落盘
# （`plan_assembler` 已删 `declared_peak_manpower[_source]` 两个键），交付物里
# **一个字都不出现**「申报峰值」。
#
# ⚠ 只删「申报峰值」这一项：**峰值人数的口径区分必须保留**
#    （`user` → 用户给定上限；`resource_curve` → 资源曲线口径）。
# ================================================================
_CONF_META = {
    "norm_coverage": {"total": 10, "bound": 4, "bound_pct": 40.0},
    "credibility": {"user": 0.0, "kb": 0.3, "ai": 0.7},
}


def _plan_with_declared(src, peak_source="resource_curve"):
    """展示值是曲线峰值 11 人，另有申报值 120 人 —— 正是用户实测那份计划的形状。

    E1 之后这两个申报键**已不再是计划端会产出的键**；这里刻意留着它们，
    用来证明**即使旧计划 / 恶意数据里带着它们，交付物也一个字都不印**。
    """
    return _plan(meta=dict(_CONF_META),
                 resource_plan={"total_manpower_days": 55, "peak_manpower": 11,
                                "curve_peak_manpower": 11,
                                "declared_peak_manpower": 120,
                                "declared_peak_manpower_source": src,
                                "peak_manpower_source": peak_source,
                                "equipment_peak": {}})
def test_申报峰值不再是展示项():
    for src in ("model", "user", None, "expert_guess"):
        plan = _plan_with_declared(src)
        cal = _peak_caliber(plan, _view(plan))
        assert "declared" not in cal, cal
        assert "declared_text" not in cal and "declared_value_text" not in cal, cal
        html = _html_text(plan)
        word = _doc_text(plan)
        assert "申报峰值" not in html, (src, html)
        assert "申报峰值" not in word, (src, word)
        assert "计划里另记申报峰值" not in html
        assert "计划里另记申报峰值" not in word


def test_峰值人数的口径区分必须保留():
    """E1 删的是「申报峰值」，不是峰值人数的口径标注（用户限额 / 资源曲线）。"""
    curve = _plan(resource_plan={"total_manpower_days": 55, "peak_manpower": 11,
                                 "peak_manpower_source": "resource_curve"})
    assert _peak_caliber(curve, _view(curve))["label"] == "资源曲线口径"
    assert "资源曲线口径" in _html_text(curve)
    user = _plan(resource_plan={"total_manpower_days": 55, "peak_manpower": 30,
                                "peak_manpower_source": "user"})
    assert _peak_caliber(user, _view(user))["label"] == "用户给定上限"
    assert "用户给定上限" in _html_text(user)
    assert "用户给定上限" in _doc_text(user)


# ================================================================
# D4：数据来源与置信度章节
# ================================================================
def test_置信度章节写明工作面容量的来源与条数():
    plan = _plan(meta={
        "norm_coverage": {"total": 10, "bound": 4, "bound_pct": 40.0,
                          "unbound": 6, "unbound_pct": 60.0,
                          "workface_saturated": 1, "workface_ceiling_raised": 2},
        "credibility": {"user": 0.0, "kb": 0.3, "ai": 0.7},
    })
    word = _doc_text(plan)
    assert "工作面容量数据来源" in word
    assert "ai_estimate / LOW（经验标定，非规范来源，按工程量联动）" in word
    assert "其中顶到每施工段人数上限" in word and "1 条" in word
    assert "上限按第 39 轮口径抬高" in word and "2 条" in word


def test_缺workface统计键时那两行整行不出():
    """Agent B 没写 / 旧计划 → 只出"数据来源"一行，不许编 0 条。"""
    plan = _plan(meta={"norm_coverage": {"total": 10, "bound": 4, "bound_pct": 40.0},
                       "credibility": {"user": 0.0, "ai": 1.0}})
    word = _doc_text(plan)
    assert "工作面容量数据来源" in word
    assert "其中顶到每施工段人数上限" not in word
    assert "上限按第 39 轮口径抬高" not in word


def test_没有工作面容量时数据来源那行也不出():
    plan = _plan(meta={"norm_coverage": {"total": 10, "bound": 4, "bound_pct": 40.0},
                       "credibility": {"user": 0.0, "ai": 1.0}},
                 resource_demand={"tasks": []})
    assert "工作面容量数据来源" not in _doc_text(plan)


# ================================================================
# D5：设备对账（未生效必须显眼）
# ================================================================
_EB = {
    "塔吊": {"declared": 2, "bound_to": None, "effective": False,
             "note": "用户申报的「塔吊」未匹配到计划中的任何机械资源，该限额未生效"},
    "混凝土输送泵车": {"declared": 1, "bound_to": "混凝土输送泵车", "effective": True,
                       "note": "已绑定到计划资源「混凝土输送泵车」，限额 1 生效"},
}


def test_设备对账把未生效的限额显形():
    plan = _plan(meta={"equipment_binding": _EB})
    word = _doc_text(plan)
    assert "用户申报设备限额对账" in word
    assert "塔吊" in word and "未生效" in word
    assert "混凝土输送泵车" in word and "生效" in word
    assert "未匹配的设备限额没有参与排程，如需生效请在计划里给它们安排工序。" in word
    # 未生效的排在最前面（用户先看到问题项）
    assert word.index("塔吊") < word.index("已绑定到计划资源")
    html = _html_text(plan)
    assert "用户申报设备限额对账" in html
    assert "未匹配的设备限额没有参与排程" in html


def test_设备对账兼容列表形态与缺字段():
    rows = _equipment_binding_rows(_plan(meta={"equipment_binding": [
        {"name": "塔吊", "quantity": 2, "bound_to": None,
         "note": "用户申报的「塔吊」未匹配到计划中的机械资源，该限额未生效"}]}))
    assert len(rows) == 1 and rows[0]["effective"] is False and rows[0]["quantity"] == 2
    # 没有这段数据 → 整段不出
    assert _equipment_binding_rows(_plan()) == []
    assert "用户申报设备限额对账" not in _doc_text(_plan())
    assert "用户申报设备限额对账" not in _html_text(_plan())


def test_设备对账全生效时不打那句提示():
    plan = _plan(meta={"equipment_binding": {
        "混凝土输送泵车": _EB["混凝土输送泵车"]}})
    word = _doc_text(plan)
    assert "用户申报设备限额对账" in word
    assert "未匹配的设备限额没有参与排程" not in word


# ================================================================
# 辅助：直接拿 build_plan_docx/build_plan_html 用的那份 view
# ================================================================
def _view(plan):
    from pipeline.nodes.delivery import _compute_view
    return _compute_view(plan)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
