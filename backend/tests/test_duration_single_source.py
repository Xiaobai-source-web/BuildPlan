# -*- coding: utf-8 -*-
"""任务工期只有一个含义（P0-B）—— `duration_days` 必须与自己的日期同源

用户审计实测（真计划 `backend/plans/plan_sample3_after_allfix.json`）：
  · `all_tasks_schedule[*].duration_days` 与自身起止日期**对不上 144/310 条**
    （`2.1.1 预应力管桩施工` 字段 15 天 / 日期 `2026-06-29→2026-07-04` 只有 6 天；
      `3.2.1 基坑土方开挖` 字段 25 天 / 日期 3 天）。原因：字段装的是 **WBS 目标天数**，
      日期是**排程跨度**，两个不同的东西挤在同一个键里；
  · `critical_path_tasks` 是 **89 条任务**，交付物却印「关键路径长度 89」（把条数当天数）；
  · 同一份计划并存四个天数（总工期 608 / 理论最短 608 / 纯 CPM 644 / 关键链合计 712），
    交付物一个都没标语义。

本文件守住：
  1. 新口径下 `duration_days == (finish_date - start_date).days + 1`，**真计划逐行 0 例外**；
  2. `wbs_target_days` 保留原值（310/310 与旧字段一致，也与 WBS 叶子一致）；
  3. 关键路径的**条数**与**工期**两行各自正确（数字来自数据，不是编的）；
  4. 四个天数的语义与来源键都印出来；取不到的行**不出现**。

运行：python -m pytest backend/tests/test_duration_single_source.py -q -p no:cacheprovider
"""

import datetime
import json
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

import pytest

REAL_PLAN = BACKEND / "plans" / "plan_sample3_after_allfix.json"


# ==================== 夹具 ====================
def _load_real():
    return json.loads(REAL_PLAN.read_text(encoding="utf-8"))


def _assemble(plan):
    """用**真** `plan_assembler.build_parts` 重新组装一遍（这就是修好的那条代码路径）。

    返回 `(parts, 合并后的 cpm_result)`：后者带上了口径键
    （`cpm_duration_basis` / `cpm_duration_note`），新计划里就是这份数据。
    """
    from pipeline.nodes import plan_assembler as PA

    ov = plan.get("overview") or {}
    params = dict(plan.get("extracted_params") or {})
    params.setdefault("planned_start_date", ov.get("planned_start_date") or "2026-06-01")
    params.setdefault("project_name", ov.get("project_name") or "未命名")
    ctx = {
        "wbs": json.loads(json.dumps(plan.get("wbs") or {})),
        "cpm_result": json.loads(json.dumps(plan.get("cpm_result") or {})),
        "resource_demand": json.loads(json.dumps(plan.get("resource_demand") or {})),
        "extracted_params": params,
        "boundary_conditions": plan.get("boundary_conditions"),
    }
    return PA.build_parts(ctx), ctx.get("cpm_result")


def _span(row):
    d0 = datetime.date.fromisoformat(str(row["start_date"]))
    d1 = datetime.date.fromisoformat(str(row["finish_date"]))
    return (d1 - d0).days + 1


@pytest.fixture(scope="module")
def real():
    if not REAL_PLAN.exists():
        pytest.skip("真计划不存在：%s" % REAL_PLAN)
    plan = _load_real()
    parts, cpm = _assemble(plan)
    fresh = json.loads(json.dumps(plan))
    fresh["all_tasks_schedule"] = parts["all_tasks_schedule"]
    fresh["critical_path_tasks"] = parts["critical_path_tasks"]
    fresh["cpm_result"] = cpm
    return {"raw": plan, "parts": parts, "fresh": fresh}


def _plan_two_tasks():
    """两行的小计划：字段与日期**故意不同源**，用来盯住"同一行两个工期"。"""
    return {
        "plan_id": "dur_two",
        "overview": {"project_name": "工期口径测试", "total_duration_days": 30,
                     "planned_start_date": "2026-01-01", "planned_end_date": "2026-01-30",
                     "critical_path_length": 2},
        "meta": {"audit_status": "未审计",
                 "audit_rounds": [{"round": 1, "name": "WBS 结构", "passed": True}],
                 "schedule_versions": {"theory_min_days": 30, "resource_ok_days": 35,
                                       "delta_days": 5}},
        "cpm_result": {"total_duration_days": 35, "cpm_total_duration_days": 48,
                       "critical_path": ["1.1", "1.2"],
                       "schedule": [{"task_id": "1.1", "es": 0, "ef": 5},
                                    {"task_id": "1.2", "es": 5, "ef": 35}]},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "1", "name": "wp", "sub_packages": [
                {"id": "1.1", "name": "甲", "quantity": 1, "unit": "t", "duration_days": 20},
                {"id": "1.2", "name": "乙", "quantity": 1, "unit": "t", "duration_days": 28}]}]}]},
        "dependencies": [],
        "all_tasks_schedule": [
            {"task_id": "1.1", "task_name": "甲", "start_date": "2026-01-01",
             "finish_date": "2026-01-05", "duration_days": 5, "wbs_target_days": 20,
             "assigned_resources": {"普工": 3}},
            {"task_id": "1.2", "task_name": "乙", "start_date": "2026-01-06",
             "finish_date": "2026-02-04", "duration_days": 30, "wbs_target_days": 28,
             "assigned_resources": {"普工": 3}}],
        "critical_path_tasks": [],
        "key_milestones": [],
        "resource_plan": {"peak_manpower": 3, "total_manpower_days": 90.0,
                          "equipment_peak": {}, "material_summary": []},
        "risks": [], "report": "报告",
    }


def _docx_parts(plan):
    """→ (doc, 全文, **表格文字**)。

    表格文字单独取一份：监督报告那一段是**存档的原文**（真计划里是修好之前生成的
    Markdown，仍写着「关键路径长度」），本文件只对"交付物自己算出来的表格/段落"断言。
    """
    from docx import Document

    from pipeline.nodes.delivery import build_plan_docx
    doc = Document(build_plan_docx(plan, draft=False))
    text = [p.text for p in doc.paragraphs]
    cells = []
    for t in doc.tables:
        for row in t.rows:
            cells.extend(c.text for c in row.cells)
    return doc, "\n".join(text + cells), "\n".join(cells)


def _html_text(plan):
    """`build_plan_html` 落盘后返回**路径**（它是导出节点），这里读回文本。"""
    from pipeline.nodes.delivery import build_plan_html
    return Path(build_plan_html(plan)).read_text(encoding="utf-8")


# ==================== 1. duration_days 与日期同源 ====================
def test_real_plan_duration_matches_date_span(real):
    """**逐行**断言真计划：新口径下 `duration_days` == 日期跨度，0 例外。"""
    rows = real["parts"]["all_tasks_schedule"]
    assert len(rows) == 310, len(rows)
    bad = [(t["task_id"], t.get("duration_days"), _span(t)) for t in rows
           if int(t.get("duration_days") or 0) != _span(t)]
    assert not bad, "仍有字段与日期对不上的行：%s" % bad[:5]


def test_real_plan_old_format_had_144_mismatches(real):
    """冻结旧字段的缺陷指纹：144/310 条对不上（这条用例是防止"缺陷样本悄悄变了"）。

    如果这个数变了，说明真计划被重新生成过 —— 那时要重新核对本文件的基线，
    而不是把断言改成"随便多少都行"。
    """
    raw = real["raw"]["all_tasks_schedule"]
    bad = [t for t in raw
           if int(t.get("duration_days") or 0) != _span(t)]
    assert len(raw) == 310
    assert len(bad) == 144, "旧字段的不一致条数变了：%d" % len(bad)
    assert [t["task_id"] for t in bad[:3]] == ["2.1.1", "2.1.2", "2.3.1"]
    assert (bad[0]["duration_days"], _span(bad[0])) == (15, 6)


def test_wbs_target_days_preserves_original_value(real):
    """`wbs_target_days` = WBS 目标天数原值（与新字段 `duration_days` 是两个东西）。"""
    old = {str(t["task_id"]): t.get("duration_days")
           for t in real["raw"]["all_tasks_schedule"]}
    rows = real["parts"]["all_tasks_schedule"]
    assert all(t.get("wbs_target_days") == old.get(str(t["task_id"])) for t in rows)
    # 与 WBS 叶子上的原值也必须一致（不是二次加工出来的数）
    leaves = {}
    for ph in (real["raw"].get("wbs") or {}).get("phases", []):
        for wp in ph.get("work_packages", []):
            for sub in wp.get("sub_packages", []):
                leaves[str(sub.get("id"))] = sub.get("duration_days")
    assert all(t.get("wbs_target_days") == leaves.get(str(t["task_id"])) for t in rows)
    # 真计划里两者确实不同（否则这条用例没有意义）
    assert sum(1 for t in rows if t.get("wbs_target_days") != t.get("duration_days")) == 144


def test_real_plan_critical_chain_two_sums(real):
    """89 条关键任务的两种合计：WBS 目标 712 / 排程跨度 673（都不是"工期"）。"""
    from pipeline.nodes.delivery import _critical_chain_sums

    fresh = real["fresh"]
    sums = _critical_chain_sums(fresh)
    assert sums["count"] == 89
    assert sums["wbs_target_days"] == 712
    assert sums["span_days"] == 673
    # 关键路径**条数**与**工期**都必须来自数据
    ov = fresh["overview"]
    cpm = fresh["cpm_result"]
    assert ov["critical_path_length"] == 89 != ov["total_duration_days"]
    assert cpm["total_duration_days"] == 608
    assert cpm["cpm_total_duration_days"] == 644


# ==================== 2. 关键路径：条数 ≠ 工期 ====================
def _overview_table_text(doc):
    """计划总览那张表（含「审计状态」行）的文字。

    只断言这一张表：真计划的**监督报告是存档原文**（修好之前生成的 Markdown，
    仍写着「关键路径长度」），交付物照抄存档内容不算它自己撒谎；复检脚本按全文
    匹配，所以那一段要等计划重新生成才会消失。
    """
    for t in doc.tables:
        cells = [c.text for row in t.rows for c in row.cells]
        if any("审计状态" in c for c in cells):
            return "\n".join(cells)
    raise AssertionError("找不到计划总览表")


def test_real_plan_docx_separates_count_and_days(real):
    doc, text, _tables = _docx_parts(real["fresh"])
    ov = _overview_table_text(doc)
    assert "关键路径长度" not in ov, "「长度 89」把条数读成了天数"
    assert "关键路径任务数" in ov
    assert "89 个（来源 overview.critical_path_length）" in ov
    assert "608 天（来源 cpm_result.total_duration_days）" in ov
    assert "644 天（来源 cpm_result.cpm_total_duration_days" in ov


def test_template_report_says_task_count_not_length():
    """模板监督报告也不许写「关键路径长度」（`plan_assembler.template_report`）。

    这是复检脚本 `q_audit_check.py` 的禁语检查项：`关键路径长度` 一旦出现在交付物里，
    读者就会把 89 **条**读成 89 **天**。
    """
    from pipeline.nodes import plan_assembler as PA

    parts = {
        "overview": {"project_name": "X", "total_duration_days": 100,
                     "planned_start_date": "2026-01-01", "planned_end_date": "2026-04-10",
                     "critical_path_length": 89},
        "key_milestones": [], "critical_path_tasks": [],
        "resource_plan": {"peak_manpower": 3, "total_manpower_days": 90.0,
                          "equipment_peak": {}},
        "risks": [],
        "display_granularity": {"note": "逐任务"},
    }
    report = PA.template_report(parts)
    assert "关键路径长度" not in report
    assert "关键路径任务数：89 个" in report


def test_real_plan_docx_labels_four_durations(real):
    _doc, text, _tables = _docx_parts(real["fresh"])
    for label in ("总工期（排程实排）", "理论最短（排程版）", "纯 CPM（未计资源约束）",
                  "关键链合计（WBS 目标天数）", "关键链合计（排程跨度）"):
        assert label in text, label
    assert "608 天（来源 overview.total_duration_days" in text
    assert "608 天（来源 meta.schedule_versions.theory_min_days" in text
    assert "644 天（来源 cpm_result.cpm_total_duration_days" in text
    assert "712 天（89 条关键任务的 wbs_target_days 之和" in text
    assert "673 天（89 条关键任务的日期跨度之和" in text


def test_missing_keys_render_nothing():
    """取不到的数据**整行不出现** —— 绝不用别的数顶上。"""
    plan = _plan_two_tasks()
    plan["cpm_result"] = {}                      # 没有纯 CPM 值
    plan["meta"].pop("schedule_versions", None)  # 没有理论最短
    plan["critical_path_tasks"] = []
    _doc, text, _tables = _docx_parts(plan)
    assert "纯 CPM" not in text
    assert "理论最短（排程版）" not in text
    assert "关键链合计" not in text


# ==================== 3. 同一行不许出现两个互相矛盾的工期 ====================
def test_docx_task_row_span_and_wbs_target_are_both_right():
    plan = _plan_two_tasks()
    doc, text, _tables = _docx_parts(plan)
    hdr, rows = None, None
    for t in doc.tables:
        head = [c.text.strip() for c in t.rows[0].cells]
        if "工期(天·排程)" in head and "WBS 目标(天)" in head:
            hdr, rows = head, t.rows[1:]
            break
    assert hdr, "找不到甘特/明细表"
    i_id, i_span, i_wbs = hdr.index("任务 ID"), hdr.index("工期(天·排程)"), hdr.index("WBS 目标(天)")
    by_id = {str(t["task_id"]): t for t in plan["all_tasks_schedule"]}
    checked = 0
    for row in rows:
        tid = row.cells[i_id].text.strip()
        if tid not in by_id:
            continue
        task = by_id[tid]
        assert int(row.cells[i_span].text) == _span(task), (tid, row.cells[i_span].text)
        assert int(row.cells[i_wbs].text) == task["wbs_target_days"], (tid, row.cells[i_wbs].text)
        checked += 1
    assert checked == 2, checked
    # 1.1：排程跨度 5 天 / WBS 目标 20 天 —— 同一行两个数，但列名各说各的
    assert "工期(天·排程)" in text and "WBS 目标(天)" in text


def test_cpm_caliber_is_recorded_together_with_the_number(real):
    """644 与 608 的**口径**必须随数字落盘（否则下游只会看到"理想值比现实值还长"）。

    `q_audit_check.py` 的「CPM 与排程口径可解释」按 `cpm_total > theory_min` 判 FAIL。
    本条用例记录的是：**这个不等式本身是对的**（两边不同源），口径键与说明都已落盘；
    把这个不等式当成"算错"去改数才是错的。
    """
    from pipeline.nodes import plan_assembler as PA

    cpm = real["fresh"]["cpm_result"]
    assert cpm["cpm_total_duration_days"] == 644
    assert cpm["total_duration_days"] == 608
    assert cpm["cpm_duration_basis"] == "wbs_target_days+dependencies"
    assert "不同源" in cpm["cpm_duration_note"]
    # 排程版的口径（由 build_meta 落进 meta.schedule_versions）
    meta = PA.build_meta({"schedule_versions": {"theory_min": {"total_duration_days": 608},
                                                "resource_ok": {"total_duration_days": 608},
                                                "compare": {"delta_days": 0}},
                          "extracted_params": {"planned_start_date": "2026-01-01"}})
    sv = meta["schedule_versions"]
    assert sv["theory_min_days"] == 608
    assert sv["theory_min_basis"] == "scheduler_task_durations"
    assert sv["comparable_to_cpm"] is False


def test_wbs_target_unknown_is_reported_as_unknown():
    """缺 `wbs_target_days` 又**判不出**语义时，交付物印「—」，不拿跨度冒充目标。

    `recompute.py`（重排/修订路径）重建 `all_tasks_schedule` 时只写跨度口径、
    不带 `wbs_target_days`。老计划同样缺这个键，但那时 `duration_days` 装的是 WBS 目标。
    两种情形只能靠"字段是否等于本行日期跨度"区分；相等就是分不清 → 报不知道。
    """
    from pipeline.nodes.delivery import _task_wbs_target

    same = {"task_id": "a", "start_date": "2026-01-01", "finish_date": "2026-01-05",
            "duration_days": 5}
    assert _task_wbs_target(same) is None
    old = {"task_id": "b", "start_date": "2026-01-01", "finish_date": "2026-01-06",
           "duration_days": 20}
    assert _task_wbs_target(old) == 20, "字段≠日期跨度 → 老计划里这就是 WBS 目标"
    new = dict(old, wbs_target_days=7)
    assert _task_wbs_target(new) == 7, "新计划只读新键"

    plan = _plan_two_tasks()
    for t in plan["all_tasks_schedule"]:
        t.pop("wbs_target_days", None)
    doc, _text, _tables = _docx_parts(plan)
    for tbl in doc.tables:
        head = [c.text.strip() for c in tbl.rows[0].cells]
        if "工期(天·排程)" in head and "WBS 目标(天)" in head:
            idx = head.index("WBS 目标(天)")
            values = [r.cells[idx].text.strip() for r in tbl.rows[1:]]
            assert values and all(v == "—" for v in values), values
            return
    raise AssertionError("找不到甘特表")


def test_cpm_with_schedule_durations_is_not_longer_than_the_schedule(real):
    """`644 > 608` 的**唯一**成因是逐任务工期不同源 —— 不是算错（本条是它的定量证据）。

    同一张依赖图，只把每任务的工期换成 scheduler 的排程跨度（`ef - es`）再正推一遍：
    得到 **605 ≤ 608**（scheduler 自己的总工期）。也就是说"无资源约束的理想值 ≤
    资源约束版"这条定律在**同一套工期**下是成立的；644 之所以更大，是因为它用的是
    WBS 目标天数，而两者在 **144/310** 条任务上不同（`2.1.1` WBS 15 天 / 排程 6 天）。
    """
    from pipeline.nodes.cpm import calculate_cpm

    raw = real["raw"]
    sched = {s["task_id"]: s for s in raw["cpm_result"]["schedule"]}
    wbs = json.loads(json.dumps(raw["wbs"]))
    n = 0
    for ph in wbs["phases"]:
        for wp in ph["work_packages"]:
            for sub in wp["sub_packages"]:
                row = sched.get(sub["id"])
                if row is not None:
                    sub["duration_days"] = max(1, int(row["ef"]) - int(row["es"]))
                    n += 1
    assert n == 310, n
    same_basis = calculate_cpm(wbs, {"dependencies": raw["dependencies"]})
    assert same_basis["total_duration_days"] == 605
    assert same_basis["total_duration_days"] <= raw["cpm_result"]["total_duration_days"]
    # 两个口径的工期差异面（真计划基线）
    assert raw["cpm_result"]["cpm_total_duration_days"] == 644
    assert raw["cpm_result"]["total_duration_days"] == 608


def test_dashboard_task_table_uses_wbs_target_not_span():
    """看板逐条工序表的「WBS 目标(天)」列不能印排程跨度（否则两列同值=白加一列）。"""
    plan = _plan_two_tasks()
    html = _html_text(plan)
    assert "WBS 目标" in html
    assert ">20</td>" in html or ">20<" in html, "WBS 目标 20 天必须出现在表里"
    assert "关键路径任务数" in html


def _kv_cards(html):
    """看板顶部摘要卡片 → ``{键: 值}``（只取 `.kv` 那一段，不碰页内表格）。"""
    import re
    i = html.index("class='kv'")
    seg = html[i:html.index("</div></div>", i)]
    return dict((k, v) for v, k in
                re.findall(r"<div><b>(.*?)</b><span>(.*?)</span></div>", seg))


def test_dashboard_summary_cards_show_result_only(real):
    """摘要卡片**只写结果**：不许把来源键（`overview.total_duration_days` 这种内部字段名）
    印上去。

    用户实测原话：「这个计划看板的摘要栏怎么这么怪异，摘要栏不需要把每条数据的来源也写进去，
    只写结果」—— 当时卡片上是「604 天（来源 overview.total_duration_days；满足工作面/资源
    约束后的实排工期）」。来源与口径的去处没变：Word 计划总览表仍带「（来源 …）」（见
    `test_real_plan_docx_separates_count_and_days`），看板的口径说明在「资源计划」卡片与
    施工组织口径段里。
    """
    html = _html_text(real["fresh"])
    cards = _kv_cards(html)
    assert cards, "找不到看板顶部的摘要卡片"
    bad = {k: v for k, v in cards.items() if "来源" in v or "（" in v or "(" in v}
    assert not bad, "摘要卡片里还有来源/注解：" + repr(bad)
    # 数还是那些数（只去注解，不改数）；总工期卡补上单位
    assert cards["关键路径任务数"] == "89 个"
    assert cards["关键路径工期·排程版"] == "608 天"
    assert cards["关键路径工期·纯 CPM（按 WBS 目标天数计价）"] == "644 天"
    assert cards["总工期(天)"] == "608 天"
    assert cards["关键链合计（WBS 目标天数）"] == "712 天"
    # Word 计划总览的来源键一个字都不许少（可追溯性没有降级）
    _doc, text, _tables = _docx_parts(real["fresh"])
    assert "608 天（来源 cpm_result.total_duration_days）" in text
    assert "89 个（来源 overview.critical_path_length）" in text


def test_dashboard_labels_four_durations(real):
    html = _html_text(real["fresh"])
    for label in ("总工期（排程实排）", "理论最短（排程版）", "纯 CPM（未计资源约束）",
                  "关键链合计（WBS 目标天数）", "关键链合计（排程跨度）"):
        assert label in html, label


# ==================== 4. 「该有的也不许静默丢」：WBS 叶子缺排程行 ====================
def _mini_ctx(task_ids, scheduled_ids):
    """小 ctx：`task_ids` 全在 WBS 里，但只有 `scheduled_ids` 有排程行。"""
    leaves = [{"id": tid, "name": "工序" + tid, "quantity": 1, "unit": "t",
               "duration_days": 5} for tid in task_ids]
    sched = [{"task_id": tid, "es": 5 * i, "ef": 5 * (i + 1)}
             for i, tid in enumerate(scheduled_ids)]
    return {
        "wbs": {"phases": [{"phase": "P", "work_packages": [
            {"id": "1", "name": "wp", "sub_packages": leaves}]}]},
        "cpm_result": {"total_duration_days": 5 * len(scheduled_ids),
                       "critical_path": list(scheduled_ids), "schedule": sched},
        "resource_demand": {"tasks": []},
        "extracted_params": {"planned_start_date": "2026-01-01", "project_name": "缺行测试"},
        "dependencies": {"dependencies": []},
    }


def test_unscheduled_leaf_is_recorded_not_silently_dropped():
    """WBS 有叶子、排程没给行 → **不许静默 continue**：必须记账 + 告警 + 落 meta。"""
    from pipeline.nodes import plan_assembler as PA

    ctx = _mini_ctx(["1.1", "1.2", "1.3"], ["1.1", "1.2"])
    parts = PA.build_parts(ctx)
    ids = [t["task_id"] for t in parts["all_tasks_schedule"]]
    assert ids == ["1.1", "1.2"], ids
    uns = parts["unscheduled_tasks"]
    assert [u["task_id"] for u in uns] == ["1.3"], uns
    assert uns[0]["task_name"] == "工序1.3"
    assert "排程" in uns[0]["reason"]
    assert "未进入" in uns[0]["effect"]

    # 走真节点：告警必须发出去（引擎会把它收进 meta.node_warnings）
    node = PA.PlanAssemblerNode()
    events = []
    node._emit = lambda e, d: events.append((e, d))
    node._run_id = "unsched_test"
    out = node.run(ctx)
    warns = [d for e, d in events if e == "warning"]
    assert len(warns) == 1, events
    assert "1 条 WBS 工序没有排程结果" in warns[0]["message"]
    assert "1.3" in warns[0]["detail"] and "工序1.3" in warns[0]["detail"]
    assert "没有进入交付物" in warns[0]["detail"]
    assert "1 条 WBS 工序没有排程结果" in node.done_summary, "摘要里也要看得见"
    # meta 里恒有这份账（差集是唯一真源）
    assert out["plan_json"]["meta"]["unscheduled_tasks"][0]["task_id"] == "1.3"
    # 报告模板也要写出来（"表格里没有"不等于"这活不用干"）
    assert "未进入交付物" in PA.template_report(parts)


def test_unscheduled_leaf_missing_es_ef_is_also_recorded():
    """排程行在、但缺 es/ef → 同样排不出日期，也要记账（不能瞎写一个日期）。"""
    from pipeline.nodes import plan_assembler as PA

    ctx = _mini_ctx(["1.1", "1.2"], ["1.1"])
    ctx["cpm_result"]["schedule"].append({"task_id": "1.2"})
    parts = PA.build_parts(ctx)
    assert [t["task_id"] for t in parts["all_tasks_schedule"]] == ["1.1"]
    assert [u["task_id"] for u in parts["unscheduled_tasks"]] == ["1.2"]
    assert "缺少 es/ef" in parts["unscheduled_tasks"][0]["reason"]


def test_unscheduled_warning_reaches_the_deliverable():
    """告警要一路走到交付物（`meta.node_warnings` → 看板的「节点级告警」卡片）。

    只 emit 不够：`engine.collect_node_warning` 把它收进 ctx，`build_meta` 再落进 meta，
    交付物读 meta —— 这条链路任何一环断了，用户就只看到一张少了任务的表。
    """
    plan = _plan_two_tasks()
    plan["meta"]["node_warnings"] = [{
        "node": "assembler",
        "message": "有 1 条 WBS 工序没有排程结果，未进入交付物",
        "detail": "  · 3.3.1 基坑监测 —— 排程结果里没有这条工序的行；"
                  "该任务未进入 all_tasks_schedule —— 交付物里看不到它，也不占用工期",
        "count": 1}]
    plan["meta"]["node_warning_count"] = 1
    html = _html_text(plan)
    assert "3.3.1" in html and "基坑监测" in html
    assert "没有排程结果" in html


def test_real_plan_has_no_unscheduled_tasks(real):
    """正常计划 310/310 → 差集为空、零告警（这条字段不能恒有内容）。"""
    from pipeline.nodes import plan_assembler as PA

    parts = real["parts"]
    assert parts["unscheduled_tasks"] == []
    assert PA.unscheduled_effect_text([]) == ""
    # 310 条 WBS 叶子全都进了排程
    leaves = 0
    for ph in (real["raw"].get("wbs") or {}).get("phases", []):
        for wp in ph.get("work_packages", []):
            leaves += len(wp.get("sub_packages") or [])
    assert leaves == 310 and len(parts["all_tasks_schedule"]) == 310
