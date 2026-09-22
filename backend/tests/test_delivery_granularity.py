# -*- coding: utf-8 -*-
"""A2a 回归：展示粒度必须真的作用到交付物，且**只改展示**。

背景：用户在粒度门里选的两维粒度（① 工序拆解深度 ② 楼层分组）此前**只影响对话**
—— `plan_level` 记了选择、终端显示了行数，但导出的 Word/看板仍然铺全部叶子
（`display_granularity` / `group_rows` 在生产代码里没有任何消费者）。选择被问了、
被记了、却没被执行。

修法：`meta.display_granularity` 落进计划；`delivery` 据此合并展示行。

**三条铁律（本文件就是在守这三条）**：
  ① 未合并时行为与旧版逐字节一致（默认路径不变）；
  ② 合并行的工期取组内任务的**排程时间跨度**，绝不重算
     → 否则粗粒度会给出与细粒度不同的总工期，同一份计划自相矛盾；
  ③ 资源曲线**不合并**（人员/设备是物理量，与"看多粗"无关）。

运行：python -m pytest backend/tests/test_delivery_granularity.py -q
"""

import json
import shutil
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from pipeline.nodes import delivery as D  # noqa: E402

START = "2026-03-01"


def _leaf(tid, name, location, step, wtype, qty, unit, dur):
    return {"id": tid, "name": name, "location": location, "duration_days": dur,
            "quantity": qty, "unit": unit, "work_type": wtype, "_step_name": step,
            "_qty_source": "参数推算", "_qty_formula": "测试公式"}


_LEAVES = [
    # 两层 × 两道工序（同工种跨层 → 粗粒度可合并）
    _leaf("5.1.1.1", "Ⅰ区 1-1层 钢筋绑扎", "Ⅰ区 1-1层", "钢筋绑扎", "钢筋工程", 10.0, "t", 2),
    _leaf("5.1.1.2", "Ⅰ区 1-1层 模板安装", "Ⅰ区 1-1层", "模板安装", "模板工程", 100.0, "m²", 2),
    _leaf("5.1.2.1", "Ⅰ区 2-2层 钢筋绑扎", "Ⅰ区 2-2层", "钢筋绑扎", "钢筋工程", 12.0, "t", 3),
    _leaf("5.1.2.2", "Ⅰ区 2-2层 模板安装", "Ⅰ区 2-2层", "模板安装", "模板工程", 120.0, "m²", 3),
    # 分层外（全楼）→ 不参与楼层分组
    _leaf("8.1.1.1", "全楼 内墙抹灰", "全楼", "内墙抹灰", "抹灰工程", 900.0, "m²", 5),
]

_SCHED = [
    # (task_id, es, ef, 人)
    ("5.1.1.1", 0, 2, {"钢筋工": 4}),
    ("5.1.1.2", 2, 4, {"模板工": 6}),
    ("5.1.2.1", 4, 7, {"钢筋工": 5}),
    ("5.1.2.2", 7, 10, {"模板工": 6}),
    ("8.1.1.1", 10, 15, {"抹灰工": 3}),
]


def _day(n):
    import datetime
    return (datetime.date.fromisoformat(START) + datetime.timedelta(days=n)).isoformat()


def _plan(granularity=None):
    tasks = [{"task_id": t, "task_name": t, "start_date": _day(a),
              "finish_date": _day(b), "duration_days": max(1, b - a),
              "assigned_resources": dict(r)} for t, a, b, r in _SCHED]
    meta = {"audit_status": "已审计"}
    if granularity is not None:
        meta["display_granularity"] = granularity
    return {
        "plan_id": "plan_test_a2a",
        "overview": {"project_name": "测试", "total_duration_days": 15,
                     "planned_start_date": START, "planned_end_date": _day(15),
                     "critical_path_length": 1},
        "wbs": {"phases": [
            {"phase": "地上主体结构", "work_packages": [
                {"id": "5.1", "name": "主体结构", "sub_packages": _LEAVES[:4]}]},
            {"phase": "装饰装修", "work_packages": [
                {"id": "8.1", "name": "装饰装修", "sub_packages": _LEAVES[4:]}]},
        ]},
        "dependencies": [],
        "cpm_result": {
            "total_duration_days": 15, "critical_path": ["5.1.1.1"],
            "schedule": [{"task_id": t, "es": a, "ef": b, "ls": a, "lf": b}
                         for t, a, b, _ in _SCHED],
        },
        "all_tasks_schedule": tasks,
        "resource_demand": {"tasks": []},
        "resource_plan": {"total_manpower_days": 0, "peak_manpower": 0,
                          "equipment_peak": {}, "material_summary": []},
        "report": "测试报告",
        "meta": meta,
    }


COARSE_WHOLE = {"depth": "coarse", "floor_grouping": "whole", "rows": 3}


# ---------------- ① 默认路径必须一字不变 ----------------
class TestDefaultUnchanged:
    def test_未选粒度时不合并(self):
        p = _plan()
        assert D._is_rolled_up(p) is False
        assert D.rolled_rows(p) == []
        assert D.granularity_note(p) == ""
        assert len(D._compute_view(p)["gantt"]) == len(_LEAVES)

    def test_未知粒度值逐维回退默认(self):
        """两个维度**互相独立**：一个非法只影响它自己，不牵连另一个。"""
        cases = [
            # 非法 depth + 合法 grouping → depth 回退工序级，grouping 保留整栋
            ({"depth": "bogus", "floor_grouping": "whole"}, ("component", "whole")),
            # 合法 depth + 非法 grouping → grouping 回退按层
            ({"depth": "coarse", "floor_grouping": "bogus"}, ("coarse", "per_floor")),
            # 都取不到 → 完全回到默认（逐叶子，不合并）
            ({"depth": None, "floor_grouping": None}, ("component", "per_floor")),
        ]
        for bad, want in cases:
            assert D._display_granularity(_plan(bad)) == want, bad
        assert D._is_rolled_up(_plan({"depth": None, "floor_grouping": None})) is False
        # 单维有效 → 仍然要合并（这是合法组合，不是"整体失效"）
        assert D._is_rolled_up(_plan({"depth": "bogus", "floor_grouping": "whole"})) is True


# ---------------- ② 合并后：行数下降、工期不重算 ----------------
class TestRolledUp:
    def test_横道行数真的下降(self):
        p = _plan(dict(COARSE_WHOLE))
        assert D._is_rolled_up(p) is True
        view = D._compute_view(p)
        assert len(view["gantt"]) == 3, [g["name"] for g in view["gantt"]]
        assert len(view["gantt"]) < len(_LEAVES)

    def test_总工期不变(self):
        base = D._compute_view(_plan())
        rolled = D._compute_view(_plan(dict(COARSE_WHOLE)))
        assert rolled["total_days"] == base["total_days"], \
            "上卷绝不能改变总工期（否则同一份计划自相矛盾）"

    def test_合并行取组内时间跨度不重算(self):
        p = _plan(dict(COARSE_WHOLE))
        rolled = D.rolled_rows(p)
        # 钢筋绑扎跨 1、2 层：起止应为第 0 天 → 第 7 天（跨度 7），不是 2+3=5
        g = [x for x in rolled if "钢筋" in (x.get("工序/工种") or "")][0]
        assert g["工期"] == 7, g["工期"]
        assert g["工期口径"] == "排程时间跨度"

    def test_全部叶子都被覆盖且不重复(self):
        p = _plan(dict(COARSE_WHOLE))
        covered = [i for g in D.rolled_rows(p) for i in g["ids"]]
        assert sorted(covered) == sorted(l["id"] for l in _LEAVES), \
            "上卷只能重新分组，不能丢任务也不能重复计"

    def test_工程量合计不变(self):
        p = _plan(dict(COARSE_WHOLE))
        rolled = D.rolled_rows(p)
        for unit, total in (("t", 22.0), ("m²", 1120.0)):
            got = sum(g["工程量"] for g in rolled if g["单位"] == unit)
            assert abs(got - total) < 0.01, (unit, got, total)

    def test_口径说明写清楚且含行数(self):
        note = D.granularity_note(_plan(dict(COARSE_WHOLE)))
        assert "展示粒度" in note and "3 行" in note
        assert "排程时间跨度" in note
        assert "总工期与逐叶子口径一致" in note


# ---------------- ③ 资源曲线绝不能因合并而变 ----------------
class TestCurvesUntouched:
    def test_人员与设备曲线逐日相同(self):
        base = D._compute_view(_plan())
        rolled = D._compute_view(_plan(dict(COARSE_WHOLE)))
        assert rolled["labor_daily"] == base["labor_daily"], \
            "人员曲线是物理量，合并展示行不得改动"
        assert rolled["equip_daily"] == base["equip_daily"]
        assert rolled["peak_total"] == base["peak_total"]
        assert rolled["peak_trade"] == base["peak_trade"]


# ---------------- ④ 真的落到交付物文件里 ----------------
class TestDeliverablesHonorGranularity:
    def test_HTML的行数随粒度变化(self):
        import re
        out = None
        try:
            p = _plan(dict(COARSE_WHOLE))
            out = Path(D.build_plan_html(p))
            html = out.read_text(encoding="utf-8")
            seg = html.split("工作分解结构")[1].split("</table>")[0]
            assert len(re.findall(r"<tr>", seg)) == 4, "3 行数据 + 1 行表头"
            assert "展示粒度：工种级（粗） × 整栋（3 行）" in html
        finally:
            if out is not None:
                shutil.rmtree(out.parent, ignore_errors=True)

    def test_Word的横道行数随粒度变化(self):
        import docx
        out = None
        try:
            p = _plan(dict(COARSE_WHOLE))
            out = Path(D.build_plan_docx(p))
            doc = docx.Document(str(out))
            rows = None
            for t in doc.tables:
                if "任务 ID" in [c.text.strip() for c in t.rows[0].cells]:
                    rows = len(t.rows) - 1
                    break
            assert rows == 3, "Word 横道应按粒度合并为 3 行，实际 %s" % rows
            text = "\n".join(x.text for x in doc.paragraphs)
            assert "展示粒度：工种级（粗） × 整栋（3 行）" in text
        finally:
            if out is not None:
                shutil.rmtree(out.parent, ignore_errors=True)

    def test_未选粒度时交付物仍是逐叶子(self):
        import docx
        out = None
        try:
            p = _plan()
            out = Path(D.build_plan_docx(p))
            doc = docx.Document(str(out))
            rows = None
            for t in doc.tables:
                if "任务 ID" in [c.text.strip() for c in t.rows[0].cells]:
                    rows = len(t.rows) - 1
                    break
            assert rows == len(_LEAVES)
            text = "\n".join(x.text for x in doc.paragraphs)
            assert "展示粒度" not in text
        finally:
            if out is not None:
                shutil.rmtree(out.parent, ignore_errors=True)


class TestConfidenceSectionKeepsLeafGranularity:
    """WS5：置信度章节的降级清单必须**永远按叶子任务**说事。

    为什么这条挂在粒度用例里：`_norm_degradations()` 读的是 WBS 叶子的
    `norm_binding`（`all_tasks_schedule` 排程行按契约不带定额锚定）。展示粒度把
    横道并成"工种级 × 整栋"之后，如果降级清单跟着合并，用户就再也看不到
    "哪一条任务的定额单位对不上" —— 而这条信息恰恰是**合并显示的依据**。
    粒度只能改展示，不能改这一章的颗粒度。
    """

    def _plan_with_bindings(self, granularity):
        leaves = []
        for lf, nb_unit in zip(_LEAVES[:2], ("工日/m³", "工日/m³")):
            leaf = dict(lf)
            leaf["norm_binding"] = {"task_id": lf["id"], "mode": "labor",
                                    "unit": nb_unit, "source_code": "GD_2018_A1_1"}
            leaves.append(leaf)
        plan = _plan(granularity)
        base = plan["wbs"]["phases"][0]["work_packages"][0]
        base["sub_packages"] = leaves
        # 只留第一条 phase，让叶子就是这两条
        plan["wbs"]["phases"] = [plan["wbs"]["phases"][0]]
        plan["all_tasks_schedule"] = plan["all_tasks_schedule"][:2]
        plan["meta"]["norm_coverage"] = {"total": 2, "bound": 2, "bound_pct": 100.0,
                                         "unbound": 0, "unbound_pct": 0.0}
        return plan

    def _degradation_rows(self, doc):
        for t in doc.tables:
            if [c.text.strip() for c in t.rows[0].cells] == ["任务 ID", "任务", "原因"]:
                return [[c.text.strip() for c in r.cells] for r in t.rows[1:]]
        return None

    def test_粗粒度下降级清单仍逐叶子(self):
        import docx
        out = None
        try:
            out = Path(D.build_plan_docx(self._plan_with_bindings(dict(COARSE_WHOLE))))
            doc = docx.Document(str(out))
            rows = self._degradation_rows(doc)
            assert rows is not None, "降级清单表格不见了"
            ids = [r[0] for r in rows]
            assert ids == ["5.1.1.1", "5.1.1.2"], ids     # 逐叶子，未合并
            # 两条叶子单位都是 t / m²（若跟着粒度合并会变成一条）
            assert all("不一致" in r[2] for r in rows), rows
        finally:
            if out is not None:
                shutil.rmtree(out.parent, ignore_errors=True)

    def test_默认粒度下降级清单也在(self):
        import docx
        out = None
        try:
            out = Path(D.build_plan_docx(self._plan_with_bindings(None)))
            doc = docx.Document(str(out))
            rows = self._degradation_rows(doc)
            assert rows is not None and len(rows) == 2, rows
        finally:
            if out is not None:
                shutil.rmtree(out.parent, ignore_errors=True)


# ---------------- ④ ECharts 交互甘特也必须跟随粒度 ----------------
class TestEchartsGanttHonorsGranularity:
    """审计缺口：同一张看板上 WBS 表写 164 行、ECharts 甘特画 307 条。

    历史实现（`echarts_page` 模块 docstring）以"要真实日历轴"为由让甘特永远读
    `all_tasks_schedule`，于是用户**看不出自己的粒度选择生效了没有**。
    现在改走 `delivery.rolled_gantt_rows`：组行同样带真实日历
    （组内最早开始 → 最晚完成），口径与 Word / WBS 表同源。
    """

    def test_上卷后甘特行数与WBS表一致(self):
        from pipeline.nodes import echarts_page as E
        p = _plan(dict(COARSE_WHOLE))
        opt = E.build_chart_options(p, D._compute_view(p))["gantt"]
        assert opt["__totalTasks"] == 3, "甘特应合并成 3 行"
        assert opt["__rolledUp"] is True
        assert opt["__leafTasks"] == len(_LEAVES)
        assert len(opt["yAxis"]["data"]) == 3
        json.dumps(opt, ensure_ascii=False)          # 仍然必须是纯 JSON

    def test_组行日期取组内最早开始到最晚完成(self):
        p = _plan(dict(COARSE_WHOLE))
        rows = D.rolled_gantt_rows(p)
        assert rows and all(r["start"] <= r["finish"] for r in rows)
        assert all(r["duration"] >= 1 for r in rows)
        sched = {t["task_id"]: t for t in p["all_tasks_schedule"]}
        by_id = {r["id"]: r for r in rows}
        for g in D.rolled_rows(p):
            ids = [str(i) for i in g["ids"]]
            row = by_id[D._group_id(ids)]
            assert row["start"] == min(sched[i]["start_date"] for i in ids)
            assert row["finish"] == max(sched[i]["finish_date"] for i in ids)

    def test_未上卷时甘特逐叶子且无新标记键(self):
        from pipeline.nodes import echarts_page as E
        p = _plan()
        opt = E.build_chart_options(p, D._compute_view(p))["gantt"]
        assert opt["__totalTasks"] == len(_LEAVES)
        assert "__rolledUp" not in opt and "__leafTasks" not in opt
        assert [d[4] for d in opt["series"][0]["data"]] == \
            [t["task_id"] for t in p["all_tasks_schedule"]]

    def test_未选粒度时不上卷(self):
        assert D.rolled_gantt_rows(_plan()) == []

    def test_看板HTML里甘特与WBS表行数一致(self):
        out = None
        try:
            p = _plan(dict(COARSE_WHOLE))
            out = Path(D.build_plan_html(p))
            html = out.read_text(encoding="utf-8")
            assert '"__totalTasks": 3' in html, "内联 option 仍是逐叶子"
            assert '"__rolledUp": true' in html
            assert "按展示粒度合并，逐叶子 5 项" in html
        finally:
            if out is not None:
                shutil.rmtree(out.parent, ignore_errors=True)


# ---------------- ⑤ 监督报告必须自证"表格是多少行" ----------------
class TestReportCaliber:
    """审计缺口：交付物合并成 164 行，`plan["report"]` 4928 字一个字没提粒度。

    修法：`plan_assembler.display_caliber` 给出唯一口径句 → 进 `build_parts`
    （LLM 看得到数字）→ `template_report` 写上 → `reporter` 对 LLM 产出**强制**
    钉一条（不靠模型自觉）。
    """

    def _ctx(self, granularity):
        p = _plan(granularity)
        return {"wbs": p["wbs"], "cpm_result": p["cpm_result"],
                "resource_demand": p["resource_demand"],
                "extracted_params": {"project_name": "测试",
                                     "planned_start_date": START},
                "display_granularity": (p.get("meta") or {}).get(
                    "display_granularity") or {}}

    def test_口径块进了parts(self):
        from pipeline.nodes.plan_assembler import build_parts
        cal = build_parts(self._ctx(dict(COARSE_WHOLE)))["display_granularity"]
        assert cal["merged"] is True and cal["rows"] == 3
        assert cal["leaves"] == len(_LEAVES)
        assert "工种级（粗） × 整栋" in cal["note"]
        assert "3 行" in cal["note"]

    def test_模板报告带口径句(self):
        from pipeline.nodes.plan_assembler import build_parts, template_report
        rep = template_report(build_parts(self._ctx(dict(COARSE_WHOLE))))
        assert "展示口径" in rep
        assert "工种级（粗） × 整栋" in rep

    def test_未选粒度时报告也说明逐任务行数(self):
        from pipeline.nodes.plan_assembler import build_parts, template_report
        parts = build_parts(self._ctx(None))
        assert parts["display_granularity"]["merged"] is False
        rep = template_report(parts)
        assert "展示口径" in rep and "逐任务 5 行" in rep

    def test_LLM漏写时由节点强制钉上(self):
        from pipeline.nodes.reporter import ReporterNode
        from pipeline.nodes.plan_assembler import build_parts, with_caliber

        class _FakeLLM:
            def chat_text(self, *a, **kw):
                return "# 监督报告\n\n## 一、总体情况\n- 总工期：15 天\n"

        ctx = self._ctx(dict(COARSE_WHOLE))
        ctx["plan_parts"] = build_parts(ctx)
        rep = ReporterNode(llm=_FakeLLM()).run(ctx)["report"]
        assert "展示口径" in rep
        assert "工种级（粗） × 整栋" in rep
        # 幂等：再插一次不会出现第二条
        assert with_caliber(rep, "X").count("展示口径") == 1

    def test_旧计划重导出时交付侧补口径(self):
        """`plan["report"]` 是上一次运行留下的（没有口径句），交付时也要补上。

        否则"打开历史计划 / 重导出"这条路径永远缺凭证 —— 而这正是用户核对
        "选择有没有体现在成果里"时最常走的路。
        """
        from pipeline.nodes.plan_assembler import with_caliber
        p = _plan(dict(COARSE_WHOLE))
        p["report"] = "# 监督报告\n\n## 一、总体情况\n- 老报告，只字未提粒度\n"
        text = D._report_text(p)
        assert "展示口径" in text and "工种级（粗） × 整栋" in text
        # 已经写过的报告不重复插
        assert D._report_text(dict(p, report=text)).count("展示口径") == 1
        assert with_caliber(text, "X").count("展示口径") == 1
        assert with_caliber(text, "X").strip() == text.strip()

    def test_没有标题时也不丢口径(self):
        from pipeline.nodes.plan_assembler import with_caliber
        out = with_caliber("没有标题的一段话", "口径句")
        assert out.startswith("- **展示口径**：口径句")


if __name__ == "__main__":
    for cls in (TestDefaultUnchanged, TestRolledUp, TestCurvesUntouched,
                TestDeliverablesHonorGranularity,
                TestConfidenceSectionKeepsLeafGranularity,
                TestEchartsGanttHonorsGranularity, TestReportCaliber):
        for name in sorted(dir(cls)):
            if name.startswith("test_"):
                getattr(cls(), name)()
                print("ok  %s.%s" % (cls.__name__, name))
    print("全部通过")
