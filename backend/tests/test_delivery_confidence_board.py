# -*- coding: utf-8 -*-
"""「数据来源与置信度」章节进**看板**的回归护栏。

背景 —— 实测：Word 交付物里有「数据来源与置信度」整章（`delivery.add_confidence_section`
消费 `confidence_section_blocks`），而同一份计划的看板里「数据来源与置信度」0 次、
「定额覆盖率」0 次。用户已拍板：看板必须补上这一段，并且：

① **单一真源**：内容只许来自同一个 `confidence_section_blocks(plan, view)`，
   绝不许另写一份内容逻辑、绝不许新造数字 —— 本文件因此逐块对照函数的返回值；
② **关键数字一眼可见**：总数 / 已绑定 / 未绑定 / 其中未审定放行条数留在可见处；
③ **长清单折进 `<details>`**：逐条未绑定原因、来源代码清单等展开可见；
④ **没数据就一个字不加**：`has_confidence_meta` 为假时看板里连标题都不出现（与 Word 同口径）；
⑤ **LLM 编排页也要有结构性兜底**：`_ensure_confidence_section`（照 `_ensure_org_section`
   的现成模式，判据常量与 `DELIVERY_MARKERS` 互相独立），追加后**留痕**。

运行：python -m pytest tests/test_delivery_confidence_board.py -q
"""

import json
import re
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from pipeline import config                          # noqa: E402
from pipeline.nodes import delivery as D             # noqa: E402

_REAL_CANDIDATES = (BACKEND / "plans" / "plan_sample3_after_org.json",
                    BACKEND / "plans" / "plan_sample3_after_fix.json")
REAL_PLAN = next((p for p in _REAL_CANDIDATES if p.exists()), None)


# ══════════════════════ 公共样例 ══════════════════════

def _plan(meta=None):
    """一份最小计划（两条 5 天任务，日期重叠 → 曲线峰值 8 人）。"""
    def _tasks():
        return [
            {"task_id": "1.1.1", "task_name": "柱浇筑", "start_date": "2026-09-01",
             "finish_date": "2026-09-06", "duration_days": 5,
             "assigned_resources": {"钢筋工": 4}},
            {"task_id": "1.1.2", "task_name": "梁浇筑", "start_date": "2026-09-01",
             "finish_date": "2026-09-06", "duration_days": 5,
             "assigned_resources": {"木工": 4}},
        ]

    plan = {
        "plan_id": "plan_confidence_board_test",
        "overview": {"project_name": "置信度看板测试", "total_duration_days": 10,
                     "planned_start_date": "2026-09-01",
                     "planned_end_date": "2026-09-11", "critical_path_length": 2},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "1.1", "name": "结构", "sub_packages": [
                {"id": "1.1.1", "name": "柱浇筑", "duration_days": 5, "quantity": 10,
                 "unit": "m³", "work_type": "混凝土工程"},
                {"id": "1.1.2", "name": "梁浇筑", "duration_days": 5, "quantity": 12,
                 "unit": "m³", "work_type": "混凝土工程"}]}]}]},
        "dependencies": [],
        "cpm_result": {"total_duration_days": 10, "critical_path": ["1.1.1"],
                       "schedule": [{"task_id": "1.1.1", "es": 0, "ef": 5,
                                     "ls": 0, "lf": 5}]},
        "resource_demand": {"tasks": _tasks()},
        "key_milestones": [{"name": "开工", "date": "2026-09-01",
                            "task_id": "1.1.1", "description": "开工"}],
        "critical_path_tasks": _tasks(),
        "all_tasks_schedule": _tasks(),
        "resource_plan": {"total_manpower_days": 40, "peak_manpower": 8,
                          "equipment_peak": {"混凝土输送泵车": 1},
                          "material_summary": []},
        "risks": [],
        "report": "# 监督报告",
    }
    if meta is not None:
        plan["meta"] = meta
    return plan


# 六个置信度键全给：覆盖率数字、放行条数、来源代码、两版工期、用户目标都在。
_CONF_META = {
    "norm_coverage": {"total": 333, "bound": 176, "bound_pct": 52.9,
                      "unbound": 157, "unbound_pct": 47.1,
                      "released_unapproved": 85, "released_unapproved_pct": 25.5,
                      "by_confidence": {"parsed": 60, "verified": 25},
                      "by_reason": {"AI估算定额": 120, "单位不一致": 30,
                                    "定额口径不符": 5, "工程量量级不可信": 2},
                      "by_reason_pct": {"AI估算定额": 36.0, "单位不一致": 9.0,
                                        "定额口径不符": 1.5, "工程量量级不可信": 0.6}},
    "credibility": {"user": 0.0, "kb": 0.29, "ai": 0.71},
    "data_sources": ["AI_ESTIMATE_V1", "GD_2018_A1_3", "LD_T72_2_2008"],
    "schedule_versions": {"theory_min_days": 967, "resource_ok_days": 1012,
                          "delta_days": 45},
    "boundary_conditions": {"project_duration_days": 900},
}

# 六项全空的 meta（只有审计状态，与 `test_delivery.py` 的 MINI 同形状）。
_NO_CONF_META = {"audit_status": "未审计"}

_DETAILS_RE = re.compile(r"<details\b.*?</details>", re.S)


def _visible(html_text):
    """去掉 `<details>` 展开区之后剩下的、用户不点开就能看到的正文。"""
    return _DETAILS_RE.sub("", html_text)


def _board(plan):
    return Path(D.build_plan_html(plan)).read_text(encoding="utf-8")


def _card(plan):
    return D._confidence_section_html(plan, D._compute_view(plan))


def _g5_normalized(text):
    """按交付物侧同一条 G5 归一收口处理文本（`build_plan_html` 落盘前就是这么做的）。

    历史计划里的 `_unit_assumed` 人话串还写着 `㎡`(U+33A1)，交付物会把它归一到 `m²`；
    测试要拿"同一份成品"比，就得走同一道归口。
    """
    from pipeline.nodes import plan_assembler as _PA
    return _PA.normalize_cjk_compat_square_metre_in_text(text)[0]


def _word_all_text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parts.extend(c.text for c in r.cells)
    return "\n".join(parts)


class _FakeDoc:
    """`add_confidence_section` 只用到 `doc.add_paragraph`（本节没有表格/图片）。"""

    def __init__(self):
        self.paras = []

    def add_paragraph(self, text):
        self.paras.append(str(text))


class _FakeLLM:
    def __init__(self, text):
        self.text = text
        self.calls = 0

    def chat_text(self, *a, **kw):
        self.calls += 1
        return self.text


BARE_LLM_HTML = ("<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'>"
                 "<title>模型编排</title></head><body><h1>施工进度计划看板</h1>"
                 "<p>模型只写了个壳</p></body></html>")


@pytest.fixture()
def tmp_deliverables(tmp_path, monkeypatch):
    """把交付物目录指向 tmp，避免测试往 `输出结果/` 里写东西。"""
    d = tmp_path / "deliverables"
    d.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(config, "DELIVERABLES_DIR", d)
    monkeypatch.setattr(D.config, "DELIVERABLES_DIR", d)
    return d


# ══════════════════════ ① 确定性看板：卡片真的在 ══════════════════════

class TestDeterministicBoard:
    def test_看板含数据来源与置信度和定额覆盖率(self, tmp_deliverables):
        html = _board(_plan(dict(_CONF_META)))
        assert "数据来源与置信度" in html
        assert "定额覆盖率" in html

    def test_卡片紧接资源卡之后且在关键路径明细之前(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        html = _board(plan)
        resource_card = D._resource_card_html(plan, D._compute_view(plan),
                                             D._peak_caliber(plan, D._compute_view(plan)))
        card = _card(plan)
        assert card and card in html
        assert html.index(resource_card) < html.index(card)
        assert html.index(card) < html.index("关键路径明细")

    def test_覆盖率关键数字在可见处(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        html = _board(plan)
        card = _card(plan)
        assert card in html
        visible = _visible(card)
        for txt in ("定额口径任务总数", "已绑定定额", "未绑定定额",
                    "其中：未经人工审定的真人定额（已放行）"):
            assert txt in visible, txt
        for txt in ("333", "176（52.9%）", "157（47.1%）", "85（25.5%）"):
            assert txt in visible, txt

    def test_卡片逐块照抄confidence_section_blocks(self, tmp_deliverables):
        """单一真源：卡片里必须出现函数返回的每一个 kv 值与 grid 单元格。"""
        plan = _plan(dict(_CONF_META))
        view = D._compute_view(plan)
        blocks = D.confidence_section_blocks(plan, view)
        assert blocks, "样例计划必须能算出内容块"
        card = _card(plan)
        assert card
        kinds = [k for k, _ in blocks]
        assert "kv" in kinds and "grid" in kinds
        import html as _h
        for kind, payload in blocks:
            if kind == "h3":
                assert _h.escape(str(payload)) in card, payload
            elif kind == "kv":
                for label, value in payload:
                    assert _h.escape(str(label)) in card, label
                    assert _h.escape(str(value)) in card, value
            elif kind == "grid":
                headers, rows = payload[0], payload[1]
                for h in headers:
                    assert _h.escape(str(h)) in card, h
                for row in rows:
                    for cell in row:
                        assert _h.escape(str(cell)) in card, cell
        # 源文案里的 **强调** 必须变成 <b>，不许把 markdown 泄进看板
        assert "**" not in card
        assert "<b>按来源放行</b>" in card


# ══════════════════════ ② 长清单折进 <details> ══════════════════════

class TestLongListsFoldIntoDetails:
    def test_逐条未绑定原因在details内(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        card = _card(plan)
        by_reason = plan["meta"]["norm_coverage"]["by_reason"]
        first_reason = next(iter(by_reason))
        assert "<details" in card
        assert card.index("<details") < card.index(first_reason), \
            "逐条未绑定原因必须在 <details> 之后（折进可展开区）"
        assert first_reason not in _visible(card)
        # summary 文案 = 真实行数
        assert ("<summary>共 %d 条，展开查看</summary>" % len(by_reason)) in card

    def test_来源代码清单与逐条清单都在details内(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        card = _card(plan)
        assert card.count("<details") >= 2
        for code in plan["meta"]["data_sources"]:
            assert code in card
            assert code not in _visible(card), code
        assert "共 %d 条，展开查看" % len(plan["meta"]["data_sources"]) in card


# ══════════════════════ ③ 没数据就一个字不加 ══════════════════════

class TestNoMetaNoCard:
    def test_六个置信度键全空时看板一个字不加(self, tmp_deliverables):
        plan = _plan(dict(_NO_CONF_META))
        assert D.has_confidence_meta(plan) is False
        assert _card(plan) == ""
        html = _board(plan)
        assert "数据来源与置信度" not in html
        assert "定额覆盖率" not in html
        assert D.CONFIDENCE_FALLBACK_COMMENT not in html

    def test_单靠峰值一行也不出空壳(self, tmp_deliverables):
        """`resource_plan.peak_manpower` 几乎每份计划都有 —— 不许因此多出一章空壳。"""
        plan = _plan(dict(_NO_CONF_META))
        plan["resource_plan"]["declared_peak_manpower"] = 120
        plan["resource_plan"]["declared_peak_manpower_source"] = "model"
        assert _card(plan) == ""
        assert "数据来源与置信度" not in _board(plan)


# ══════════════════════ ④ LLM 编排页的结构性兜底 ══════════════════════

class TestAgentFallback:
    def test_缺标记时追加确定性卡片并留痕(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        ctx = {}
        path, used_agent = D.build_plan_html_agent(plan, _FakeLLM(BARE_LLM_HTML), ctx=ctx)
        assert used_agent is True, "追加段走的是 LLM 编排页面这一路"
        html = Path(path).read_text(encoding="utf-8")
        assert "数据来源与置信度" in html and "定额覆盖率" in html
        assert D.CONFIDENCE_FALLBACK_COMMENT in html
        # 追加段必须在 </body> 之前（页面仍合法）
        assert html.rfind(D.CONFIDENCE_FALLBACK_COMMENT) < html.rfind("</body>")
        # 追加的就是确定性卡片本身（同一函数，一字不差）
        assert _card(plan) in html
        # 模型原有内容不许被丢掉
        assert "模型只写了个壳" in html
        joined = " ".join(ctx.get("wbs_warnings") or [])
        assert "已追加确定性置信度段落" in joined, ctx.get("wbs_warnings")

    def test_已含标记时原样不动(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        view = D._compute_view(plan)
        card = _card(plan)
        rich = ("<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'></head><body>"
                + card + "</body></html>")
        merged, added = D._ensure_confidence_section(rich, plan, view)
        assert added is False
        assert merged == rich, "标记齐备就不该动模型页面"
        assert merged.count("数据来源与置信度") == 1

    def test_agent路不重复追加(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        card = _card(plan)
        page = ("<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'></head><body>"
                "<h1>模型编排</h1>依据 工作面容量 主要机械峰值 施工组织口径 组织缺口"
                + card + "</body></html>")
        path, used_agent = D.build_plan_html_agent(plan, _FakeLLM(page), ctx={})
        assert used_agent is True
        html = Path(path).read_text(encoding="utf-8")
        assert html.count("数据来源与置信度") == 1
        assert D.CONFIDENCE_FALLBACK_COMMENT not in html

    def test_没有置信度元数据时不追加也不留痕(self, tmp_deliverables):
        plan = _plan(dict(_NO_CONF_META))
        ctx = {}
        path, used_agent = D.build_plan_html_agent(plan, _FakeLLM(BARE_LLM_HTML), ctx=ctx)
        assert used_agent is True
        html = Path(path).read_text(encoding="utf-8")
        assert "数据来源与置信度" not in html
        assert D.CONFIDENCE_FALLBACK_COMMENT not in html
        joined = " ".join(ctx.get("wbs_warnings") or [])
        assert "置信度" not in joined, joined

    def test_判据常量与DELIVERY_MARKERS互相独立(self):
        assert D.CONFIDENCE_TITLE not in D.DELIVERY_MARKERS
        assert D.CONFIDENCE_MARKERS[0] not in D.DELIVERY_MARKERS
        assert set(D.CONFIDENCE_MARKERS).isdisjoint(set(D.ORG_MARKERS))
        assert D.DELIVERY_MARKERS == ("依据", "工作面容量", "主要机械峰值")

    def test_覆盖率标记只在有计划数据时才当判据(self):
        """只有来源构成 / 两版工期的计划（没有 norm_coverage）也补得上这一段。"""
        meta = {k: v for k, v in _CONF_META.items() if k != "norm_coverage"}
        plan = _plan(meta)
        assert D.has_confidence_meta(plan) is True
        assert "定额覆盖率" not in _card(plan)
        merged, added = D._ensure_confidence_section(BARE_LLM_HTML, plan, D._compute_view(plan))
        assert added is True
        assert "数据来源与置信度" in merged
        assert D.CONFIDENCE_FALLBACK_COMMENT in merged


# ══════════════════════ ⑤ Word 那条路径不受影响 ══════════════════════

class TestWordPathUnchanged:
    def test_Word章节仍是二级标题且内容齐(self, tmp_deliverables):
        plan = _plan(dict(_CONF_META))
        path = D.build_plan_docx(plan)
        text = _word_all_text(path)
        assert "数据来源与置信度" in text
        from docx import Document
        titles = [(p.style.name, p.text) for p in Document(path).paragraphs
                  if (p.style.name or "").startswith("Heading")]
        assert ("Heading 2", "数据来源与置信度") in titles, titles
        for sub in ("1. 定额覆盖率", "2. 来源构成", "3. 两版工期与用户目标"):
            assert sub in text, sub
        assert "333" in text and "176" in text and "157" in text and "85" in text
        assert "AI估算定额" in text

    def test_Word没有置信度元数据时不出(self, tmp_deliverables):
        plan = _plan(dict(_NO_CONF_META))
        assert "数据来源与置信度" not in _word_all_text(D.build_plan_docx(plan))

    def test_add_confidence_section判据未变(self):
        view = D._compute_view(_plan(dict(_CONF_META)))
        # 六项全空 → 一个字都不写，返回 False
        calls = []
        empty_doc = _FakeDoc()
        assert D.add_confidence_section(
            empty_doc, _plan(dict(_NO_CONF_META)), view,
            lambda *a, **k: calls.append(("h", a, k)),
            lambda *a: calls.append(("kv", a)),
            lambda *a: calls.append(("grid", a))) is False
        assert calls == [] and empty_doc.paras == []

        # 有数据 → 先二级标题，再逐块（h3 → kv → para → grid）
        rich_doc = _FakeDoc()
        calls = []
        ok = D.add_confidence_section(
            rich_doc, _plan(dict(_CONF_META)), view,
            lambda *a, **k: calls.append(("h", a, k)),
            lambda *a: calls.append(("kv", a)),
            lambda *a: calls.append(("grid", a)))
        assert ok is True
        assert calls[0] == ("h", ("数据来源与置信度",), {"level": 2})
        assert calls[1][0] == "h" and calls[1][1] == ("1. 定额覆盖率",)
        assert calls[1][2] == {"level": 3}
        assert any(c[0] == "kv" for c in calls)
        assert any(c[0] == "grid" for c in calls)


# ══════════════════════ ⑥ 真实计划渲染 ══════════════════════

@pytest.mark.skipif(REAL_PLAN is None, reason="本机没有真实计划样本")
class TestRealPlanRender:
    def test_真实计划看板有置信度卡片且数字可见(self, tmp_deliverables):
        plan = json.loads(REAL_PLAN.read_text(encoding="utf-8"))
        cov = plan["meta"]["norm_coverage"]
        card = _g5_normalized(_card(plan))
        assert card
        html = _board(plan)
        assert card in html
        assert "数据来源与置信度" in html and "定额覆盖率" in html
        visible = _visible(card)
        for txt in (str(cov["total"]), str(cov["bound"]), str(cov["unbound"])):
            assert txt in visible, txt
        assert D.CONFIDENCE_CARD_LEAD in card
        assert re.search(r"<details><summary>共 \d+ 条，展开查看</summary>", card), card[:2000]
        # 逐条未绑定原因只出现在展开区里
        by_reason = cov.get("by_reason") or {}
        if by_reason:
            first_reason = next(iter(by_reason))
            assert first_reason in card
            assert first_reason not in visible


# ══════════════════════ ⑦ AI 经验估算定额条数（政策变更 2026-09-20） ══════════════════════
# 政策变更（用户 2026-09-20）：AI 经验估算定额（KB `AI_ESTIMATE_V1`）由「只作参考」
# 改为**照用**（与真人定额同等参与工期与班组计算）。交付物这一层唯一不能省的就是
# **如实给出条数**并把这一行留在看板可见处（不能折进 <details>）。
# 条数来源：优先 `meta.norm_coverage.released_ai`（上游新口径）；上游还没写这个键时
# 由交付侧从任务级数据自己数；数不出来 → 整行不出，**绝不用 0 充数**。

def test_AI条数行在可见处且随数据变化(tmp_deliverables):
    def _ai_rd(tid):
        return {"task_id": tid, "task_name": "AI 定额任务 " + tid, "quantity": 10,
                "planned_duration_days": 5,
                "resources": {"瓦工": {"per_day": 4, "total_days": 20}},
                "_resource_source": {"瓦工": {"origin": "ai_estimate",
                                              "ref": "AI_ESTIMATE_V1"}},
                "_norm_applied": {"mode": "labor", "norm_value": 0.1,
                                  "source_code": "AI_ESTIMATE_V1",
                                  "match_type": "ai", "origin": "ai"}}

    for n in (1, 3):
        plan = _plan(dict(_CONF_META))
        plan["resource_demand"] = {"tasks": [_ai_rd("1.1.%d" % i) for i in range(1, n + 1)]}
        card = _card(plan)
        visible = _visible(card)
        assert D.AI_NORM_ROW_LABEL_LEGACY in visible, visible[:1200]
        assert "%d 条（本次运行已据此算出班组与工日）" % n in visible, visible[:1200]
        # 上游没写 released_ai → 不许冒用"已按来源放行"的政策标签
        assert D.AI_NORM_ROW_LABEL not in visible

    # 上游给了 `released_ai` → 照抄条数并用政策标签
    plan = _plan(dict(_CONF_META, norm_coverage=dict(
        _CONF_META["norm_coverage"], released_ai=9, released_ai_pct=2.7)))
    visible = _visible(_card(plan))
    assert D.AI_NORM_ROW_LABEL in visible
    assert "9（2.7%）" in visible


def test_没有AI数据时AI条数行不出也不写0(tmp_deliverables):
    """数据干净（计划里没有一条 AI 经验估算定额）→ 整行不出，绝不写 0 条占位。"""
    plan = _plan(dict(_CONF_META))
    visible = _visible(_card(plan))
    assert D.AI_NORM_ROW_LABEL not in visible
    assert D.AI_NORM_ROW_LABEL_LEGACY not in visible
    assert "AI 经验估算定额" not in visible, visible[:1200]
