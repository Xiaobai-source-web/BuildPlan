# -*- coding: utf-8 -*-
"""第 2 批 · 域 2（代码侧）验收护栏：参数键增删 + 交付物声明 + 装配式中断。

覆盖本批的硬要求（每条都有断言，报告里的"已删除/已存在"结论都从这里可复现）：
  2.1 `foundation_type`（基础类型）登记到全部落点；**缺了即中断**（连试算也不放行）；
  2.2 `total_infill_wall`（m³）/ `total_pile`（桩，**不预设单位**）同样登记；
  2.3/2.4 `total_wall` / `total_precast` 全链路 0 引用（`backend/prompts/` 除外）；
  2.6 `materials`（材料清单）不再要求/接受/展示，但 `material_transport`（材料运输**工序**）
      与 `_materialize_unit_assumption`（"落实假设值"）**必须还在**；
  2.7 交付物声明在**看板 HTML 与 Word 两种产物**里都出现；
  2.9 装配式建筑 → 报错返回、不出计划；而"管桩/预制桩"**不是**装配式标志。

运行：python -m pytest backend/tests/test_batch2_domain2_param_keys.py -q
"""

import re
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))
# 终端侧的表在 `terminal/renderer.py`（终端进程不 import 后端 → 两张表要分别检查）
for _p in (BACKEND, BACKEND.parent / "terminal"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import pytest  # noqa: E402

from pipeline.nodes import boundary as B  # noqa: E402
from pipeline.nodes import delivery as D  # noqa: E402
from pipeline.nodes import plan_assembler as PA  # noqa: E402
from pipeline import ratio_scope as RS  # noqa: E402
from pipeline.nodes.extractor import (  # noqa: E402
    ExtractorNode, _NUM_KEYS, detect_prefab_system, match_foundation_type)

REPO_ROOT = BACKEND.parent
SOURCE_SCAN_DIRS = (BACKEND / "pipeline", BACKEND.parent / "terminal")


def _iter_sources():
    for base in SOURCE_SCAN_DIRS:
        if not base.exists():
            continue
        for p in base.rglob("*.py"):
            if "__pycache__" in p.parts:
                continue
            yield p


# ══════════════════════════════════════════════════════════════════
# 2.1 / 2.2 登记落点
# ══════════════════════════════════════════════════════════════════
NEW_KEYS = ("foundation_type", "total_infill_wall", "total_pile")
REMOVED_KEYS = ("total_wall", "total_precast")


class TestRegistration:
    def test_新增键进核心参数与回退档(self):
        for k in NEW_KEYS:
            assert k in B.CORE_KEYS, ("CORE_KEYS 缺 %s" % k)
        assert "foundation_type" in B.REQUIRED_KEYS, B.REQUIRED_KEYS
        for k in ("total_infill_wall", "total_pile"):
            assert k in B.FALLBACK_KEYS, ("FALLBACK_KEYS 缺 %s" % k)

    def test_新增键都有中文名且_total_pile不带单位(self):
        for k in NEW_KEYS:
            assert B.param_label(k) != k, ("没有中文名：%s" % k)
        assert B.param_label("foundation_type") == "基础类型"
        assert B.param_label("total_infill_wall") == "填充墙(m³)"
        # **不预设单位**：标签里不许出现 m / m³ / t / 吨 / 根 / 米
        label = B.param_label("total_pile")
        assert label == "桩", label
        for unit_word in ("m", "m³", "t", "吨", "根", "米"):
            assert unit_word not in label, (unit_word, label)

    def test_终端参数名表与后端同步(self):
        import renderer
        for k in NEW_KEYS:
            assert k in renderer._PARAM_LABELS, ("renderer 缺 %s" % k)
            assert renderer._PARAM_LABELS[k] == B.PARAM_LABELS[k], k
        # total_pile 在终端侧同样不带单位
        assert renderer._PARAM_LABELS["total_pile"] == "桩"
        assert "total_pile" in renderer._OPTIONAL_PARAMS
        assert "total_infill_wall" in renderer._OPTIONAL_PARAMS
        # 装基础类型是硬必要键，不属于"不填也能编"
        assert "foundation_type" not in renderer._OPTIONAL_PARAMS

    def test_extractor数值键表同步(self):
        assert "total_infill_wall" in _NUM_KEYS and "total_pile" in _NUM_KEYS
        for k in REMOVED_KEYS:
            assert k not in _NUM_KEYS, ("_NUM_KEYS 仍含已删除键 %s" % k)
        # foundation_type 是**文本键**，绝不许进数值归一表（会把「筏板基础」毁掉）
        assert "foundation_type" not in _NUM_KEYS

    def test_ratio_scope那处已按域6收口(self):
        """【域 6 · 6.4】`pile_foundation` 已是**占比表外的直接量**，不再是拆分工种。

        第 2 批曾把它临时接续到 `total_pile`（占位）；域 6 的裁决是
        「**桩基不是独立项，它就是「基础」栏**」⇒ 该键**必须**移出
        `GROUP_TOTAL_PARAMS`（占比表里 `pile_foundation` 实测 0 行，
        留着只会让桩基 33 个 L4 全部被判「异常缺行」→ 量 0 出局）。
        桩量由 `ratio_scope._bind_foundation_column` 从「基础」栏改投得到。
        """
        assert "pile_foundation" not in RS.GROUP_TOTAL_PARAMS, RS.GROUP_TOTAL_PARAMS
        assert "total_wall" not in RS.GROUP_TOTAL_PARAMS.values()
        # `total_pile` 仍是「桩」的直接量来源，只是换了消费点（不在占比表对账里）
        assert RS.PILE_TOTAL_PARAM == "total_pile"
        src = (BACKEND / "pipeline" / "ratio_scope.py").read_text(encoding="utf-8")
        assert "域 6" in src, "该处必须留一句域 6 的收口说明"

    def test_两处L3映射口径一致(self):
        """`kb_scope._quantity_strengthened_l3` 仍是「用户给量 → L3 强化为必须」的真源。

        ⚠️ 与 `ratio_scope.GROUP_TOTAL_PARAMS` 的分工**在域 6 变了**：
          · `kb_scope` 那张表管的是「**哪个 L3 被强化为 REQUIRED**」——桩基当然要
            （用户给了桩量，桩基工种必须进树），所以 `("total_pile", "pile_foundation")`
            **必须留着**；
          · `ratio_scope.GROUP_TOTAL_PARAMS` 管的是「**哪个工种按占比表拆量**」——
            桩基不在占比表里（0 行），所以**必须移出**。
        两处不再是"逐字一致"，而是"各管一段"，这条断言把分工钉住。
        """
        from pipeline.nodes import kb_scope as KS
        src = (BACKEND / "pipeline" / "nodes" / "kb_scope.py").read_text(encoding="utf-8")
        assert '("total_pile", "pile_foundation")' in src, src[:0]
        # 反向钉住：kb_scope 侧**必须**保留桩基强化（否则桩基工种不进树，
        # 改投过去的量就没有落地的工序了）。
        strengthened = KS._quantity_strengthened_l3({"total_pile": 8000})
        assert strengthened.get("pile_foundation") == ("total_pile", 8000), strengthened
        # 而 ratio_scope 侧**必须**不含桩基（占比表里没有它的分组）
        assert "pile_foundation" not in RS.GROUP_TOTAL_PARAMS

    def test_earthwork项不许顺手修正(self):
        """已知现象：`GROUP_TOTAL_PARAMS` 里 earthwork 项**不在** `Component_Ratio` 表内。

        本批**不许**顺手"修正"它（父代理点名要求）。
        """
        assert RS.GROUP_TOTAL_PARAMS["earthwork"] == "total_earthwork"


# ══════════════════════════════════════════════════════════════════
# 2.1 「提取不到基础类型 → 报错返回、不出计划」
# ══════════════════════════════════════════════════════════════════
class TestFoundationTypeHardGate:
    def test_缺基础类型判为绝对必要(self):
        # 【第 2 批收口】基础类型与结构形式是**同档的两条独立硬必要**。都不给时两条都该
        # 出现在 missing_required / missing_absolute 里；单看一条的断言在下面两个用例。
        comp = B.params_completeness({"floors": 18, "total_area": 14200})
        assert comp["ok"] is False
        assert comp["missing_required"] == ["foundation_type", "structure_type"], \
            comp["missing_required"]
        assert comp["missing_absolute"] == ["foundation_type", "structure_type"], \
            comp["missing_absolute"]

    def test_给了结构形式就只剩基础类型缺失(self):
        """【第 2 批收口 · 用户裁决】「结构各类型和基础类型都是，如果没有输入，那就报错。"""
        comp = B.params_completeness({"floors": 18, "total_area": 14200,
                                      "structure_type": "框架-剪力墙结构"})
        assert comp["ok"] is False
        assert comp["missing_required"] == ["foundation_type"], comp["missing_required"]
        assert comp["missing_absolute"] == ["foundation_type"], comp["missing_absolute"]

    def test_给了基础类型就只剩结构形式缺失(self):
        comp = B.params_completeness({"floors": 18, "total_area": 14200,
                                      "foundation_type": "筏板基础"})
        assert comp["ok"] is False
        assert comp["missing_required"] == ["structure_type"], comp["missing_required"]
        assert comp["missing_absolute"] == ["structure_type"], comp["missing_absolute"]

    def test_绝对必要键是必要键的子集(self):
        assert set(B.ABSOLUTE_KEYS) <= set(B.REQUIRED_KEYS), B.ABSOLUTE_KEYS

    def test_封闭词表认得标准基础形式(self):
        for text, want in (("基础类型：筏板基础", "筏板基础"),
                           ("采用独立基础", "独立基础"),
                           ("桩承台基础施工", "桩承台基础"),
                           ("桩筏基础", "桩筏基础"),
                           ("地上18层，剪力墙结构", None)):
            assert match_foundation_type(text) == want, (text, want)

    def test_封闭词表长名优先(self):
        # "桩筏基础" 必须赢过 "桩基础"（否则会误判成普通桩基础）
        assert match_foundation_type("桩筏基础") == "桩筏基础"

    def test_extractor把基础类型落进参数(self):
        from pipeline.nodes.extractor import extract_by_regex
        assert extract_by_regex("基础类型：筏板基础")["foundation_type"] == "筏板基础"

    def test_提示词里的基础类型优先于兜底(self):
        from pipeline.nodes.extractor import normalize_params
        merged = normalize_params({"foundation_type": "箱形基础"}, "筏板基础")
        assert merged["foundation_type"] == "箱形基础", merged


# ══════════════════════════════════════════════════════════════════
# 2.9 装配式 → 中断；管桩/预制桩 → 不中断
# ══════════════════════════════════════════════════════════════════
class TestPrefabGate:
    @pytest.mark.parametrize("text", [
        "本项目为装配式建筑，预制率 40%",
        "采用预制装配整体式框架，装配率 50%",
        "PC构件由工厂生产",
        "装配式",
    ])
    def test_装配式体系命中即中断(self, text):
        node = ExtractorNode(llm=None)
        node._emit = lambda e, d: None
        out = node.run({"prompt": text, "intent": "plan"})
        assert "_stop" in out, (text, out)
        assert "装配式" in out["_stop"], out["_stop"]

    @pytest.mark.parametrize("text", [
        "约120根管桩，桩基工程",
        "预制桩施工，基础类型：桩基础",
        "劳动力：装配式安装工 20 人",       # 工种名，不是建筑体系
        "剪力墙结构，筏板基础，地上18层",
    ])
    def test_管桩与预制桩不是装配式标志(self, text):
        assert detect_prefab_system(text) == "", text
        node = ExtractorNode(llm=None)
        node._emit = lambda e, d: None
        out = node.run({"prompt": text, "intent": "plan"})
        assert "_stop" not in out, (text, out)


# ══════════════════════════════════════════════════════════════════
# 2.3 / 2.4 / 2.6 删除：全链路 0 引用（`backend/prompts/` 除外）
# ══════════════════════════════════════════════════════════════════
def _grep_sources(pattern):
    hits = []
    rx = re.compile(pattern)
    for p in _iter_sources():
        for i, line in enumerate(p.read_text(encoding="utf-8").splitlines(), 1):
            if rx.search(line):
                hits.append((str(p.relative_to(REPO_ROOT)).replace("\\", "/"), i, line.strip()))
    return hits


def _code_hits(hits):
    """只保留**可执行代码**行（`#` 注释行不算）——注释里说明"删了什么"是允许的。"""
    return [h for h in hits if not h[2].startswith("#")]


class TestRemovedKeysZeroReference:
    def test_total_wall与total_precast在代码里0引用(self):
        """**可执行代码**里 0 引用。

        注释里保留少量"本批删掉了这个键"的说明（便于后人查证），所以判据是
        "非 `#` 注释行 0 命中"；完整 grep 计数见交付报告。
        `backend/prompts/**` 不归本批代码侧负责（另一个代理独占），不在扫描范围。
        """
        hits = _code_hits(_grep_sources(r"total_wall|total_precast"))
        assert hits == [], "代码里仍有引用：\n" + "\n".join(
            "%s:%d %s" % h for h in hits)

    def test_生产代码里materials只剩删除机制本身(self):
        """材料清单键的**功能**引用只允许出现在白名单处（逐条有理由，见交付报告）。

        白名单：
          · `boundary.py`：`REMOVED_BOUNDARY_KEYS`（删除机制本身）与相关注释；
          · `norm_bind.py`：`_materials_text` —— **换算参数**（墙厚/容重）的既有来源
            之一，不是"材料需求量"清单；删它会改掉面积↔体积的换算回退行为，
            并打破 `test_norm_basis` / `test_unit_area_volume` 的既有护栏。
            新计划已不再产出该键 → 该路径对新计划**惰性**，仅历史计划仍可用；
          · `schemas.py`：`material_summary` 字段（向后兼容，历史计划校验要过）；
          · `delivery.py`：历史计划的 `material_summary` 仍要做 U+33A1 归一（G5 硬断言）。
        """
        allow = {
            "backend/pipeline/nodes/boundary.py",
            "backend/pipeline/nodes/norm_bind.py",
            "backend/pipeline/schemas.py",
            "backend/pipeline/nodes/delivery.py",
        }
        offenders = [h for h in _code_hits(_grep_sources(r"\bmaterials\b|material_summary"))
                     if h[0] not in allow]
        assert offenders == [], "生产代码里仍有未登记的 materials 引用：\n" + "\n".join(
            "%s:%d %s" % h for h in offenders)

    def test_材料清单不再进_source且被源头剔除(self):
        assert "materials" not in B.SOURCE_KEYS, B.SOURCE_KEYS
        assert "materials" in B.REMOVED_BOUNDARY_KEYS, B.REMOVED_BOUNDARY_KEYS


# ══════════════════════════════════════════════════════════════════
# 2.6 绝不能误删的两样
# ══════════════════════════════════════════════════════════════════
class TestMustSurvive:
    def test_material_transport工种仍在(self):
        """`material_transport` 是「材料运输」类**工序**（知识库里 116 个 L4）。"""
        hits = _grep_sources(r"material_transport")
        assert hits, "material_transport 不见了！"
        # 它必须仍被当作 WORK TYPE（而不是被删成普通词）
        for rel, _ln, line in hits:
            if rel.endswith("kb_scope.py") or rel.endswith("wbs_gen.py"):
                return
        # 允许只出现在测试/其它模块，但至少要有一处生产代码
        assert any("pipeline" in h[0] for h in hits), hits

    def test_materialize_unit_assumption仍在(self):
        """`_materialize_unit_assumption` 是"物化/落实假设值"，**与材料无关**。"""
        hits = _grep_sources(r"_materialize_unit_assumption")
        assert hits, "_materialize_unit_assumption 不见了！"
        from pipeline.nodes import resource as R
        assert hasattr(R, "_materialize_unit_assumption"), "resource 里没有这个函数"

    def test_材料运输L4数量仍是116(self):
        """`material_transport` 名下 116 个 L4 不许少（删错键的最终判据）。"""
        from pipeline import kb
        rows = kb._query_all(  # noqa: SLF001
            "SELECT COUNT(*) FROM L4_Activity_Dictionary WHERE work_type_id='material_transport'")
        n = rows[0][0] if rows else 0
        assert n == 116, ("material_transport 名下 L4 数量变了：%s" % n)


# ══════════════════════════════════════════════════════════════════
# 2.7 交付物声明：HTML + Word 两种产物里都要有
# ══════════════════════════════════════════════════════════════════
class TestMaterialsNotice:
    def test_声明原文与本批口径一字不差(self):
        assert D.MATERIALS_EXCLUDED_NOTICE == (
            '本计划不含材料计划。材料按"管够"处理，不参与工期与资源计算。')

    def test_word里有这句声明(self):
        from test_delivery import MINI
        from docx import Document
        path = D.build_plan_docx(MINI)
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.append(c.text)
        joined = "\n".join(parts)
        assert D.MATERIALS_EXCLUDED_NOTICE in joined, "Word 里没有材料计划声明"

    def test_看板html里有这句声明(self):
        from test_delivery import MINI
        html = Path(D.build_plan_html(MINI)).read_text(encoding="utf-8")
        assert D.MATERIALS_EXCLUDED_NOTICE in html, "看板里没有材料计划声明"

    def test_看板不再有主要材料一行(self):
        from test_delivery import MINI
        html = Path(D.build_plan_html(MINI)).read_text(encoding="utf-8")
        assert "主要材料" not in html, "看板仍在展示材料清单"

    def test_llm编排页缺声明时会被追加(self):
        """结构性保证：模型页面没写这句 → `_ensure_materials_notice` 追加确定性一段。"""
        merged, added = D._ensure_materials_notice("<html><body>模型页面</body></html>")
        assert added is True
        assert D.MATERIALS_EXCLUDED_NOTICE in merged, merged
        # 已有则不重复追加
        _again, added2 = D._ensure_materials_notice(merged)
        assert added2 is False


# ══════════════════════════════════════════════════════════════════
# 附加：新键不破坏 plan_json 契约
# ══════════════════════════════════════════════════════════════════
class TestPlanJsonStillValid:
    def test_新键能随extracted_params进产物(self):
        from test_delivery import MINI
        import copy
        plan = copy.deepcopy(MINI)
        plan.setdefault("meta", {})
        plan["meta"]["extracted_params"] = {
            "foundation_type": "筏板基础", "total_infill_wall": 1800,
            "total_pile": 320, "total_masonry": 3000}
        from pipeline import schemas
        out = schemas.PlanJson.model_validate(plan).model_dump()
        got = out["meta"]["extracted_params"]
        assert got["foundation_type"] == "筏板基础"
        assert got["total_pile"] == 320
        assert PA is not None
