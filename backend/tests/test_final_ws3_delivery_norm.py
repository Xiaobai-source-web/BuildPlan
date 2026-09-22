# -*- coding: utf-8 -*-
"""终版修改 · WS3（契约 §9 / §11）验收门：D4 / D6 / D7 + 口径换算留痕。

覆盖合同 §11 给 WS3 的三条最少要求，外加结构性保证（独立判据 + 确定性追加）：
  ① **D4**：`_norm_applied is None` 的工序在看板与 Word 上都有可读条目
     「本行无定额依据，工期与人数来自模型（原因：××）」，且原因归到合同 §9.1 的类别
     （KB无定额行 / 定额口径不符 / 单位不可用 / 口径无法对齐 / 活动绑定不一致）；
  ② **D6**：每条工序标来源档次（规范台班 / 规范人工 / AI 定额（已审）/ 模型估算 / 无依据），
     并给出**关键路径规范依据覆盖率**（分子只含规范定额两档，AI 无规范依据不计入），
     写进 `meta.norm_coverage.critical_norm_coverage`（**不新增 meta 顶层键**）；
  ③ **D7**：人工覆盖入口 —— **没有覆盖文件时行为完全不变**（一个字都不多），
     有覆盖文件时逐条呈现「原值 → 覆盖值 → 覆盖人 / 时间 / 说明」。

外加合同 §9.4：口径换算留痕（`norm_binding.basis_adjust` / `basis_unconfirmed`，
WS1 并行开发 → 字段可能不存在，**缺了也必须与改造前逐字一致**）。

判据全部取数据：本文件的 task_id 一律是**故意随便取的**，改个名字/加一条任务，
断言跟着变 —— 没有任何写死的 10 条名单（见 `TestD4IsDataDriven`）。

运行：python -m pytest backend/tests/test_final_ws3_delivery_norm.py -q
"""

import json
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from pipeline import config                          # noqa: E402
from pipeline.nodes import delivery as D             # noqa: E402


# ══════════════════════ 公共样例 ══════════════════════

def _sched(tid, name, start, days):
    """排程行（起止日期 → 工期真源；`_schedule_span` 只认日期）。"""
    from datetime import date, timedelta
    d0 = date.fromisoformat(start)
    d1 = d0 + timedelta(days=days - 1)
    return {"task_id": tid, "task_name": name, "start_date": d0.isoformat(),
            "finish_date": d1.isoformat(), "duration_days": days, "assigned_resources": {}}


def _rd(tid, name, *, unit="m³", norm=None, resources=None, flag=None,
        binding=None, **kw):
    t = {"task_id": tid, "task_name": name, "quantity": 100.0, "planned_duration_days": 5}
    if norm is not None:
        t["_norm_applied"] = norm
    if resources:
        t["resources"] = resources
    if flag:
        t["_norm_flagged"] = flag
        t["_warning"] = "定额不可作证据，未计算班组：" + flag
    t.update(kw)
    return t


LABOR_NORM = {"mode": "labor", "norm_value": 0.1, "source_code": "GD_2018_A1_3",
              "match_type": "exact", "origin": "kb"}
MACHINE_NORM = {"mode": "machine", "norm_value": 1.5, "source_code": "JX_2020_2",
                "match_type": "exact", "origin": "kb"}
AI_NORM = {"mode": "labor", "norm_value": 0.1, "source_code": "AI_ESTIMATE_V1",
           "match_type": "ai", "origin": "ai"}


def _plan(rd_tasks, leaves=None, sched=None, critical=("1.1.1",), total_days=100,
          meta=None, plan_id="plan_ws3_test"):
    """一份最小计划：WBS 叶子 / 排程行 / resource_demand 三份数据一一对应。"""
    if sched is None:
        sched = [_sched(t["task_id"], t["task_name"], "2026-01-01", 5) for t in rd_tasks]
    if leaves is None:
        leaves = [{"id": t["task_id"], "name": t["task_name"], "duration_days": 5,
                   "quantity": 100.0, "unit": "m³", "work_type": "混凝土工程"}
                  for t in rd_tasks]
    plan = {
        "plan_id": plan_id,
        "overview": {"project_name": "WS3 验收样例", "total_duration_days": total_days,
                     "planned_start_date": "2026-01-01", "planned_end_date": "2026-04-10",
                     "critical_path_length": len(critical)},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "1.1", "name": "结构", "sub_packages": leaves}]}]},
        "dependencies": [],
        "cpm_result": {"total_duration_days": total_days,
                       "critical_path": list(critical),
                       "schedule": [{"task_id": s["task_id"], "es": 0, "ef": 4}
                                    for s in sched]},
        "all_tasks_schedule": sched,
        "key_milestones": [{"name": "开工", "date": "2026-01-01",
                            "task_id": critical[0] if critical else "1.1.1",
                            "description": "开工"}],
        "critical_path_tasks": [dict(sched[0])] if sched else [],
        "resource_demand": {"tasks": rd_tasks},
        "resource_plan": {"total_manpower_days": 100.0, "peak_manpower": 9,
                          "peak_manpower_source": "resource_curve",
                          "curve_peak_manpower": 9, "equipment_peak": {},
                          "material_summary": []},
        "meta": dict(meta or {"audit_status": "未审计"}),
        "report": "# 报告",
    }
    return plan


def _board(plan, tmp_deliverables):
    return Path(D.build_plan_html(plan)).read_text(encoding="utf-8")


def _word(plan, tmp_deliverables):
    from docx import Document
    doc = Document(D.build_plan_docx(plan))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parts.extend(c.text for c in r.cells)
    return "\n".join(parts)


@pytest.fixture()
def tmp_deliverables(tmp_path, monkeypatch):
    """把交付物目录指向 tmp，避免测试往 `输出结果/` 里写东西。"""
    d = tmp_path / "deliverables"
    d.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(config, "DELIVERABLES_DIR", d)
    monkeypatch.setattr(D.config, "DELIVERABLES_DIR", d)
    return d


class _FakeLLM:
    def __init__(self, text):
        self.text = text

    def chat_text(self, *a, **kw):
        return self.text


BARE_LLM_HTML = ("<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'>"
                 "<title>模型编排</title></head><body><h1>施工进度计划看板</h1>"
                 "<p>模型只写了个壳</p></body></html>")


# ═══════════════ ① D6 来源档次判据（逐条，纯数据） ═══════════════

class TestTierJudgement:
    def test_规范台班与规范人工按mode分档(self):
        assert D._norm_tier_of(_rd("a", "n", norm=MACHINE_NORM))[0] == D.NORM_TIER_MACHINE
        assert D._norm_tier_of(_rd("a", "n", norm=LABOR_NORM))[0] == D.NORM_TIER_LABOR

    def test_AI定额按来源判不是按任务名(self):
        rd = _rd("a", "名字里完全没有 AI 字样", norm=AI_NORM)
        assert D._norm_tier_of(rd)[0] == D.NORM_TIER_AI
        # 反过来：名字里带 AI、但定额是真人来源 → 仍是规范档
        rd2 = _rd("b", "AI 审计要点", norm=LABOR_NORM)
        assert D._norm_tier_of(rd2)[0] == D.NORM_TIER_LABOR

    def test_没有定额时有资源算模型估算_连资源都没有算无依据(self):
        assert D._norm_tier_of(
            _rd("a", "n", resources={"瓦工": {"per_day": 3}}))[0] == D.NORM_TIER_MODEL
        assert D._norm_tier_of(_rd("b", "n"))[0] == D.NORM_TIER_NONE
        # 空记录 / 非 dict 一律「无依据」，不许猜
        assert D._norm_tier_of({})[0] == D.NORM_TIER_NONE
        assert D._norm_tier_of(None)[0] == D.NORM_TIER_NONE

    def test_旧计划AI拦截但没算出班组时不许标已审(self):
        """`_norm_applied` 为空 = 那条 AI 定额这一回**没参与计算** → 不许标「AI 定额（已审）」。"""
        rd = _rd("a", "n", flag="AI 估算定额（只作参考，不用来算班组）")
        assert D._norm_tier_of(rd)[0] == D.NORM_TIER_NONE
        # 有资源但定额没参与 → 模型估算
        rd2 = _rd("b", "n", flag="AI 估算定额（只作参考，不用来算班组）",
                  resources={"瓦工": {"per_day": 3}})
        assert D._norm_tier_of(rd2)[0] == D.NORM_TIER_MODEL

    def test_逐条档次前缀进依据列(self):
        """D6「每条工序标注档次」的落点：依据 / 资源 列的档次前缀（Word 与看板同一真源）。"""
        for rd, want in ((_rd("a", "n", norm=MACHINE_NORM), D.NORM_TIER_MACHINE),
                         (_rd("b", "n", norm=LABOR_NORM), D.NORM_TIER_LABOR),
                         (_rd("c", "n", norm=AI_NORM), D.NORM_TIER_AI),
                         (_rd("d", "n", resources={"瓦工": {"per_day": 1}}), D.NORM_TIER_MODEL),
                         (_rd("e", "n", flag="KB 无定额行"), D.NORM_TIER_NONE)):
            assert D._evidence_text(rd).startswith("【%s】" % want), D._evidence_text(rd)
        # 什么都不知道的行**不加前缀**（贴档次就是编），既有逐字文案保持不变
        assert D._evidence_text({}) == "来源未记录（该任务没有资源记录）"
        assert D._evidence_text(_rd("f", "空记录")) == "来源未记录（无资源记录）"


# ═══════════════ ② D6 关键路径规范依据覆盖率 ═══════════════

class TestCriticalCoverage:
    @staticmethod
    def _plan_for(coverage_tasks, ai_critical=False):
        """3 条任务：1 条规范人工(10天) + 1 条 AI/规范台班(20天) + 1 条非关键(30天)。

        总工期 100 天；关键路径两条。分子只看规范两档：
          · AI 情形 → 分子 10 → 10.0%
          · 规范台班情形 → 分子 30 → 30.0%
        """
        second = AI_NORM if ai_critical else MACHINE_NORM
        rds = [_rd("k1", "关键一", norm=LABOR_NORM, resources={"瓦工": {"per_day": 3}}),
               _rd("k2", "关键二", norm=second, resources={"瓦工": {"per_day": 3}}),
               _rd("n1", "非关键", norm=LABOR_NORM, resources={"瓦工": {"per_day": 3}})]
        sched = [_sched("k1", "关键一", "2026-01-01", 10),
                 _sched("k2", "关键二", "2026-01-11", 20),
                 _sched("n1", "非关键", "2026-02-01", 30)]
        return _plan(rds, sched=sched, critical=("k1", "k2"), total_days=100)

    def test_覆盖率按规范两档算且AI不计入分子(self):
        cov = D._norm_critical_coverage(self._plan_for(None, ai_critical=True))
        assert cov["norm_days"] == 10.0 and cov["total_days"] == 100.0, cov
        assert cov["pct"] == 10.0, cov
        assert cov["tiers"][D.NORM_TIER_AI] == 20.0, cov
        cov2 = D._norm_critical_coverage(self._plan_for(None, ai_critical=False))
        assert cov2["norm_days"] == 30.0 and cov2["pct"] == 30.0, cov2

    def test_覆盖率分子只数关键路径(self):
        cov = D._norm_critical_coverage(self._plan_for(None))
        assert cov["critical_tasks"] == 2, cov          # 非关键那条 30 天不进分子
        assert cov["critical_days"] == 30.0, cov

    def test_没有关键路径或没有日期时算不出来(self):
        plan = self._plan_for(None)
        plan["cpm_result"]["critical_path"] = []
        assert D._norm_critical_coverage(plan) is None
        plan = self._plan_for(None)
        plan["all_tasks_schedule"] = []
        assert D._norm_critical_coverage(plan) is None

    def test_目标80未达标如实写(self):
        cov = D._norm_critical_coverage(self._plan_for(None, ai_critical=True))
        txt = D._norm_coverage_value_text(cov)
        assert "%s%%" % D._fnum(cov["pct"]) in txt, txt
        assert "≥80%" in txt and "未达标" in txt, txt

    def test_写进meta挂在norm_coverage下且不新增顶层键(self):
        plan = self._plan_for(None)
        plan["meta"]["norm_coverage"] = {"total": 3, "bound": 3}
        before = set(plan["meta"].keys())
        assert D._record_norm_coverage(plan) is True
        assert set(plan["meta"].keys()) == before, "不许新增 meta 顶层键"
        rec = plan["meta"]["norm_coverage"][D.NORM_COVERAGE_META_KEY]
        assert rec["pct"] == 30.0 and rec["norm_days"] == 30.0, rec
        assert rec["target_pct"] == 80.0
        # 没有 norm_coverage 这个结构 → 一个字都不写（不新建）
        plan2 = self._plan_for(None)
        assert D._record_norm_coverage(plan2) is False
        assert "norm_coverage" not in plan2["meta"]

    def test_写回只改已存在的计划文件不新建(self, tmp_deliverables, tmp_path, monkeypatch):
        monkeypatch.setattr(config, "PLANS_DIR", tmp_path / "plans")
        assert D._persist_norm_coverage(
            {"plan_id": "不存在"}, {"pct": 1.0}) == ""      # 目录都没有 → 不新建

    def test_写回把子键挂进已存在的计划JSON(self, tmp_path, monkeypatch):
        plans = tmp_path / "plans"
        plans.mkdir()
        monkeypatch.setattr(config, "PLANS_DIR", plans)
        plan = self._plan_for(None)
        plan["meta"]["norm_coverage"] = {"total": 3, "bound": 3}
        (plans / "p1.json").write_text(
            json.dumps({"plan_id": "p1", "meta": {"norm_coverage": {"total": 3}}},
                       ensure_ascii=False), encoding="utf-8")
        plan["plan_id"] = "p1"
        assert D._record_norm_coverage(plan) is True
        disk = json.loads((plans / "p1.json").read_text(encoding="utf-8"))
        rec = disk["meta"]["norm_coverage"][D.NORM_COVERAGE_META_KEY]
        assert rec["pct"] == 30.0 and rec["norm_days"] == 30.0, rec
        assert "critical_norm_coverage" not in disk, "不许新增 meta 顶层键"
        # 磁盘上的其它键一个字不动（同格式、可被 PlanJson 读回）
        assert disk["meta"]["norm_coverage"]["total"] == 3

    def test_覆盖率出现在看板与Word且同一数值(self, tmp_deliverables):
        plan = self._plan_for(None, ai_critical=True)
        plan["meta"]["norm_coverage"] = {"total": 3, "bound": 3, "bound_pct": 100.0,
                                         "unbound": 0, "unbound_pct": 0.0}
        h = _board(plan, tmp_deliverables)
        w = _word(plan, tmp_deliverables)
        for txt in (D.NORM_COVERAGE_LABEL, D.NORM_TIER_TITLE):
            assert txt in h, "看板缺：%s" % txt
            assert txt in w, "Word 缺：%s" % txt
        assert "%s%%" % D._fnum(10.0) in D._confidence_section_html(plan, D._compute_view(plan))
        assert "%s%%" % D._fnum(10.0) in w

    def test_档次表只列出现过的档次(self, tmp_deliverables):
        """计划里没有 AI 档次时，正文不该大写「AI 定额（已审）」（既有回归门钉过这条）。"""
        plan = self._plan_for(None, ai_critical=False)
        plan["meta"]["norm_coverage"] = {"total": 3, "bound": 3}
        card = D._confidence_section_html(plan, D._compute_view(plan))
        assert D.NORM_TIER_MACHINE in card and D.NORM_TIER_LABOR in card
        assert "AI 经验估算定额" not in card, card[:1500]


# ═══════════════ ③ D4 无定额依据工序 ═══════════════

class TestD4MissingNorm:
    @staticmethod
    def _missing_plan():
        """5 条 `_norm_applied` 为空的工序，原因分别命中合同 §9.1 的五个类别。"""
        reasons = {
            "x1": ("KB无定额行", "AI估算定额：KB 无定额行，只作参考"),
            "x2": ("定额口径不符", "定额口径不符：定额分母与工序口径不一致"),
            "x3": ("单位不可用",
                   "单位不可用：单位不一致且不可换算：任务「根」 vs 定额分母「m³」"),
            "x4": ("口径无法对齐", "口径无法对齐：任务口径与定额口径之间没有换算依据"),
            "x5": ("活动绑定不一致", ""),
        }
        rds, leaves = [], []
        for tid, (_code, raw) in reasons.items():
            rds.append(_rd(tid, "无定额工序 " + tid, flag=(raw or None),
                           resources={"普工": {"per_day": 2, "total_days": 10.0}}))
            nb = {}
            if tid == "x5":
                nb = {"usable": False, "match_type": "unbound",
                      "not_usable_reason": "活动绑定不一致"}
            leaves.append({"id": tid, "name": "无定额工序 " + tid, "duration_days": 5,
                           "quantity": 1.0, "unit": "项", "work_type": "其他",
                           "norm_binding": nb})
        return _plan(rds, leaves=leaves,
                     sched=[_sched(t, "无定额工序 " + t, "2026-01-01", 5)
                            for t in reasons],
                     critical=("x1",), total_days=50)

    def test_五类原因都归得出来(self):
        rows, total = D._norm_missing_rows(self._missing_plan())
        assert total == 5, rows
        got = dict((r["task_id"], r["reason_code"]) for r in rows)
        assert got == {"x1": "KB无定额行", "x2": "定额口径不符", "x3": "单位不可用",
                       "x4": "口径无法对齐", "x5": "活动绑定不一致"}, got

    def test_原因判据优先取绑定的具体原因(self):
        """`_norm_flagged` 是笼统的 AI 旧文案、而绑定说的是单位不可用 → 取具体的那条。"""
        plan = _plan(
            [_rd("y1", "n", flag="AI 估算定额（只作参考，不用来算班组）")],
            leaves=[{"id": "y1", "name": "n", "duration_days": 5, "quantity": 1.0,
                     "unit": "项", "work_type": "其他",
                     "norm_binding": {"usable": False,
                                      "not_usable_reason": "单位不可用：不可换算"}}])
        rows, _t = D._norm_missing_rows(plan)
        assert rows[0]["reason_code"] == "单位不可用", rows

    def test_拿不到原因就写来源未记录不编(self):
        plan = _plan([_rd("z1", "n", resources={"普工": {"per_day": 1}})])
        rows, _t = D._norm_missing_rows(plan)
        assert rows[0]["reason"] == "来源未记录" and rows[0]["reason_code"] == "", rows
        assert "原因：来源未记录" in rows[0]["sentence"], rows

    def test_合同要求的逐条句子出现在看板与Word(self, tmp_deliverables):
        plan = self._missing_plan()
        plan["meta"]["norm_coverage"] = {"total": 5, "bound": 0, "bound_pct": 0.0,
                                        "unbound": 5, "unbound_pct": 100.0}
        h = _board(plan, tmp_deliverables)
        w = _word(plan, tmp_deliverables)
        for txt in (D.NORM_MISSING_TITLE, "本行无定额依据，工期与人数来自模型"):
            assert txt in h, "看板缺：%s" % txt
            assert txt in w, "Word 缺：%s" % txt
        for code in ("KB无定额行", "定额口径不符", "单位不可用", "口径无法对齐",
                     "活动绑定不一致"):
            assert "原因：%s" % code in h, "看板缺原因：%s" % code
            assert "原因：%s" % code in w, "Word 缺原因：%s" % code

    def test_没有这类工序时一个块都不加(self):
        plan = _plan([_rd("ok1", "n", norm=LABOR_NORM,
                          resources={"瓦工": {"per_day": 1}})])
        _rows, total = D._norm_missing_rows(plan)
        assert total == 0
        assert D._norm_missing_card_html(plan) == ""

    def test_结构性保证_模型页面缺这段就追加(self, tmp_deliverables, monkeypatch):
        plan = self._missing_plan()
        plan["meta"]["norm_coverage"] = {"total": 5, "bound": 0}
        ctx = {}
        path, used = D.build_plan_html_agent(plan, _FakeLLM(BARE_LLM_HTML), ctx=ctx)
        html = Path(path).read_text(encoding="utf-8")
        assert used is True
        assert D.NORM_MISSING_MARKER in html, "最终页面上必须有这段（无论由哪个保证补上）"
        assert D.NORM_MISSING_TITLE in html
        # 追加段本身也要能独立触发（模型页面已有置信度段时，仍然不许缺这一节）
        merged, added = D._ensure_norm_missing_section(BARE_LLM_HTML, plan)
        assert added is True and D.NORM_MISSING_TITLE in merged

    def test_结构性保证_页面已有该标记就不重复追加(self):
        plan = self._missing_plan()
        page = (BARE_LLM_HTML.replace("</body>", "<p>本行无定额依据，工期与人数来自模型"
                                                "（原因：KB无定额行）</p></body>"))
        merged, added = D._ensure_norm_missing_section(page, plan)
        assert added is False and merged == page

    def test_结构性保证_追加失败保持原页面不动(self, monkeypatch):
        plan = self._missing_plan()

        def _boom(_plan):
            raise RuntimeError("渲染炸了")

        monkeypatch.setattr(D, "_norm_missing_card_html", _boom)
        page = BARE_LLM_HTML
        merged, added = D._ensure_norm_missing_section(page, plan)
        assert added is False and merged == page


class TestD4IsDataDriven:
    def test_名单跟着数据走不写死task_id(self):
        """合同点名的 10 条只是**当前基线**：换 id、加条数，清单跟着变。"""
        a = _plan([_rd("1.3.4", "临时设施", flag="KB 无定额行")],
                  leaves=[{"id": "1.3.4", "name": "临时设施", "duration_days": 3,
                           "quantity": 1.0, "unit": "项"}])
        rows_a, total_a = D._norm_missing_rows(a)
        assert total_a == 1 and rows_a[0]["task_id"] == "1.3.4"
        # 把 id 换成完全无关的名字，判据必须仍然成立（不看名字/不看固定清单）
        b = _plan([_rd("zz.9.9", "随便改个名字", flag="KB 无定额行")],
                  leaves=[{"id": "zz.9.9", "name": "随便改个名字", "duration_days": 3,
                           "quantity": 1.0, "unit": "项"}])
        rows_b, total_b = D._norm_missing_rows(b)
        assert total_b == 1 and rows_b[0]["task_id"] == "zz.9.9"
        # 修好一条（补上 `_norm_applied`）→ 清单自己少一条
        fixed = json.loads(json.dumps(a))
        fixed["resource_demand"]["tasks"][0]["_norm_applied"] = dict(LABOR_NORM)
        assert D._norm_missing_rows(fixed)[1] == 0
        # 再加一条没有定额的 → 清单自己多一条
        more = json.loads(json.dumps(a))
        more["resource_demand"]["tasks"].append(
            _rd("1.3.5", "新增无定额", flag="KB 无定额行"))
        more["all_tasks_schedule"].append(_sched("1.3.5", "新增无定额", "2026-02-01", 3))
        assert D._norm_missing_rows(more)[1] == 2


# ═══════════════ ④ D7 人工覆盖入口（优雅降级） ═══════════════

class TestD7Override:
    @staticmethod
    def _plan_with_conf():
        plan = _plan([_rd("o1", "被覆盖的工序", norm=LABOR_NORM,
                          resources={"瓦工": {"per_day": 3}})])
        plan["meta"]["norm_coverage"] = {"total": 1, "bound": 1}
        return plan

    def test_没有覆盖文件时行为完全不变(self, tmp_deliverables):
        plan = self._plan_with_conf()
        assert D._norm_override_path(plan) is None
        assert D._norm_overrides(plan) == ([], "")
        h = _board(plan, tmp_deliverables)
        w = _word(plan, tmp_deliverables)
        for txt in (D.NORM_OVERRIDE_TITLE, "人工覆盖留痕", "覆盖人"):
            assert txt not in h, "看板不该出现覆盖段：%s" % txt
            assert txt not in w, "Word 不该出现覆盖段：%s" % txt
        # 追加段也不该被这一项触发
        assert D.NORM_OVERRIDE_TITLE not in D._norm_tier_card_html(plan)

    def test_有覆盖文件时逐条呈现原值到覆盖值(self, tmp_deliverables):
        plan = self._plan_with_conf()
        d = D._plan_dir(plan)
        (d / D.NORM_OVERRIDE_FILE).write_text(json.dumps({"overrides": [
            {"task_id": "o1", "field": "工期", "original": 5, "value": 7,
             "by": "张三", "at": "2026-09-21T10:00:00", "note": "现场踏勘后调整"}]},
            ensure_ascii=False), encoding="utf-8")
        rows, src = D._norm_overrides(plan)
        assert len(rows) == 1 and rows[0]["by"] == "张三", rows
        assert str(d) in src
        h = _board(plan, tmp_deliverables)
        w = _word(plan, tmp_deliverables)
        card = D._norm_tier_card_html(plan)
        assert D.NORM_OVERRIDE_TITLE in h and D.NORM_OVERRIDE_TITLE in w
        assert D.NORM_OVERRIDE_TITLE in card
        for txt in ("5 → 7", "张三", "2026-09-21T10:00:00", "现场踏勘后调整"):
            assert txt in h, "看板缺：%s" % txt
            assert txt in w, "Word 缺：%s" % txt

    def test_坏文件与坏结构一律当作没有覆盖(self, tmp_deliverables):
        plan = self._plan_with_conf()
        d = D._plan_dir(plan)
        p = d / D.NORM_OVERRIDE_FILE
        p.write_text("{ 这不是 JSON", encoding="utf-8")
        assert D._norm_overrides(plan) == ([], "")
        p.write_text(json.dumps({"overrides": "不是列表"}), encoding="utf-8")
        assert D._norm_overrides(plan) == ([], "")

    def test_覆盖入口只展示不改计划数值(self, tmp_deliverables):
        """WS3 只负责交付物呈现：覆盖文件**不得**改动计划里的工期/资源数。"""
        plan = self._plan_with_conf()
        d = D._plan_dir(plan)
        (d / D.NORM_OVERRIDE_FILE).write_text(json.dumps({"overrides": [
            {"task_id": "o1", "field": "工期", "original": 5, "value": 99}]},
            ensure_ascii=False), encoding="utf-8")
        before = json.dumps([plan["resource_demand"], plan["all_tasks_schedule"]],
                            ensure_ascii=False, sort_keys=True)
        _board(plan, tmp_deliverables)
        _word(plan, tmp_deliverables)
        assert json.dumps([plan["resource_demand"], plan["all_tasks_schedule"]],
                          ensure_ascii=False, sort_keys=True) == before


# ═══════════════ ⑤ 口径换算留痕（WS1 字段，缺失即不变） ═══════════════

class TestBasisAdjust:
    @staticmethod
    def _plan_with_basis():
        leaf = {"id": "b1", "name": "口径待换算工序", "duration_days": 5, "quantity": 100.0,
                "unit": "m²", "work_type": "砌筑工程",
                "norm_binding": {"unit": "工日/m³", "usable": True,
                                 "basis_adjust": {
                                     "task_scope": "m²（墙面）",
                                     "norm_scope": "m³（砌体体积）",
                                     "task_quantity": 100.0,
                                     "adjusted_quantity": 20.0,
                                     "method": "× 墙厚 0.2 m",
                                     "note": "按图纸墙厚"}}}
        plan = _plan([_rd("b1", "口径待换算工序", norm=LABOR_NORM,
                          resources={"瓦工": {"per_day": 3}})],
                     leaves=[leaf], critical=("b1",), total_days=50)
        plan["meta"]["norm_coverage"] = {"total": 1, "bound": 1}
        return plan

    def test_WS1没写字段时整段不出_行为与改造前一致(self, tmp_deliverables):
        plan = _plan([_rd("b1", "普通工序", norm=LABOR_NORM,
                          resources={"瓦工": {"per_day": 3}})])
        plan["meta"]["norm_coverage"] = {"total": 1, "bound": 1}
        rows, total = D._norm_basis_adjust_rows(plan)
        assert rows == [] and total == 0
        h = _board(plan, tmp_deliverables)
        w = _word(plan, tmp_deliverables)
        assert D.NORM_BASIS_TITLE not in h and D.NORM_BASIS_TITLE not in w
        assert "口径换算留痕" not in h and "口径换算留痕" not in w

    def test_有换算留痕时四个要素都印出来(self, tmp_deliverables):
        plan = self._plan_with_basis()
        h = _board(plan, tmp_deliverables)
        w = _word(plan, tmp_deliverables)
        card = D._norm_lists_card_html(plan)
        assert D.NORM_BASIS_TITLE in h and D.NORM_BASIS_TITLE in w
        assert D.NORM_BASIS_TITLE in card, "合同 §9.4 要求在「定额降级/口径」区块里列出"
        for txt in ("m²（墙面）", "m³（砌体体积）", "100 → 20", "× 墙厚 0.2 m", "按图纸墙厚"):
            assert txt in h, "看板缺：%s" % txt
            assert txt in w, "Word 缺：%s" % txt

    def test_口径未确认如实打标不阻断(self):
        plan = self._plan_with_basis()
        nb = plan["wbs"]["phases"][0]["work_packages"][0]["sub_packages"][0]["norm_binding"]
        nb.pop("basis_adjust")
        nb["basis_unconfirmed"] = True
        rows, total = D._norm_basis_adjust_rows(plan)
        assert total == 1 and rows[0]["unconfirmed"] is True
        assert "口径未确认" in D._norm_basis_adjust_text(rows[0])
        # 绑定仍然可用 → 不作为降级处理（不阻断）
        assert D._norm_tier_of(plan["resource_demand"]["tasks"][0])[0] == D.NORM_TIER_LABOR

    def test_结构性保证_有留痕而页面没写就追加(self):
        plan = self._plan_with_basis()
        merged, added = D._ensure_norm_tier_section(BARE_LLM_HTML, plan)
        assert added is True
        assert D.NORM_BASIS_TITLE in merged and D.NORM_COVERAGE_LABEL in merged

    def test_结构性保证_没有留痕也没档次数据时不动页面(self):
        plan = _plan([])
        merged, added = D._ensure_norm_tier_section(BARE_LLM_HTML, plan)
        assert added is False and merged == BARE_LLM_HTML
