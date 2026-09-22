# -*- coding: utf-8 -*-
"""关键路径口径：**条数（个）** 不许被写成 **天数（天）**

真实缺陷（真计划 `backend/plans/plan_run_1789895021.json`，`overview.critical_path_length=81`、
`cpm_result.total_duration_days=604`）：落盘的 `plan["report"]`（模型写的 Markdown 监督报告）里写着

    …当前计划总工期为604天，**关键路径长度为81天**，表明非关键路径任务具有一定的浮动时间。

—— 81 是**关键路径任务的条数（个）**，被写成了天数；而同一份交付物的表格里明明已经写着
「关键路径任务数 81 个」「关键路径工期·排程版 604 天」（同一个 `report` 里还有第二处
「关键路径由81项任务组成，总时长81天」）。

本文件守住三道防线（一条都不靠模型自觉）：
  ① **输入侧去歧义**：`plan_assembler.build_parts` 交给报告模型的 overview 里**没有**
     `critical_path_length` 这个键名（换成 `critical_path_task_count` + 一句 note），
     另给 `critical_path_caliber`；`delivery._facts_bundle`（交付 facts）同样处理。
     落盘时 `assemble_plan_json` 把契约键 `critical_path_length` 还原回去 —— 契约不破。
  ② **提示词铁律**：`prompts/report.txt` / `prompts/deliver_html.txt` 点名禁用「关键路径长度」。
  ③ **确定性改写**（真正的不变量）：`plan_assembler.fix_critical_path_wording`，交付侧
     `delivery._report_text`（docx / 看板 / facts 三条出口）与
     `delivery._fix_critical_path_narrative`（模型编排的 HTML 整页）都过这一道 ——
     不管报告是这次模型写的、还是**上一次运行存档的**，印出来之前都把条数当天数的句子改对；
     数字与计划真值对不上、或真值取不到，就**一个字都不动**（宁缺勿造）。

运行：python -m pytest backend/tests/test_critical_path_wording.py -q -p no:cacheprovider
"""

import json
import re
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

import pytest

from pipeline import schemas                          # noqa: E402
from pipeline.nodes import delivery as D              # noqa: E402
from pipeline.nodes import plan_assembler as PA       # noqa: E402

REAL_PLAN = BACKEND / "plans" / "plan_run_1789895021.json"

# 真计划里那句缺陷原文（逐字，来自 `plan_run_1789895021.json` 的 `plan["report"]`）。
BAD_REPORT = (
    "# 监督报告\n\n"
    "## 一、总体情况\n"
    "- 项目：海之子\n- 总工期：604 天\n- **关键路径任务数**：81项\n\n"
    "## 三、关键路径分析\n"
    "关键路径由81项任务组成，总时长81天。该路径始于“周边环境监测”，直至竣工验收。\n"
    "*   **对总工期的影响**：关键路径上的任何任务延误都将直接导致项目总工期的延长。"
    "当前计划总工期为604天，关键路径长度为81天，表明非关键路径任务具有一定的浮动时间。\n"
)
BAD_SENTENCE = "当前计划总工期为604天，关键路径长度为81天，表明非关键路径任务具有一定的浮动时间。"


# ==================== 夹具 ====================
def _plan():
    """最小计划：**81 项关键任务 / 总工期 604 天**（条数与天数天差地别，误读一眼可见）。"""
    leaves = [{"id": "1.1.1", "name": "柱浇筑", "duration_days": 5,
               "quantity": 10, "unit": "m³", "work_type": "混凝土工程"}]
    sched = [{"task_id": "1.1.1", "task_name": "柱浇筑",
              "start_date": "2026-06-01", "finish_date": "2026-06-05", "duration_days": 5,
              "assigned_resources": {"钢筋工": 4}}]
    return {
        "plan_id": "plan_cp_wording_test",
        "overview": {"project_name": "关键路径口径测试", "total_duration_days": 604,
                     "planned_start_date": "2026-06-01", "planned_end_date": "2028-01-25",
                     "critical_path_length": 81},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "1.1", "name": "结构", "sub_packages": leaves}]}]},
        "dependencies": [],
        "cpm_result": {"total_duration_days": 604, "cpm_total_duration_days": 644,
                       # 81 条关键任务（`build_parts` 的 overview 条数就是 len(这一串)）
                       "critical_path": ["1.1.%d" % i for i in range(1, 82)],
                       "schedule": [{"task_id": "1.1.1", "es": 0, "ef": 5,
                                     "ls": 0, "lf": 5}]},
        "all_tasks_schedule": sched,
        "critical_path_tasks": [dict(sched[0])],
        "key_milestones": [{"name": "开工", "date": "2026-06-01", "task_id": "1.1.1",
                            "description": "开工"}],
        "resource_demand": {"tasks": []},
        "resource_plan": {"total_manpower_days": 20.0, "peak_manpower": 4,
                          "equipment_peak": {}, "material_summary": []},
        "meta": {"audit_status": "未审计"},
        "report": BAD_REPORT,
    }


def _ctx(plan):
    """`build_parts` 需要的 ctx（与 `test_duration_single_source._assemble` 同形）。"""
    ov = plan.get("overview") or {}
    params = {"planned_start_date": ov.get("planned_start_date") or "2026-06-01",
              "project_name": ov.get("project_name") or "未命名"}
    return {
        "wbs": json.loads(json.dumps(plan.get("wbs") or {})),
        "cpm_result": json.loads(json.dumps(plan.get("cpm_result") or {})),
        "resource_demand": json.loads(json.dumps(plan.get("resource_demand") or {})),
        "extracted_params": params,
        "boundary_conditions": plan.get("boundary_conditions"),
    }


def _docx_text(plan):
    """Word 全文（段落 + 表格单元格，拼一起）。"""
    from docx import Document

    doc = Document(D.build_plan_docx(plan))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parts.extend(c.text for c in r.cells)
    return "\n".join(parts)


def _html_text(plan):
    """`build_plan_html` 落盘后读回文本（导出节点，路径在隔离的交付目录里）。"""
    return Path(D.build_plan_html(plan)).read_text(encoding="utf-8")


class _FakeLLM:
    """只实现 `chat_text` 的假模型（与 test_delivery_org_visible 同一惯例）。"""

    def __init__(self, text):
        self.text = text
        self.calls = 0

    def chat_text(self, *args, **kwargs):
        self.calls += 1
        return self.text


# ==================== ① 输入侧去歧义 ====================
def test_build_parts_overview_drops_the_ambiguous_key():
    """给报告模型的 parts 里**不出现** `critical_path_length` 这个键名（歧义源）。"""
    plan = _plan()
    parts = PA.build_parts(_ctx(plan))
    ov = parts["overview"]

    assert "critical_path_length" not in ov, "歧义键名还在喂给模型"
    assert ov["critical_path_task_count"] == 81
    assert "不是天数" in ov["critical_path_task_count_note"]
    # 整个 parts 序列化后，也不能出现以它为**键名**的 JSON 字段
    raw = json.dumps(parts, ensure_ascii=False)
    assert '"critical_path_length"' not in raw, "仍然有以歧义键名为键的字段"
    # 两个数各自的名字 / 单位 / 来源键都在
    cal = parts["critical_path_caliber"]
    assert cal["critical_path_task_count"] == {
        "label": "关键路径任务数", "value": 81, "unit": "个",
        "meaning": "条数（个），不是天数", "source": "overview.critical_path_length"}
    assert cal["critical_path_duration_days"]["label"] == "关键路径工期·排程版"
    assert cal["critical_path_duration_days"]["value"] == 604
    assert cal["critical_path_duration_days"]["unit"] == "天"
    assert "禁用" in cal["rule"]


def test_assemble_plan_json_restores_the_contract_key():
    """落盘必须还原契约键 `critical_path_length`（schemas.Overview 硬要求）。"""
    plan = _plan()
    ctx = _ctx(plan)
    parts = PA.build_parts(ctx)
    out = PA.assemble_plan_json(ctx, parts, report="# 报告")

    assert out["overview"]["critical_path_length"] == 81
    assert "critical_path_task_count" not in out["overview"]
    assert "critical_path_task_count_note" not in out["overview"]
    # 与 `test_contracts.py:105` 同一条契约断言：条数 = len(cpm_result.critical_path)
    assert out["overview"]["critical_path_length"] == len(out["cpm_result"]["critical_path"])
    # 契约校验：型号不变（真源 `schemas.PlanJson`）
    validated = schemas.PlanJson.model_validate(out)
    assert validated.overview.critical_path_length == 81
    assert validated.overview.total_duration_days == 604


def test_template_report_accepts_both_part_shapes():
    """确定性模板报告两种 parts 形状都认，且**永不**出现「关键路径长度」。"""
    ov = {"project_name": "X", "total_duration_days": 604,
          "planned_start_date": "2026-06-01", "planned_end_date": "2028-01-25"}
    base = {"key_milestones": [], "critical_path_tasks": [],
            "resource_plan": {"peak_manpower": 3, "total_manpower_days": 20.0},
            "risks": [], "display_granularity": {"note": "逐任务"}}
    # 契约键形状（老调用方 / 手工 parts）
    r1 = PA.template_report(dict(base, overview=dict(ov, critical_path_length=81)))
    # build_parts 给模型的无歧义形状
    r2 = PA.template_report(dict(base, overview=dict(ov, critical_path_task_count=81)))
    for r in (r1, r2):
        assert "关键路径长度" not in r
        assert "关键路径任务数：81 个（条数，不是天数）" in r


def test_facts_bundle_has_no_ambiguous_key():
    """交付 facts（交付网页提示词的输入）同样去歧义，并带上带标签的口径。"""
    plan = _plan()
    facts = D._facts_bundle(plan, D._compute_view(plan))

    assert "critical_path_length" not in facts["overview"]
    assert facts["overview"]["critical_path_task_count"] == 81
    assert "不是天数" in facts["overview"]["critical_path_task_count_note"]
    assert "critical_path_length" not in facts["key_numbers"]
    assert facts["key_numbers"]["critical_path_task_count"] == 81
    assert facts["critical_path_caliber"]["critical_path_task_count"]["unit"] == "个"
    assert facts["critical_path_caliber"]["critical_path_duration_days"]["value"] == 604
    # 正文（facts["report"]）也已经过改写，模型照抄也不会抄到那句话
    assert "关键路径长度" not in facts["report"]
    assert "关键路径任务数为81 个（条数，不是天数）" in facts["report"]


def test_assemble_plan_json_fixes_the_report_of_new_runs():
    """**新运行**的落盘数据本身也修好了（不用等渲染那一道）。"""
    plan = _plan()
    ctx = _ctx(plan)
    out = PA.assemble_plan_json(ctx, PA.build_parts(ctx), report=BAD_REPORT)
    assert "关键路径长度" not in out["report"]
    assert "关键路径任务数为81 个（条数，不是天数）" in out["report"]


def test_deliver_node_rewrites_report_before_saving(monkeypatch):
    """`PlanDeliverNode` 落盘前再走一道：磁盘上的 plan_json 里不再有那句话。"""
    from pipeline import config

    node = PA.PlanDeliverNode()
    node._emit = lambda *a, **k: None
    node._run_id = "cp_wording_test"
    out = node.run({"plan_json": _plan()})

    assert out["plan_json"]["plan_id"] == "plan_cp_wording_test"
    assert "关键路径长度" not in out["plan_json"]["report"]
    saved = json.loads(
        (config.PLANS_DIR / "plan_cp_wording_test.json").read_text(encoding="utf-8"))
    assert "关键路径长度" not in saved["report"]
    assert "关键路径任务数为81 个（条数，不是天数）" in saved["report"]
    # 契约键照旧在（`_overview_contract` 还原过）
    assert saved["overview"]["critical_path_length"] == 81


# ==================== ② 提示词铁律 ====================
def test_prompts_forbid_the_ambiguous_term():
    """两份提示词都点名禁用「关键路径长度」并给出正确写法（不靠模型自觉）。"""
    from pipeline.prompts_loader import load

    report = load("report.txt")
    assert "关键路径长度" in report and "禁止" in report
    assert "关键路径任务数" in report and "个" in report
    assert "critical_path_caliber" in report
    assert "critical_path_task_count" in report

    html = load("deliver_html.txt")
    assert "关键路径长度" in html
    assert "一次都不许出现" in html
    assert "critical_path_caliber" in html
    assert "关键路径工期·排程版" in html


# ==================== ③ 确定性改写（真正的不变量） ====================
def test_fix_rewrites_count_written_as_days():
    """真计划原句：81（条数）被写成 81 天 → 改成「关键路径任务数 81 个」+「工期 604 天」。"""
    plan = _plan()
    fixed, fixes = PA.fix_critical_path_wording(plan["report"], plan)

    assert "关键路径长度" not in fixed
    assert "关键路径任务数为81 个（条数，不是天数）" in fixed
    # 同段第二处「关键路径由81项任务组成，总时长81天」→ 后半句改回真天数
    assert "关键路径由 81 项任务组成，关键路径工期·排程版为 604 天" in fixed
    assert "总时长81天" not in fixed
    assert len(fixes) == 2, fixes
    assert any("条数当天数" in f for f in fixes)
    assert any("链式句" in f for f in fixes)
    # 改的是那一处，其余原文一字不动
    assert "该路径始于“周边环境监测”，直至竣工验收。" in fixed
    assert "关键路径上的任何任务延误都将直接导致项目总工期的延长。" in fixed


def test_fix_maps_the_real_duration_number_to_the_days_label():
    """数字等于**真天数**（604）时：只把术语换名，数字不动。"""
    fixed, fixes = PA.fix_critical_path_wording("关键路径长度为604天。", _plan())
    assert fixed == "关键路径工期·排程版为604 天。"
    assert fixes and "术语改名" in fixes[0]


def test_fix_renames_a_bare_term_without_number():
    """光是一个术语（后面没数字）→ 换成条数标签，不编任何数字。"""
    fixed, fixes = PA.fix_critical_path_wording("关键路径长度与浮动时间。", _plan())
    assert fixed == "关键路径任务数与浮动时间。"
    assert len(fixes) == 1


def test_fix_leaves_unknown_numbers_alone():
    """数字与计划真值对不上 → **一个字都不动**（绝不猜、绝不拿别的数顶上）。"""
    plan = _plan()
    for text in ("关键路径长度为99天。",          # 既不是条数 81，也不是天数 604
                 "关键路径长度为604项任务。"):   # 数字对，但语义判不出来（没有「天」）
        fixed, fixes = PA.fix_critical_path_wording(text, plan)
        assert fixed == text, text
        assert fixes == [], (text, fixes)


def test_fix_does_nothing_without_truth():
    """真值取不到（老计划 / 手工文本没有 overview / cpm_result）→ 一个字都不动。"""
    for source in (None, {}, {"overview": {}}, {"overview": {"total_duration_days": 604}}):
        fixed, fixes = PA.fix_critical_path_wording(BAD_SENTENCE, source)
        assert fixed == BAD_SENTENCE
        assert fixes == []


def test_fix_is_idempotent():
    """改写两次 == 一次（交付链上 docx / 看板 / facts 可能各过一道）。"""
    plan = _plan()
    once, _ = PA.fix_critical_path_wording(plan["report"], plan)
    twice, again = PA.fix_critical_path_wording(once, plan)
    assert twice == once
    assert again == []


# ==================== ③-落地：docx / 看板 / facts / 模型编排页 ====================
def test_docx_and_dashboard_have_zero_occurrences():
    """真计划那种存档报告：Word 与看板里「关键路径长度」出现 **0 次**，且那句话被改对。"""
    plan = _plan()
    docx_text = _docx_text(plan)
    html_text = _html_text(plan)

    assert "关键路径长度" not in docx_text
    assert "关键路径长度" not in html_text
    for text in (docx_text, html_text):
        assert "关键路径任务数为81 个（条数，不是天数）" in text
        assert "关键路径由 81 项任务组成，关键路径工期·排程版为 604 天" in text
    # 交付物自己算出来的表格仍是老口径的名字（没被改写波及）
    assert "关键路径任务数" in docx_text
    assert "关键路径工期·排程版" in html_text


def test_agent_html_guard_rewrites_and_traces():
    """模型编排的整页 HTML 也过一道：改写那句话、模型原有内容不丢、痕迹进节点状态。"""
    bad_html = ("<!DOCTYPE html><html lang='zh'><head><meta charset='utf-8'>"
                "<title>模型编排</title></head><body><h1>施工进度计划看板</h1>"
                "<p>" + BAD_SENTENCE + "</p></body></html>")
    plan = _plan()
    llm = _FakeLLM(bad_html)
    ctx = {}
    path, used_agent = D.build_plan_html_agent(plan, llm, ctx=ctx)

    assert used_agent is True
    assert llm.calls == 1
    html = Path(path).read_text(encoding="utf-8")
    assert "关键路径长度为81天" not in html
    assert "关键路径任务数为81 个（条数，不是天数）" in html
    assert "当前计划总工期为604天" in html, "模型原有内容不许丢"
    joined = " ".join(ctx.get("wbs_warnings") or [])
    assert "把关键路径**条数**写成了天数" in joined


def test_agent_html_guard_is_silent_without_the_term():
    """页面里没这个词 → 一个字不动、一条告警都不发（与 `_ensure_*` 同口径）。"""
    clean = "<html><body><p>关键路径任务数 81 个；关键路径工期·排程版 604 天。</p></body></html>"
    fixed, fixes = D._fix_critical_path_narrative(clean, _plan())
    assert fixed == clean
    assert fixes == []
    # 术语在、数字对不上 → 同样不动作（宁缺勿造）
    odd = "<html><body><p>关键路径长度为99天。</p></body></html>"
    same, none = D._fix_critical_path_narrative(odd, _plan())
    assert same == odd and none == []


# ==================== 真计划（存在才跑） ====================
@pytest.fixture(scope="module")
def real_plan():
    if not REAL_PLAN.exists():
        pytest.skip("真计划不存在：%s" % REAL_PLAN)
    return json.loads(REAL_PLAN.read_text(encoding="utf-8"))


def test_real_plan_archived_report_is_repaired(real_plan):
    """缺陷指纹 + 修复结果（真计划 `plan_run_1789895021`）：存档原文被改对。"""
    raw = real_plan["report"]
    assert BAD_SENTENCE in raw, "真计划里那句缺陷原文变了，本文件的基线要重新核对"
    assert real_plan["overview"]["critical_path_length"] == 81
    assert real_plan["cpm_result"]["total_duration_days"] == 604

    fixed, fixes = PA.fix_critical_path_wording(raw, real_plan)
    assert "关键路径长度" not in fixed
    assert "关键路径任务数为81 个（条数，不是天数）" in fixed
    assert "关键路径由 81 项任务组成，关键路径工期·排程版为 604 天" in fixed
    assert len(fixes) == 2, fixes
    # 存档 JSON 本身没被改写（渲染前才改写；reporter 不在交付路径上）
    assert real_plan["report"] == raw


def test_real_plan_deliverables_show_zero_occurrences(real_plan):
    """真计划产出 Word + 看板：「关键路径长度」出现 **0 次**。"""
    docx_text = _docx_text(real_plan)
    html_text = _html_text(real_plan)
    assert docx_text.count("关键路径长度") == 0
    assert html_text.count("关键路径长度") == 0
    assert docx_text.count("关键路径任务数为81 个（条数，不是天数）") == 1
    # 表格里的老口径名字仍在（只改那句话，不改交付物自己算的表）
    assert "关键路径任务数" in docx_text
    assert re.search(r"81 个（来源 overview\.critical_path_length）", docx_text), \
        "计划总览表的条数口径被改动了"
