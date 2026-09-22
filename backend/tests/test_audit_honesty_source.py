# -*- coding: utf-8 -*-
"""审计身份（这道门是**谁**答的）—— 「已审计定稿」的判据只能有一条

用户审计 P0-A 查出的是两件事，第二件比第一件更隐蔽：
  ① 计划数据里只有 R1、R2（`audit_status = 未审计`），定稿 Word 却印
     「【已审计定稿】」「三轮回审：R1 通过、R2 通过、R3 通过」；
  ② 就算把 R3 写进数据，那也**不是真人审的** —— 三个门全是
     `devtools/rerun_sample3.py` 用 `REGISTRY.resolve(key, {"passed": True, ...})`
     脚本代答的。`passed=True` 只说明"这道门被放行了"，**不说明有人审过**。

所以本文件守住五件事：
  1. 门被应答时把「谁答的」记进审计链（`answered_by` ∈ human / script / auto / unknown）；
  2. `answered_by` 取不到 → `unknown`，**绝不默认是人工**（宁缺勿假）；
  3. 只有 **三轮都是 human** 才允许印「已审计定稿」；否则「未审计 · 待人工复审」+ 写明原因；
  4. `meta.audit_status` 自称「已审计」不算证据，拿不出人工记录就按未审计；
  5. 每轮结束把审计结论**写回磁盘上的计划数据**（真源与交付物同源，不许两个真相）。

运行：python -m pytest backend/tests/test_audit_honesty_source.py -q -p no:cacheprovider
"""

import json
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

import pytest

from pipeline import config
from pipeline.nodes.audit_gate import (
    DraftAuditNode, ScheduleAuditNode, WBSAuditNode,
    answered_by_of, audit_honesty,
)


# ==================== 夹具 ====================
ROUND_CLASSES = (WBSAuditNode, ScheduleAuditNode, DraftAuditNode)


class _FakeRegistry(object):
    """假的交互登记处：门发 payload，这里回放预设决策（含 `answered_by`）。"""

    def __init__(self, decision=None):
        self.decision = decision if decision is not None else {"passed": True}
        self.registered = []

    def register(self, key):
        self.registered.append(key)

    def wait(self, key, cancel_evt=None, timeout=None):
        return dict(self.decision)


def _plan():
    """最小可渲染计划（字段与 `plan_assembler` 的产物同名同形）。"""
    return {
        "plan_id": "audit_honesty_test",
        "overview": {"project_name": "审计身份测试", "total_duration_days": 100,
                     "planned_start_date": "2026-01-01", "planned_end_date": "2026-04-11",
                     "critical_path_length": 2},
        "meta": {"audit_status": "未审计", "audit_rounds": [], "caliber_note": "单栋口径"},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "5.1", "name": "wp", "sub_packages": [
                {"id": "5.1.1.1", "name": "钢筋", "quantity": 10, "unit": "t",
                 "duration_days": 20}]}]}]},
        "dependencies": [],
        "all_tasks_schedule": [{"task_id": "5.1.1.1", "task_name": "钢筋",
                                "start_date": "2026-01-01", "finish_date": "2026-01-21",
                                "duration_days": 20, "wbs_target_days": 20,
                                "assigned_resources": {"钢筋工": 5}}],
        "key_milestones": [], "critical_path_tasks": [],
        "resource_plan": {"peak_manpower": 5, "total_manpower_days": 100.0,
                          "equipment_peak": {}, "material_summary": []},
        "risks": [], "report": "报告正文",
    }


def _run_rounds(decisions, plan=None):
    """按顺序跑完 R1/R2/R3（真门 + 假登记表），返回 (ctx, plan)。

    `decisions` 就是"打进门里的 resolve 载荷" —— 人工通道会带 `answered_by="human"`，
    自动化脚本不带（或带 "script"），这正是要区分的东西。
    """
    plan = plan if plan is not None else _plan()
    ctx = {"plan_json": plan, "_run_id": "honesty_test"}
    for node_cls, decision in zip(ROUND_CLASSES, decisions):
        node = node_cls()
        node._registry = _FakeRegistry(decision)
        node._run_id = "honesty_test"
        node._cancel_evt = None
        node._emit = lambda e, d: None
        ctx.update(node.run(ctx) or {})
    return ctx, plan


def _docx_text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


HUMAN = {"passed": True, "answered_by": "human"}
SCRIPT = {"passed": True, "answered_by": "script"}
LEGACY = {"passed": True}                       # 老计划/自动化都不带标记


# ==================== 1. 「谁答的门」的取值 ====================
def test_answered_by_never_assumes_human():
    """取不到标记 → unknown；值不认识 → 也不猜（宁缺勿假）。"""
    assert answered_by_of({"passed": True}) == "unknown"
    assert answered_by_of({"passed": True, "decision": True, "choice": "continue",
                           "approved": True}) == "unknown", "脚本那套字段里没有人证"
    assert answered_by_of({"answered_by": "human"}) == "human"
    assert answered_by_of({"answered_by": " HUMAN "}) == "human"
    assert answered_by_of({"answered_by": "script"}) == "script"
    assert answered_by_of({"answered_by": "auto"}) == "auto"
    assert answered_by_of({"answered_by": "机器人"}) == "unknown"
    assert answered_by_of(None) == "unknown"


def test_honesty_judge_requires_three_human_rounds():
    def r(n, passed=True, src=None):
        d = {"round": n, "name": "x", "passed": passed}
        if src:
            d["answered_by"] = src
        return d

    ok = audit_honesty({"audit_rounds": [r(1, True, "human"), r(2, True, "human"),
                                         r(3, True, "human")]})
    assert ok["confirmed"] is True and ok["status"] == "已审计"
    assert ok["reasons"] == []

    # ① 三轮全通过但某轮是脚本代答 → 不算
    bad = audit_honesty({"audit_rounds": [r(1, True, "human"), r(2, True, "human"),
                                          r(3, True, "script")]})
    assert bad["confirmed"] is False and bad["status"] == "未审计"
    assert any("脚本代答" in x for x in bad["reasons"])

    # ② 缺 answered_by（老计划）→ 不算，且原因里点名"缺少记录"
    old = audit_honesty({"audit_rounds": [r(1), r(2), r(3)]})
    assert old["confirmed"] is False
    assert all("缺少「谁答的门」记录" in x for x in old["reasons"])

    # ③ 只有 R1/R2（真计划的实际状态）→ 不算
    partial = audit_honesty({"audit_rounds": [r(1), r(2)]})
    assert partial["confirmed"] is False
    assert any("R3 没有审计记录" in x for x in partial["reasons"])

    # ④ 自称已审计不算证据，必须把矛盾说出来
    lie = audit_honesty({"audit_status": "已审计",
                         "audit_rounds": [r(1, True, "script"), r(2, True, "script"),
                                          r(3, True, "script")]})
    assert lie["confirmed"] is False
    assert any("自称「已审计」" in x for x in lie["reasons"])


# ==================== 2. 门把「谁答的」记进审计链 ====================
def test_gate_records_answered_by_in_audit_chain():
    ctx, plan = _run_rounds([SCRIPT, SCRIPT, SCRIPT])
    rounds = plan["meta"]["audit_rounds"]
    assert [r["round"] for r in rounds] == [1, 2, 3]
    assert all(r["answered_by"] == "script" for r in rounds), rounds
    # 脚本代答跑完，"三轮都通过"是真的，但**不是人工复审** → 数据只能写未审计
    assert ctx.get("audited") is True, "流程上三轮确实都放行了"
    assert plan["meta"]["audit_status"] == "未审计", plan["meta"]["audit_status"]


def test_legacy_payload_records_unknown():
    ctx, plan = _run_rounds([LEGACY, LEGACY, LEGACY])
    assert all(r["answered_by"] == "unknown" for r in plan["meta"]["audit_rounds"])
    assert plan["meta"]["audit_status"] == "未审计"


def test_three_human_rounds_write_back_to_disk():
    """P0-A 的核心：R3 跑完必须把结论写回**磁盘上的计划数据**。

    历史缺陷：`_mark` 只改内存里的 `ctx["plan_json"]`，落盘发生在三个门**之前**的
    `PlanAssemblerNode`，于是磁盘文件永远停在「未审计 / 只有 R1、R2」，而定稿 Word
    从内存渲染出「R3 通过」—— 同一份计划两个真相。
    """
    ctx, plan = _run_rounds([HUMAN, HUMAN, HUMAN])
    assert plan["meta"]["audit_status"] == "已审计"
    disk = Path(config.PLANS_DIR) / ("%s.json" % plan["plan_id"])
    assert disk.exists(), "审计结论没写回磁盘：%s" % disk
    saved = json.loads(disk.read_text(encoding="utf-8"))
    assert saved["meta"]["audit_status"] == "已审计"
    assert [r["round"] for r in saved["meta"]["audit_rounds"]] == [1, 2, 3]
    assert all(r.get("answered_by") == "human" for r in saved["meta"]["audit_rounds"])
    # 内存里的计划必须与磁盘一致（同源）
    assert json.dumps(saved["meta"]["audit_rounds"], ensure_ascii=False, sort_keys=True) == \
        json.dumps(plan["meta"]["audit_rounds"], ensure_ascii=False, sort_keys=True)


# ==================== 3. 交付物如实区分三态 ====================
def _render(decisions, draft=False):
    from pipeline.nodes.delivery import build_plan_docx
    _ctx2, plan = _run_rounds(decisions)
    return _docx_text(build_plan_docx(plan, draft=draft)), plan


def test_docx_prints_audited_only_when_three_human_rounds():
    text, plan = _render([HUMAN, HUMAN, HUMAN])
    assert "【已审计定稿】" in text
    assert "R1 通过" in text and "R3 通过" in text
    assert "人工复核通过" in text, "必须写明是人工审的，不是系统代答"
    assert "【未审计 · 待人工复审】" not in text


def test_docx_refuses_audit_when_script_answered():
    text, plan = _render([SCRIPT, SCRIPT, SCRIPT])
    assert "【已审计定稿】" not in text, "脚本代答不许印已审计定稿"
    assert "【未审计 · 待人工复审】" in text
    assert "本文件按未审计口径出具" in text
    assert "脚本代答" in text, "原因必须点名是脚本代答，而不是含糊过去"
    # 用户的复检脚本按禁语做子串匹配：这三句**一个都不许出现**
    for banned in ("已审计定稿", "R3 通过", "审计状态 已审计"):
        assert banned not in text, "命中禁语 %s（q_audit_check.py）" % banned


def test_docx_refuses_audit_when_source_missing():
    text, plan = _render([LEGACY, LEGACY, LEGACY])
    assert "【已审计定稿】" not in text
    assert "【未审计 · 待人工复审】" in text
    assert "缺少「谁答的门」记录" in text


def test_docx_refuses_audit_when_only_two_rounds():
    """真计划的形状：只有 R1、R2，一条 R3 都没有（用户审计里的定稿却印了 R3 通过）。"""
    from pipeline.nodes.delivery import build_plan_docx
    plan = _plan()
    plan["meta"]["audit_status"] = "已审计"        # 数据自称已审计 —— 不算证据
    plan["meta"]["audit_rounds"] = [
        {"round": 1, "name": "WBS 结构", "passed": True},
        {"round": 2, "name": "两版工期", "passed": True}]
    text = _docx_text(build_plan_docx(plan, draft=False))
    assert "【已审计定稿】" not in text
    assert "R3 通过" not in text, "数据里没有 R3，交付物一个字都不许编"
    assert "【未审计 · 待人工复审】" in text
    assert "自称「已审计」" in text


def test_draft_is_always_unaudited_even_with_three_human_rounds():
    """第 3 轮要看的**草案**：哪怕前两轮已经人工通过，它也必须是未审计口径。"""
    text, _plan_ = _render([HUMAN, HUMAN, HUMAN], draft=True)
    assert "【草案 · 未审计】" in text
    assert "【已审计定稿】" not in text


# ==================== 4. 看板与 Word 同一判据 ====================
def _html_text(plan):
    """`build_plan_html` 落盘后返回**路径**（它是导出节点），这里读回文本。"""
    from pipeline.nodes.delivery import build_plan_html
    return Path(build_plan_html(plan)).read_text(encoding="utf-8")


def test_dashboard_badge_three_states():
    _c1, plan_human = _run_rounds([HUMAN, HUMAN, HUMAN])
    html_ok = _html_text(plan_human)
    assert "已审计 · R1/R2/R3 真人通过" in html_ok

    _c2, plan_script = _run_rounds([SCRIPT, SCRIPT, SCRIPT])
    html_bad = _html_text(plan_script)
    assert "未审计 · 待人工复审" in html_bad
    assert "脚本代答" in html_bad
