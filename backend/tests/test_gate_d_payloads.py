# -*- coding: utf-8 -*-
"""§D 冻结数据契约回归 —— 门要带"用户能据以决策的实物内容"（工作流 B）

用户原话：
  「我最不满意的是你每道门的返回内容……WBS 门应该返回详细的 WBS 树的，为什么只返回
    摘要？你打算让用户依据什么来决定是否继续计划？」

契约见 `资料/终端界面改造规格.md` §D。本文件守住五件事：
  1. §D1 `wbs_tree`：结构照契约；限流 6 阶段 / 40 叶子（**从头截取**）；
     `truncated_leaves` 如实报被截掉的数量；叶子字段 id/name/qty/unit/duration_days/kb_activity_id；
  2. §D1 同一棵树要出现在 **R1 结构审计门** 与 **WBS 复评人工门** 两个事件里；
  3. §D2 `schedule_compare`：两版字段齐全，且**逐字段等于 ctx 里的现值**（不新算、不编数）；
  4. §D3 `draft_outline`：拿不到数据时**不报错且不带该字段**；有数据时目录与真实草案 Word 一致；
  5. 老字段（summary / highlights / issues / output_summary / context_summary）一个不丢。

运行：python -m pytest backend/tests/test_gate_d_payloads.py -q
"""

import json
import os
import shutil
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

import pytest

from pipeline.nodes.audit_gate import (
    SCHEDULE_TOP_TASKS, WBS_TREE_MAX_LEAVES, WBS_TREE_MAX_PHASES,
    DraftAuditNode, ScheduleAuditNode, WBSAuditNode,
    draft_outline_payload, schedule_compare_payload, wbs_tree_payload,
)
from pipeline.nodes.wbs_agent import WBSAgentNode


# ==================== 夹具 ====================
class _Reg(object):
    """最小登记处替身：记录注册、回放预设决策。"""

    def __init__(self, decision=None):
        self.decision = decision if decision is not None else {"passed": True}
        self.registered = []

    def register(self, key):
        self.registered.append(key)

    def wait(self, key, cancel_evt=None, timeout=None):
        return dict(self.decision)


def _emit_gate(node, ctx, decision=None):
    reg = _Reg(decision)
    events = []
    node._registry = reg
    node._run_id = "t"
    node._cancel_evt = None
    node._emit = lambda e, d: events.append((e, d))
    node.run(ctx)
    return events, reg


def _payload(events):
    evs = [d for e, d in events if e in ("param_review", "node_paused")]
    assert evs, "门必须发出交互事件"
    return evs[0]


def _big_wbs(n_phases=10, leaves_per_phase=42):
    """415 条量级的 WBS（阶段数/叶子数都可调），用于验证限流。"""
    phases = []
    for p in range(1, n_phases + 1):
        leaves = [{"id": "%d.1.%d" % (p, i), "name": "阶段%d叶子%d" % (p, i),
                   "quantity": round(i * 1.5, 2), "unit": "m³",
                   "duration_days": i, "kb_activity_id": "ACT_%d" % p}
                  for i in range(1, leaves_per_phase + 1)]
        phases.append({"phase": "阶段%d" % p,
                       "work_packages": [{"id": "%d.1" % p, "name": "工作包%d" % p,
                                          "sub_packages": leaves}]})
    return {"phases": phases}


def _flat_leaves(tree):
    return [l for ph in tree["phases"] for wp in ph["work_packages"] for l in wp["leaves"]]


# ==================== 1. §D1 wbs_tree：结构与限流 ====================
def test_d1_contract_shape_exactly_as_frozen():
    tree = wbs_tree_payload(_big_wbs(n_phases=2, leaves_per_phase=3))
    assert set(tree) == {"counts", "phases", "shown_leaves", "truncated_leaves"}
    assert set(tree["counts"]) == {"phases", "work_packages", "leaves"}
    assert set(tree["phases"][0]) == {"phase", "work_packages"}
    wp = tree["phases"][0]["work_packages"][0]
    assert set(wp) == {"id", "name", "leaves"}
    assert set(wp["leaves"][0]) == {"id", "name", "qty", "unit", "duration_days",
                                    "kb_activity_id"}
    assert tree["counts"] == {"phases": 2, "work_packages": 2, "leaves": 6}


def test_d1_leaf_values_come_from_wbs_verbatim():
    leaf = wbs_tree_payload(_big_wbs(n_phases=1, leaves_per_phase=2))["phases"][0] \
        ["work_packages"][0]["leaves"][1]
    assert leaf["id"] == "1.1.2" and leaf["name"] == "阶段1叶子2"
    assert leaf["qty"] == 3.0 and leaf["unit"] == "m³" and leaf["duration_days"] == 2
    assert leaf["kb_activity_id"] == "ACT_1"


def test_d1_leaf_without_kb_activity_omits_the_key_not_null():
    tree = wbs_tree_payload({"phases": [{"phase": "P", "work_packages": [
        {"id": "1.1", "name": "w", "sub_packages": [
            {"id": "1.1.1", "name": "n", "quantity": 2, "unit": "t",
             "duration_days": 4}]}]}]})
    leaf = tree["phases"][0]["work_packages"][0]["leaves"][0]
    assert "kb_activity_id" not in leaf, "没有锚定活动就不该塞一个 null 给界面"
    assert leaf["qty"] == 2 and leaf["duration_days"] == 4


def test_d1_leaf_cap_40_and_truncated_count():
    """415 条量级的树：只许带 40 条叶子，其余如实报数（不许塞进 SSE 帧）。"""
    wbs = _big_wbs(n_phases=10, leaves_per_phase=42)          # 420 条叶子
    tree = wbs_tree_payload(wbs)
    assert tree["counts"]["leaves"] == 420 and tree["counts"]["phases"] == 10
    assert len(_flat_leaves(tree)) == tree["shown_leaves"] == WBS_TREE_MAX_LEAVES == 40
    assert tree["truncated_leaves"] == 420 - 40 == 380


def test_d1_truncation_takes_from_the_head_in_order():
    """截断必须**从头按顺序取** —— 用户看到的是计划的开始部分。"""
    tree = wbs_tree_payload(_big_wbs(n_phases=10, leaves_per_phase=42))
    leaves = _flat_leaves(tree)
    assert leaves[0]["id"] == "1.1.1", "第一条必须是整棵树的第一条"
    assert leaves[-1]["id"] == "1.1.40", "取满 40 条即止（本阶段内的顺序不能乱）"
    assert [l["id"] for l in leaves] == ["1.1.%d" % i for i in range(1, 41)]


def test_d1_phase_cap_6_when_leaves_are_few():
    wbs = _big_wbs(n_phases=10, leaves_per_phase=1)           # 10 阶段 × 1 叶子
    tree = wbs_tree_payload(wbs)
    assert len(tree["phases"]) == WBS_TREE_MAX_PHASES == 6
    assert tree["shown_leaves"] == 6 and tree["truncated_leaves"] == 4
    assert [p["phase"] for p in tree["phases"]] == ["阶段%d" % i for i in range(1, 7)]


def test_d1_exact_40_leaves_is_not_truncated():
    tree = wbs_tree_payload(_big_wbs(n_phases=4, leaves_per_phase=10))
    assert tree["shown_leaves"] == 40 and tree["truncated_leaves"] == 0
    assert len(tree["phases"]) == 4


def test_d1_empty_phase_is_kept_in_the_tree():
    """空阶段（无叶子）在结构审计里必须看得见，不能被悄悄丢掉。"""
    wbs = {"phases": [{"phase": "空阶段", "work_packages": []},
                      {"phase": "有叶子", "work_packages": [{"id": "2.1", "name": "wp",
                        "sub_packages": [{"id": "2.1.1", "name": "n", "quantity": 1,
                                          "unit": "项", "duration_days": 1}]}]}]}
    tree = wbs_tree_payload(wbs)
    assert [p["phase"] for p in tree["phases"]] == ["空阶段", "有叶子"]
    assert tree["counts"] == {"phases": 2, "work_packages": 1, "leaves": 1}


def test_d1_non_finite_numbers_never_leak_into_the_frame():
    """NaN / inf 会让 SSE 帧变成**非法 JSON** —— 一律退回 0，绝不漏出去。"""
    tree = wbs_tree_payload({"phases": [{"phase": "P", "work_packages": [
        {"id": "1.1", "name": "w", "sub_packages": [
            {"id": "1.1.1", "name": "n", "quantity": float("nan"), "unit": "m³"},
            {"id": "1.1.2", "name": "n2", "quantity": float("inf"),
             "duration_days": float("-inf")}]}]}]})
    leaves = tree["phases"][0]["work_packages"][0]["leaves"]
    assert leaves[0]["qty"] == 0 and leaves[0]["duration_days"] == 0
    assert leaves[1]["qty"] == 0 and leaves[1]["duration_days"] == 0
    json.dumps(tree, allow_nan=False)          # 严格 JSON 必须能序列化


def test_d1_unavailable_data_returns_none_never_raises():
    for bad in ({}, None, [], "x", {"phases": None}, {"phases": []}, 42,
                {"phases": [None]}, {"phases": [{"work_packages": "烂数据"}]}):
        wbs_tree_payload(bad)          # 不抛异常即可


def test_r1_gate_event_carries_the_tree_and_keeps_old_fields():
    ctx = {"wbs": _big_wbs(n_phases=8, leaves_per_phase=60)}
    events, _reg = _emit_gate(WBSAuditNode(), ctx, {"passed": True})
    payload = _payload(events)
    assert payload["purpose"] == "audit" and payload["round"] == 1
    # 老字段一个不少（界面有兼容路径）
    assert payload["summary"] and payload["highlights"] and payload["next_hint"]
    assert payload["highlights"]["leaves"] == 480
    tree = payload["wbs_tree"]
    assert tree["shown_leaves"] == 40 and tree["truncated_leaves"] == 440
    assert tree["counts"]["leaves"] == 480


def test_r1_gate_without_wbs_omits_the_field_but_still_asks():
    events, _reg = _emit_gate(WBSAuditNode(), {"wbs": "烂数据"}, {"passed": True})
    payload = _payload(events)
    assert "wbs_tree" not in payload, "取不到就不带该字段（界面优雅退化）"
    assert payload["summary"], "没有树也必须有摘要，用户不能被卡住"
    assert payload["highlights"] == {}, "老字段照旧给出（此处本来就取不到 WBS）"


# ==================== 2. §D1 WBS 复评人工门 ====================
def _review_gate_events(wbs):
    node = WBSAgentNode()
    node._registry = _Reg({"action": "continue"})       # 视为打 Y
    node._run_id = "t"
    node._cancel_evt = None
    events = []
    node._emit = lambda e, d: events.append((e, d))
    node._human_gate([{"severity": "HIGH", "dimension": "层数", "target": "1.1",
                       "finding": "只统计了一层", "suggestion": "×层数"}], wbs, {})
    return [d for e, d in events if e == "node_paused"]


def test_wbs_review_gate_carries_the_same_tree_and_keeps_old_fields():
    wbs = _big_wbs(n_phases=10, leaves_per_phase=42)
    payload = _review_gate_events(wbs)[0]
    assert payload["issues"] and payload["output_summary"] and payload["context_summary"], \
        "老的 issues / output_summary / context_summary 必须保留（只增不改）"
    assert payload["wbs_tree"] == wbs_tree_payload(wbs), \
        "复评门与 R1 门给用户的必须是同一棵树"


def test_wbs_review_gate_without_tree_still_asks():
    payload = _review_gate_events(None)[0]
    assert "wbs_tree" not in payload
    assert payload["issues"] and payload["output_summary"]


# ==================== 3. §D2 schedule_compare ====================
def _sv_ctx():
    """两版排程结果 + 定额口径警告，全部是**预先写死的哨兵值**。"""
    def rows(shift):
        return [
            {"task_id": "1.1.1", "es": 0 + shift, "ef": 3 + shift,
             "crew": {"钢筋工": 45}, "capped": False},
            {"task_id": "1.1.2", "es": 3 + shift, "ef": 20 + shift,
             "crew": {"模板工": 12}, "capped": True},
            {"task_id": "1.1.3", "es": 20 + shift, "ef": 30 + shift,
             "crew": {"瓦工": 8}, "capped": False},
        ]
    return {
        "wbs": _big_wbs(n_phases=1, leaves_per_phase=3),      # 名字只在 WBS 里有
        "schedule_versions": {
            "theory_min": {"total_duration_days": 616, "schedule": rows(0),
                           "critical_path": ["1.1.1", "1.1.2"], "peak_labor": 70},
            "resource_ok": {"total_duration_days": 616, "schedule": rows(0),
                            "critical_path": ["1.1.1", "1.1.2"], "peak_labor": 70},
            "compare": {"delta_days": 0},
            "warnings": ["定额工日需求合计 12,969 人日（按工种：架子工 225、钢筋工 7,091）"
                         "——这是定额算出来的客观量，工期 = 工日需求 ÷ 班组人数"],
        },
    }


def test_d2_contract_shape_and_values_are_taken_from_ctx():
    ctx = _sv_ctx()
    cmp_ = schedule_compare_payload(ctx)
    assert set(cmp_) == {"theory_min", "resource_ok", "limit_note", "top_tasks", "labor_note"}
    for key in ("theory_min", "resource_ok"):
        assert set(cmp_[key]) == {"total_duration_days", "leaves", "critical", "peak_labor"}
    # 逐字段等于 ctx 里的现值
    assert cmp_["theory_min"] == {"total_duration_days": 616, "leaves": 3,
                                 "critical": 2, "peak_labor": 70}
    assert cmp_["resource_ok"] == cmp_["theory_min"]
    assert cmp_["limit_note"] == "用户未给资源限额，两版一致"
    # labor_note 原样取自 ctx 的口径警告（砍掉"——"之后的解释段）
    assert cmp_["labor_note"] == ("定额工日需求合计 12,969 人日"
                                  "（按工种：架子工 225、钢筋工 7,091）")


def test_d2_is_not_recomputed_from_rows():
    """ctx 说 999 天就报 999 —— 门里绝不许自己用 ef/es 重算一遍工期。"""
    ctx = _sv_ctx()
    ctx["schedule_versions"]["theory_min"]["total_duration_days"] = 999
    ctx["schedule_versions"]["theory_min"]["peak_labor"] = 12345
    cmp_ = schedule_compare_payload(ctx)
    assert cmp_["theory_min"]["total_duration_days"] == 999
    assert cmp_["theory_min"]["peak_labor"] == 12345
    assert cmp_["theory_min"]["leaves"] == 3, "叶子数 = 该版排程行数（取自 ctx）"


def test_d2_top_tasks_sorted_by_days_desc_and_capped_at_6():
    ctx = _sv_ctx()
    ctx["schedule_versions"]["resource_ok"]["schedule"] = [
        {"task_id": "T%d" % i, "es": 0, "ef": i, "crew": {}, "capped": False}
        for i in range(1, 10)]
    cmp_ = schedule_compare_payload(ctx)
    assert len(cmp_["top_tasks"]) == SCHEDULE_TOP_TASKS == 6
    days = [t["days"] for t in cmp_["top_tasks"]]
    assert days == sorted(days, reverse=True) == [9, 8, 7, 6, 5, 4]
    assert set(cmp_["top_tasks"][0]) == {"task_id", "task_name", "days", "es", "ef",
                                         "crew", "capped"}


def test_d2_task_name_comes_from_wbs_and_days_from_es_ef():
    cmp_ = schedule_compare_payload(_sv_ctx())
    top = {t["task_id"]: t for t in cmp_["top_tasks"]}
    assert top["1.1.2"]["task_name"] == "阶段1叶子2"          # 名字来自 ctx["wbs"]
    assert top["1.1.2"]["days"] == 17 and top["1.1.2"]["crew"] == {"模板工": 12}
    assert top["1.1.2"]["capped"] is True
    assert top["1.1.1"]["capped"] is False
    assert [t["task_id"] for t in cmp_["top_tasks"]] == ["1.1.2", "1.1.3", "1.1.1"], \
        "按工期降序：17 / 10 / 3 天"


def test_d2_limit_note_follows_user_limits_in_ctx():
    ctx = _sv_ctx()
    ctx["schedule_versions"]["resource_ok"] = {
        "total_duration_days": 743, "schedule": [], "critical_path": [], "peak_labor": 44}
    ctx["schedule_versions"]["compare"] = {"delta_days": 127}
    ctx["boundary_conditions"] = {"labor": {"by_trade": [{"trade": "钢筋工", "quantity": 8}]}}
    cmp_ = schedule_compare_payload(ctx)
    assert "资源限额" in cmp_["limit_note"] and "127" in cmp_["limit_note"]
    assert cmp_["resource_ok"]["total_duration_days"] == 743


def test_d2_unavailable_data_returns_none_or_omits_labor_note():
    for bad in ({}, {"schedule_versions": None}, {"schedule_versions": "烂数据"},
                {"schedule_versions": []}):
        assert schedule_compare_payload(bad) is None
    # 有排程但没有那条口径警告 → 不带 labor_note，其余照常
    ctx = _sv_ctx()
    ctx["schedule_versions"]["warnings"] = ["别的不相关警告"]
    cmp_ = schedule_compare_payload(ctx)
    assert "labor_note" not in cmp_ and cmp_["top_tasks"]


def test_r2_gate_event_carries_compare_and_keeps_old_fields():
    events, _reg = _emit_gate(ScheduleAuditNode(), _sv_ctx(), {"passed": True})
    payload = _payload(events)
    assert payload["round"] == 2
    assert payload["summary"] and payload["highlights"]
    assert payload["highlights"]["theory_min_days"] == 616, "老 highlights 不变"
    assert payload["schedule_compare"]["theory_min"]["total_duration_days"] == 616
    assert len(payload["schedule_compare"]["top_tasks"]) == 3


# ==================== 4. §D3 draft_outline ====================
def _plan_for_outline(risks=True):
    return {
        "plan_id": "outline_test",
        "overview": {"project_name": "目录测试", "total_duration_days": 30,
                     "planned_start_date": "2026-03-01", "planned_end_date": "2026-03-31",
                     "critical_path_length": 2},
        "meta": {"audit_status": "未审计", "caliber_note": "单栋项目，按整栋口径编制",
                 "audit_rounds": [], "audit_comments": []},
        "wbs": {"phases": [{"phase": "主体", "work_packages": [
            {"id": "5.1", "name": "wp", "sub_packages": [
                {"id": "5.1.1.1", "name": "钢筋", "quantity": 10, "unit": "t",
                 "duration_days": 10}]}]}]},
        "dependencies": [],
        "all_tasks_schedule": [
            {"task_id": "5.1.1.1", "task_name": "钢筋",
             "start_date": "2026-03-01", "finish_date": "2026-03-11",
             "duration_days": 10, "assigned_resources": {"钢筋工": 5}},
            {"task_id": "5.1.1.2", "task_name": "模板",
             "start_date": "2026-03-05", "finish_date": "2026-03-15",
             "duration_days": 10, "assigned_resources": {"模板工": 4, "塔吊": 1}},
        ],
        "key_milestones": [{"name": "开工", "date": "2026-03-01", "description": "开工"}],
        "critical_path_tasks": [],
        "resource_plan": {"peak_manpower": 5, "total_manpower_days": 50.0,
                          "equipment_peak": {"塔吊": 1}, "material_summary": []},
        "risks": ([{"risk_name": "雨季施工", "mitigation": "提前排水"}] if risks else []),
        "report": "# 监督报告\n\n- 第一条\n- 第二条",
    }


def test_d3_contract_shape():
    ctx = {"plan_json": _plan_for_outline(),
           "norm_coverage": {"total": 415, "bound": 278, "bound_pct": 67.0}}
    out = draft_outline_payload(ctx)
    assert set(out) == {"sections", "tables", "figures", "coverage", "note"}
    assert set(out["sections"][0]) == {"title", "lines"}
    assert set(out["tables"][0]) == {"title", "rows"}
    assert all(isinstance(f, str) for f in out["figures"])
    assert out["coverage"] == "定额口径覆盖率 67.0%（278/415）"
    assert out["note"] == "草案未审计、不含图表；通过后才出定稿与看板"
    assert all(s["lines"] >= 0 for s in out["sections"])


def test_d3_coverage_omitted_when_no_norm_coverage():
    out = draft_outline_payload({"plan_json": _plan_for_outline()})
    assert "coverage" not in out, "拿不到覆盖率就不带该字段"
    assert out["sections"] and out["tables"]


def test_d3_risk_section_follows_the_plan():
    with_risks = [s["title"] for s in draft_outline_payload(
        {"plan_json": _plan_for_outline(risks=True)})["sections"]]
    without = [s["title"] for s in draft_outline_payload(
        {"plan_json": _plan_for_outline(risks=False)})["sections"]]
    assert "七、主要风险与应对" in with_risks
    assert "七、主要风险与应对" not in without
    assert "八、施工监督报告" in without


def test_d3_unavailable_data_returns_none_never_raises():
    for bad in ({}, None, {"plan_json": None}, {"plan_json": {}}, {"plan_json": []},
                {"plan_json": "烂数据"}, {"plan_json": {"meta": "x"}}):
        draft_outline_payload(bad)     # 不抛异常即可


def test_r3_gate_without_plan_json_omits_the_field_and_does_not_crash():
    """③ 拿不到数据时**不报错且不带该字段**（界面优雅退化）。"""
    ctx = {"wbs": {"phases": []}}                      # 没有 plan_json
    events, _reg = _emit_gate(DraftAuditNode(), ctx, {"passed": True})
    payload = _payload(events)
    assert "draft_outline" not in payload
    assert payload["round"] == 3 and payload["summary"] and payload["next_hint"]


def test_r3_gate_event_carries_outline_and_keeps_old_fields():
    ctx = {"plan_json": _plan_for_outline(), "wbs": _plan_for_outline()["wbs"],
           "artifacts": {"docx_draft": "D:/tmp/草案.docx"},
           "norm_coverage": {"total": 100, "bound": 67, "bound_pct": 67.0}}
    events, _reg = _emit_gate(DraftAuditNode(), ctx, {"passed": True})
    payload = _payload(events)
    assert payload["summary"] and payload["highlights"]
    assert payload["highlights"]["leaves"] == 1
    assert payload["draft_outline"]["tables"], "R3 门要带草案目录"
    assert any(f == "甘特图" for f in payload["draft_outline"]["figures"])


# ==================== 5. 目录必须与真实草案 Word 对得上 ====================
@pytest.fixture()
def docx_dir(monkeypatch):
    from pipeline import config

    root = BACKEND / "_test_tmp" / ("d_payload_p%d" % os.getpid())
    shutil.rmtree(str(root), ignore_errors=True)
    root.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(config, "DELIVERABLES_DIR", root, raising=False)
    yield root
    shutil.rmtree(str(root), ignore_errors=True)


def test_d3_sections_match_the_real_draft_docx(docx_dir):
    """目录是**镜像**草案 Word 的：章节标题与甘特行数必须与真产物逐一对上。

    没有这条护栏，`build_plan_docx` 改了章节而门里的目录没跟着改，用户看到的目录
    就会与手上的草案不是一份东西。
    """
    from docx import Document

    from pipeline.nodes.delivery import build_plan_docx

    plan = _plan_for_outline()
    path = build_plan_docx(plan, draft=True)
    doc = Document(path)
    headings = [p.text for p in doc.paragraphs
                if (p.style.name or "").startswith("Heading")]
    outline = draft_outline_payload({"plan_json": plan})
    assert [s["title"] for s in outline["sections"]] == headings, \
        "目录章节必须与草案 Word 的章节标题一致"

    gantt_tables = [t for t in doc.tables
                    if t.rows and t.rows[0].cells[0].text == "任务 ID"]
    assert gantt_tables, "草案里应有甘特表"
    rows_in_doc = len(gantt_tables[0].rows) - 1
    gantt_section = [t for t in outline["tables"] if t["title"].startswith("横道图")][0]
    assert gantt_section["rows"] == rows_in_doc == len(plan["all_tasks_schedule"])


# ==================== 6. 门不许因为取数失败而崩 ====================
@pytest.mark.parametrize("node_cls", [WBSAuditNode, ScheduleAuditNode, DraftAuditNode])
def test_gates_never_crash_and_never_emit_null_fields(node_cls):
    for ctx in ({}, {"wbs": 123}, {"schedule_versions": "x", "wbs": [1, 2]},
                {"plan_json": []}, {"wbs": {"phases": [None]},
                                    "schedule_versions": {"theory_min": None},
                                    "plan_json": {"meta": None}},
                {"wbs": {"phases": [{"work_packages": [{"sub_packages": "烂"}]}]}},
                None):
        events, _reg = _emit_gate(node_cls(), ctx, {"passed": True})
        payload = _payload(events)
        assert payload["summary"], "摘要必须有（哪怕是空 ctx）"
        for key in ("wbs_tree", "schedule_compare", "draft_outline"):
            # 取不到就干脆不带这个键 —— 绝不能带一个 None 给界面
            assert payload.get(key, "（缺）") is not None, "%s 不许是 None" % key
