# -*- coding: utf-8 -*-
"""只读：dump 技术说明文档的章节大纲 + 定位「十二、总结」+ 表格清单。

用法：python devtools/_dump_doc_outline.py "建策BuildPlan_技术说明文档_v3.0.docx"
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from docx import Document            # noqa: E402
from docx.table import Table         # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402


def iter_block_items(parent):
    from docx.oxml.ns import qn
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


def main():
    path = Path(sys.argv[1] if len(sys.argv) > 1 else "建策BuildPlan_技术说明文档_v3.0.docx")
    doc = Document(str(path))
    print("文件: %s" % path)
    print("段落数 %d / 表格数 %d" % (len(doc.paragraphs), len(doc.tables)))
    print("-" * 78)
    for i, item in enumerate(iter_block_items(doc)):
        if isinstance(item, Table):
            head = " | ".join((c.text or "").strip()[:12] for c in item.rows[0].cells)
            print("[%4d] TABLE  %dx%d  %s" % (i, len(item.rows), len(item.columns), head[:70]))
            continue
        t = (item.text or "").strip()
        st = item.style.name if item.style is not None else ""
        if not t:
            continue
        if st.lower().startswith("heading") or st.startswith("标题") or t[:3] in (
                "十二", "十一", "十三", "十四", "一、", "二、", "三、", "四、", "五、",
                "六、", "七、", "八、", "九、", "十、"):
            print("[%4d] %-14s %s" % (i, st, t[:70]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
