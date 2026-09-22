# -*- coding: utf-8 -*-
"""把《建策BuildPlan_产品说明文档_补充附录.md》的内容插进技术说明文档 docx。

插入位置：现有「十二、总结」**之前**；插入后原「十二、总结」改名为「十四、总结」。
新章节：十二、竞品对比分析 / 十三、产品不足与发展展望 / 补充数据（实测输出）。

设计要点（为什么要这么写）：
  · 该文档**所有段落都是 Normal 样式**、靠手工格式区分层级（无 Heading），所以本脚本
    从文档里**现取三个格式模板**（章节标题 / 小节标题 / 正文）复制段落属性，而不是
    自己造样式 —— 否则插进去的东西一眼就看出是"外来的"。
  · 表格用文档里现有表格的 style 新建，再做 XML 级移动（`addprevious`）保证落点。
  · **幂等**：文档里已出现「十二、竞品对比分析」就什么都不做。
  · **跳过附录里"怎么插"的说明块**（使用说明 / 建议插入位置 / 新增章节内容 三个标题、
    以及末尾"文档版本/生成时间/用途"三行元信息）——那些是给操作者的指示，不是正文。

用法：
    python devtools\\append_doc_appendix.py                     # dry-run，只打印计划
    python devtools\\append_doc_appendix.py --apply             # 写盘（自动备份）
    python devtools\\append_doc_appendix.py --apply --docx <路径> --md <路径>
"""
import argparse
import copy
import datetime
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from docx import Document      # noqa: E402

DEFAULT_DOCX = ROOT / "建策BuildPlan_技术说明文档_v3.0.docx"
DEFAULT_MD = ROOT / "建策BuildPlan_产品说明文档_补充附录.md"
ANCHOR = "十二、总结"
RENAMED = "十四、总结"
NEW_MARK = "十二、竞品对比分析"
SKIP_TITLES = {"使用说明", "建议插入位置", "新增章节内容"}
SKIP_TAIL = ("文档版本", "生成时间", "用途：")


def clean(text):
    t = text.strip()
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = t.replace("`", "")
    t = re.sub(r"^#+\s*", "", t)
    return t.strip()


def parse_md(path):
    """→ [{kind: 'title'|'sub'|'body'|'bullet'|'table'|'sep', ...}]"""
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):                       # 表格
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [clean(c) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                blocks.append({"kind": "table", "rows": rows})
            continue
        if s.startswith(">"):                       # 引用块 = 操作说明 → 跳过
            i += 1
            continue
        if s.startswith("---"):                     # 分隔线
            blocks.append({"kind": "sep"})
            i += 1
            continue
        text = clean(s)
        if any(text.startswith(k) for k in SKIP_TAIL):
            i += 1
            continue
        if s.startswith("#### "):
            blocks.append({"kind": "sub", "text": text, "level": 4})
        elif s.startswith("### "):
            # 三级标题里「十二、竞品对比分析」这种带中文序号的是**章**，
            # 其余（「测试项目：…」「输出文件」）是节。
            kind = "title" if re.match(r"^十[一二三四五六七八九十]+、", text) else "sub"
            blocks.append({"kind": kind, "text": text, "level": 3})
        elif s.startswith("## "):
            blocks.append({"kind": "title", "text": text, "level": 2})
        elif s.startswith("# "):
            i += 1                                  # 文档大标题，不要
        elif s.startswith("- "):
            blocks.append({"kind": "bullet", "text": clean(s[2:])})
        else:
            blocks.append({"kind": "body", "text": text})
        i += 1
    # 只保留**正文内容**：从第一个「十二、…」章节标题开始，前面全是"怎么插"的说明
    # （含围栏代码块、引用块、三个说明小标题）——不靠逐条认词，避免漏网。
    start = next((n for n, b in enumerate(blocks)
                  if b["kind"] == "title" and b["text"].startswith("十二、")), None)
    if start is None:
        return []
    return blocks[start:]


def find_by_text(doc, want):
    for p in doc.paragraphs:
        if (p.text or "").strip() == want:
            return p
    return None


def insert_before(anchor, doc, text, template, bold=None):
    """在 anchor 前插一个段落，文本格式复制 template（可覆盖 run 加粗）。"""
    new_p = anchor.insert_paragraph_before(text)
    try:
        new_p.style = template.style
    except Exception:
        pass
    src_runs = template.runs
    if src_runs:
        rpr = src_runs[0]._element.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
        if rpr is not None:
            for run in new_p.runs:
                run._element.insert(0, copy.deepcopy(rpr))
        if bold is not None:
            for run in new_p.runs:
                run.bold = bold
    ppr = template._p.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
    if ppr is not None and new_p._p.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr") is None:
        new_p._p.insert(0, copy.deepcopy(ppr))
    return new_p


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--docx", default=str(DEFAULT_DOCX))
    ap.add_argument("--md", default=str(DEFAULT_MD))
    args = ap.parse_args()

    docx_path = Path(args.docx)
    if not docx_path.exists():
        print("找不到 docx：%s" % docx_path)
        return 2
    md_path = Path(args.md)
    if not md_path.exists():
        print("找不到 md：%s" % md_path)
        return 2

    try:
        open(docx_path, "r+b").close()
    except PermissionError:
        print("❌ docx 被占用（很可能是 WPS/Word 打开着）→ 请先关闭再跑。")
        return 3

    doc = Document(str(docx_path))
    if find_by_text(doc, NEW_MARK) is not None:
        print("幂等：文档里已有「%s」，什么都不做。" % NEW_MARK)
        return 0

    anchor = find_by_text(doc, ANCHOR)
    if anchor is None:
        print("找不到锚点段落「%s」——请人工确认文档结构。" % ANCHOR)
        return 4

    tpl_title = find_by_text(doc, "十一、部署与运维")
    tpl_sub = find_by_text(doc, "11.1 运行环境")
    tpl_body = None
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if len(t) > 60 and p.style is not None and p.style.name == "Normal":
            tpl_body = p
            break
    if tpl_title is None or tpl_sub is None or tpl_body is None:
        print("模板段落取不到：title=%s sub=%s body=%s"
              % (tpl_title is not None, tpl_sub is not None, tpl_body is not None))
        return 5

    blocks = parse_md(md_path)
    n_t = sum(1 for b in blocks if b["kind"] == "title")
    n_s = sum(1 for b in blocks if b["kind"] == "sub")
    n_b = sum(1 for b in blocks if b["kind"] in ("body", "bullet"))
    n_tb = sum(1 for b in blocks if b["kind"] == "table")
    trows = sum(len(b["rows"]) for b in blocks if b["kind"] == "table")
    print("解析 md：章节标题 %d / 小节标题 %d / 段落 %d / 表格 %d（共 %d 行）"
          % (n_t, n_s, n_b, n_tb, trows))
    for b in blocks:
        if b["kind"] == "title":
            print("   [章节] %s" % b["text"])
        elif b["kind"] == "sub":
            print("      [小节] %s" % b["text"])
        elif b["kind"] == "table":
            print("      [表格] %dx%d  %s" % (len(b["rows"]), len(b["rows"][0]),
                                              " | ".join(b["rows"][0])[:60]))
    old_tables = len(doc.tables)
    old_paras = len(doc.paragraphs)
    print("插入前：段落 %d / 表格 %d；锚点后原「%s」将改名为「%s」"
          % (old_paras, old_tables, ANCHOR, RENAMED))

    if not args.apply:
        print("\n（dry-run，未写盘。加 --apply 执行。）")
        return 0

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = docx_path.with_name(docx_path.stem + ".bak_%s_appendix.docx" % ts)
    shutil.copy2(docx_path, backup)
    print("已备份 → %s" % backup.name)

    tpl_table_style = doc.tables[-1].style if doc.tables and doc.tables[-1].style else None
    made = 0
    for b in blocks:
        if b["kind"] == "sep":
            continue
        if b["kind"] == "title":
            insert_before(anchor, doc, b["text"], tpl_title, bold=True)
        elif b["kind"] == "sub":
            insert_before(anchor, doc, b["text"], tpl_sub, bold=True)
        elif b["kind"] == "bullet":
            insert_before(anchor, doc, "•  " + b["text"], tpl_body, bold=False)
        elif b["kind"] == "body":
            insert_before(anchor, doc, b["text"], tpl_body, bold=False)
        elif b["kind"] == "table":
            rows = b["rows"]
            tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
            if tpl_table_style is not None:
                try:
                    tbl.style = tpl_table_style
                except Exception:
                    pass
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    tbl.cell(r, c).text = cell
            anchor._p.addprevious(tbl._tbl)
            made += 1

    anchor.text = RENAMED
    doc.save(str(docx_path))
    doc2 = Document(str(docx_path))
    has = find_by_text(doc2, NEW_MARK) is not None
    print("已写盘：段落 %d → %d，表格 %d → %d，新表格 %d 个，锚点已改名：%s"
          % (old_paras, len(doc2.paragraphs), old_tables, len(doc2.tables), made,
             find_by_text(doc2, RENAMED) is not None))
    print("校验：文档含「%s」= %s；仍含旧「%s」= %s" % (
        NEW_MARK, has, ANCHOR, find_by_text(doc2, ANCHOR) is not None))
    return 0 if has else 1


if __name__ == "__main__":
    sys.exit(main())
