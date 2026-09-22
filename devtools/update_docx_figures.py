# -*- coding: utf-8 -*-
"""E：把重画好的两张图贴回 docx，并把显示宽度放大到**满版心**。

为什么要单独写：python-docx 不支持替换已有图片，也不便改显示尺寸。
这里直接做 zip 级重写 —— 除两个 media 文件与 document.xml 里的两处尺寸外，
其余条目**逐字节保持原样**，docx 的其它内容零风险。

用法：
    python tools/update_docx_figures.py --dry-run
    python tools/update_docx_figures.py
"""
import argparse
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "建策BuildPlan_技术说明文档_v3.0.docx"


def _pick(*names):
    """按优先级挑图源：**优先 SVG 版**（用户拍板"图就用现在的 SVG"），
    没有 SVG 版时才退回旧的 Pillow 版 —— 这样图 1 还没重画完也能先跑。
    """
    for n in names:
        p = ROOT / "docs" / n
        if p.exists():
            return p
    return ROOT / "docs" / names[-1]


IMG1 = _pick("图1_技术栈图_SVG版.png", "图1_技术栈图.png")
IMG2 = _pick("图2_系统架构图_SVG版.png", "图2_系统架构图.png")

EMU_PER_CM = 360000.0
TARGET_W_CM = 15.24          # 满版心（页面 21.59 − 左右各 3.18）


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--width-cm", type=float, default=TARGET_W_CM)
    args = ap.parse_args()

    for p in (DOC, IMG1, IMG2):
        if not p.exists():
            raise SystemExit("缺少文件：%s" % p)

    from PIL import Image
    sizes = {}
    for name, p in (("image1", IMG1), ("image2", IMG2)):
        im = Image.open(p)
        sizes[name] = im.size
        print("  %s  %s  %.0f KB  → 显示 %.2f x %.2f cm"
              % (p.name, im.size, p.stat().st_size / 1024,
                 args.width_cm, args.width_cm * im.size[1] / im.size[0]))

    with zipfile.ZipFile(DOC) as z:
        names = z.namelist()
        xml = z.read("word/document.xml").decode("utf-8")
    media = [n for n in names if n.startswith("word/media/")]
    print("  文档内 media：%s" % media)

    # ---- 计算新的 extent（保持纵横比）----
    w_emu = int(round(args.width_cm * EMU_PER_CM))
    new = []
    order = ["image1", "image2"] if len(media) == 2 else []
    for i, key in enumerate(order):
        iw, ih = sizes[key]
        h_emu = int(round(w_emu * ih / iw))
        new.append((w_emu, h_emu))
        print("  extent[%d] → cx=%d cy=%d (%.2f x %.2f cm)"
              % (i + 1, w_emu, h_emu, w_emu / EMU_PER_CM, h_emu / EMU_PER_CM))

    # ---- 替换 document.xml 里的两处 wp:extent 与两处 a:ext ----
    def _sub_all(pattern, repl_list, text):
        idx = [0]

        def _r(m):
            i = idx[0]
            idx[0] += 1
            if i >= len(repl_list):
                return m.group(0)
            return m.group(1) + str(repl_list[i][0]) + m.group(3) + str(repl_list[i][1]) + m.group(5)

        return re.sub(pattern, _r, text)

    wp_pat = r'(<wp:extent cx=")(\d+)(" cy=")(\d+)(")'
    a_pat = r'(<a:ext cx=")(\d+)(" cy=")(\d+)(")'
    before_wp = re.findall(wp_pat, xml)
    before_a = re.findall(a_pat, xml)
    print("  document.xml 里 wp:extent %d 处、a:ext %d 处" % (len(before_wp), len(before_a)))

    xml_new = _sub_all(wp_pat, new, xml)
    xml_new = _sub_all(a_pat, new, xml_new)
    print("  改后 wp:extent: %s" % re.findall(wp_pat, xml_new))
    print("  改后 a:ext    : %s" % re.findall(a_pat, xml_new))

    if args.dry_run:
        print("\n[dry-run] 未写盘。")
        return

    bak = DOC.with_name("建策BuildPlan_技术说明文档_v3.0.bak_%s.docx"
                        % datetime.now().strftime("%Y%m%d_%H%M%S"))
    shutil.copy2(DOC, bak)
    print("\n已备份 -> %s" % bak.name)

    tmp = DOC.with_suffix(".tmp.docx")
    with zipfile.ZipFile(DOC) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = xml_new.encode("utf-8")
            elif item.filename == "word/media/image1.png":
                data = IMG1.read_bytes()
            elif item.filename == "word/media/image2.png":
                data = IMG2.read_bytes()
            zout.writestr(item, data)
    tmp.replace(DOC)
    print("已写回 -> %s（%.0f KB）" % (DOC.name, DOC.stat().st_size / 1024))

    # ---- 自检：能重新打开、且尺寸已改 ----
    import docx
    d = docx.Document(str(DOC))
    with zipfile.ZipFile(DOC) as z:
        x = z.read("word/document.xml").decode("utf-8")
        m1 = z.read("word/media/image1.png")
        m2 = z.read("word/media/image2.png")
    print("自检：段落 %d、表格 %d" % (len(d.paragraphs), len(d.tables)))
    print("自检：extent = %s" % re.findall(wp_pat, x))
    print("自检：新图字节相同 = %s / %s"
          % (m1 == IMG1.read_bytes(), m2 == IMG2.read_bytes()))


if __name__ == "__main__":
    main()
