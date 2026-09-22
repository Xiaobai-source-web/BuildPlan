# -*- coding: utf-8 -*-
"""把技术说明文档 §5.5「展示粒度」从**两个维度**改成**六个组合选项**。

背景（第 33 轮，用户原话）：「粒度选择不要做成"X+X"两轴选项，直接给用户六个选项」。
代码已改（plan_level._picker_payload 出六个组合号），文档若不改就会与实现相反。

只改文字，不动版式：每个目标段落/单元格都把新文字写进 run[0]、清空其余 run
（run[0] 原本就承载全部可见文字，见改造前的探查输出）。

用法：python devtools/update_doc_granularity.py
"""
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "backend"))

from docx import Document  # noqa: E402

DOC = os.path.join(ROOT, "建策BuildPlan_技术说明文档_v3.0.docx")

PARAS = {
    108: "5.5 展示粒度（六个组合直接选）",
    109: "计划表可以有粗有细。系统把「楼层分段」和「工序细度」两件事一次问清："
         "计划细度门直接列出六个组合选项（每种组合都带真实行数），"
         "敲一个数字就定完，不用先选维度再选档：",
    110: "•  楼层分段（三种施工段）：按层（最细，能逐层核对）/ 每 5 层一组（常用，表不长）"
         "/ 整栋（只做总控）。",
    111: "•  工序细度（两种）：工序级（细，每道工序一行）/ 工种级（粗，同一工种合并成一行）。",
    112: "六个编号：1 按层·工序级、2 按层·工种级、3 每 5 层一组·工序级、4 每 5 层一组·工种级、"
         "5 整栋·工序级、6 整栋·工种级；界面会先量出每种组合的真实行数再让你选。"
         "以某 12 栋项目（415 条工序叶子）为例：①415 行、②338 行、③116 行、④99 行、"
         "⑤25 行、⑥22 行——可见楼层分段才是行数的主杠杆。认不出楼层部位的任务"
         "（基础、室外等）归「分层外」，不参与楼层分段，所以「整栋」并不等于一行一个工序。",
}

CELLS = {
    "选展示粒度：工序拆解深度 × 楼层分组，界面先给出每种组合的真实行数":
        "选展示粒度：六个组合（三种楼层分段 × 两种工序细度），界面先给出每种组合的真实行数",
    "选「工序拆解深度 × 楼层分组」，界面给出每种组合的真实行数":
        "选「六个组合（楼层分段 × 工序细度）」，界面给出每种组合的真实行数",
    "展示粒度（两个维度）+ 真实行数": "展示粒度（六个组合）+ 真实行数",
    "按「工序拆解深度 × 楼层分组」分组、算真实行数、组内时间跨度；不改树不改量":
        "按「六个组合（楼层分段 × 工序细度）」分组、算真实行数、组内时间跨度；不改树不改量",
}


def set_text(par, new):
    if not par.runs:
        par.add_run(new)
        return
    par.runs[0].text = new
    for r in par.runs[1:]:
        r.text = ""


def main():
    doc = Document(DOC)
    for i, new in PARAS.items():
        set_text(doc.paragraphs[i], new)
    hit = 0
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                if c.text in CELLS:
                    set_text(c.paragraphs[0], CELLS[c.text])
                    hit += 1
    doc.save(DOC)
    print("段落改了 %d 段；表格单元格改了 %d 处（应 4）" % (len(PARAS), hit))


if __name__ == "__main__":
    main()
